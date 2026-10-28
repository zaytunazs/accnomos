"""Read the standard exports of Xero, QuickBooks Online, Sage 50, Sage Accounting, FreeAgent or a spreadsheet.

Five kinds of file feed the opening balances (app/import_history.py):

* **trial balance**: one row per account with Debit / Credit columns (or one signed balance column);
* **chart of accounts**: code, name and type per account, used to place each account (asset, liability, …);
* **contacts**: customers and suppliers with their details;
* **open items**: unpaid sales invoices and unpaid bills at the conversion date (the aged debtors / creditors);
* **bank balances**: the balance of each bank account at the conversion date.

Every vendor exports CSV or Excel with a few title rows above the header and a total row under it, so a file is read
as a grid of text (CSV, or the first sheet of an .xlsx), the header row is found by its column names, and every
field is looked up by a list of the names the vendors use for it (`FIELDS`). A file whose header has none of them is
refused with a sentence saying what we looked for.

What each vendor's help says the files contain (checked 28 Oct 2026 through web search summaries of the pages below;
their sites could not be opened from the build machine, so the exact header spellings are **not** confirmed against
a real export and the parser accepts every spelling listed in `FIELDS`; docs/HONESTY.md):

* Xero. Trial Balance report, Export > Excel: account code, account name, account type, Debit, Credit, YTD Debit,
  YTD Credit (https://tinytax.co.uk/support/export-xero-trial-balance, a third-party guide; Xero's own page is
  https://central.xero.com/s/article/Trial-Balance-report). Some layouts show the account as "Sales (200)": the code
  in brackets is split off. We prefer the YTD columns, which hold the balance at the report date. Chart of accounts
  export / import CSV: *Code, *Name, *Type, *Tax Code, Description (mandatory columns starred;
  https://central.xero.com/s/article/Import-a-chart-of-accounts). Account types such as BANK, CURRENT, FIXED,
  CURRLIAB, TERMLIAB, EQUITY, REVENUE, SALES, DIRECTCOSTS, EXPENSE, OVERHEADS, DEPRECIATN.
* QuickBooks Online. Reports > Trial Balance > Export to Excel (save as CSV): a first column of account names under
  an empty header, then Debit and Credit, with title rows above
  (https://quickbooks.intuit.com/learn-support/en-uk/reports-and-accounting/i-export-my-trial-balance-to-excel-to-prepare-management-reports/00/269621).
  Chart of accounts report: Account #, Account, Type, Detail type, Balance; sub-accounts written "Parent:Child"
  (https://quickbooks.intuit.com/learn-support/en-us/help-article/chart-accounts/import-chart-accounts/L9Res1eb1_US_en_US).
  Open Invoices report (Reports > Who owes you > Open Invoices, Export to Excel): Date, Transaction type, Num,
  Customer, Due date, Amount, Open balance.
* Sage 50 Accounts. Reports > Financials > Trial Balance > Export > To File (CSV): N/C, Name, Debit, Credit
  (https://tinytax.co.uk/guides/sage-ct600-integration). Sage 50's default nominal codes are the same four-digit
  layout as Accnomos (1100 Debtors control, 1200 Bank current account, 2100 Creditors control, 2200 Sales tax
  control, 3000 Capital, 4000 Sales, 5000 Purchases, 7xxx overheads), so its codes are kept as they are.
* Sage Accounting (cloud). Reporting > Trial Balance > Export CSV; Reporting > More > Aged Debtors / Aged Creditors >
  Detailed > Export CSV lists the unpaid invoices
  (https://help.sbc.sage.com/en-gb/accounting/reporting/financial_reports/extra-aged-creditors-report.html);
  Contacts > Customers / Suppliers > Export CSV with balances.
* FreeAgent. Accounting > Reports > Trial Balance > Export Report > Summary CSV (balance per account) or Chart of
  Accounts CSV; codes 001-599 are profit and loss, 600-999 balance sheet
  (https://support.freeagent.com/hc/en-gb/articles/115001222384,
  https://support.freeagent.com/hc/en-gb/articles/4412362325394).
* Spreadsheet: our own templates (`TEMPLATES`, served by `GET /import/templates/{kind}.csv`).

Money is read as pounds with up to two decimal places ("1,234.56", "£1,234.56", "(12.00)" and "12.00 CR" are
negative / credit, "-" is nothing) and kept as integer pence; a third decimal place is refused rather than rounded.
Dates are UK day-first ("31/03/2026", "31-Mar-2026", "31 Mar 2026") or ISO; Excel cells that are dates are read as
dates. Nothing in this module touches storage.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from collections.abc import Sequence
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Literal

from pydantic import BaseModel

MAX_FILE = 5 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLS = 80  # an .xlsx is read this wide at most: the exports have a dozen columns, a zip bomb has thousands
HEADER_SEARCH = 40  # rows scanned for the header (report exports have a title block above it)

FileKind = Literal["trial_balance", "chart_of_accounts", "contacts", "open_items", "bank_balances"]
Vendor = Literal["xero", "quickbooks", "sage50", "sage_accounting", "freeagent", "spreadsheet"]
Side = Literal["sale", "purchase"]
KINDS: tuple[FileKind, ...] = ("trial_balance", "chart_of_accounts", "contacts", "open_items", "bank_balances")
VENDORS: tuple[Vendor, ...] = ("xero", "quickbooks", "sage50", "sage_accounting", "freeagent", "spreadsheet")
VENDOR_NAMES: dict[Vendor, str] = {
    "xero": "Xero",
    "quickbooks": "QuickBooks Online",
    "sage50": "Sage 50 Accounts",
    "sage_accounting": "Sage Accounting",
    "freeagent": "FreeAgent",
    "spreadsheet": "a spreadsheet",
}


class ImportFormatError(ValueError):
    """The file cannot be read as the kind asked for. The message is shown to the user as is."""


def _norm(header: str) -> str:
    """'*Account Code ' -> 'account code'; 'N/C' -> 'n/c'; 'Account #' -> 'account #'."""
    return re.sub(r"\s+", " ", re.sub(r"[*:_.()\[\]]", " ", header or "")).strip().lower()


#: field -> the header spellings the vendors use for it (normalised by `_norm`).
FIELDS: dict[str, tuple[str, ...]] = {
    "code": (
        "code",
        "account code",
        "account #",
        "account no",
        "account number",
        "n/c",
        "nominal code",
        "nominal",
        "ledger account code",
        "nominal account",
        "a/c code",
        "gl code",
    ),
    "name": (
        "name",
        "account",
        "account name",
        "ledger account",
        "category",
        "description",
        "nominal name",
        "account description",
        "ledger account name",
    ),
    "type": (
        "type",
        "account type",
        "detail type",
        "category type",
        "class",
        "account class",
        "report type",
        "category group",
    ),
    "debit": ("debit", "debits", "debit £", "dr"),
    "credit": ("credit", "credits", "credit £", "cr"),
    "ytd_debit": ("ytd debit", "ytd debits", "year to date debit"),
    "ytd_credit": ("ytd credit", "ytd credits", "year to date credit"),
    "balance": ("balance", "closing balance", "total", "amount", "ytd balance", "net", "value", "balance £"),
    # contacts
    "contact": (
        "contactname",
        "contact name",
        "contact",
        "customer",
        "supplier",
        "vendor",
        "customer name",
        "supplier name",
        "vendor name",
        "company name",
        "company",
        "name",
        "organisation",
        "a/c name",
        "account name",
    ),
    "contact_ref": (
        "account reference",
        "a/c",
        "reference",
        "account ref",
        "customer ref",
        "supplier ref",
        "contact reference",
        "accountnumber",
        "account number",
    ),
    "email": ("email", "emailaddress", "email address", "e-mail", "main email", "contact email"),
    "phone": ("phone", "telephone", "phonenumber", "phone number", "mobile", "tel", "main phone"),
    "address1": (
        "address",
        "address line 1",
        "poaddressline1",
        "street1",
        "street",
        "billing address",
        "address 1",
        "street address",
        "billing street",
    ),
    "address2": ("address line 2", "poaddressline2", "street2", "address 2"),
    "town": ("town", "city", "pocity", "town/city", "billing city", "address line 3", "address 3"),
    "postcode": (
        "postcode",
        "post code",
        "popostalcode",
        "postal code",
        "zip",
        "billing postcode",
        "billing zip",
        "address line 5",
        "address 5",
    ),
    "vat_number": (
        "vat number",
        "taxnumber",
        "tax number",
        "vat reg no",
        "vat registration number",
        "vat no",
        "vat registration",
        "tax registration number",
    ),
    "contact_kind": ("contact type", "kind", "is customer", "role", "customer or supplier"),
    "terms_days": ("terms", "payment terms", "credit terms", "terms days", "due date terms"),
    # open items
    "number": (
        "invoicenumber",
        "invoice number",
        "num",
        "no",
        "number",
        "invoice no",
        "reference",
        "ref",
        "document number",
        "bill number",
        "invoice",
        "transaction number",
        "inv ref",
        "doc no",
    ),
    "doc_date": (
        "invoicedate",
        "invoice date",
        "date",
        "transaction date",
        "issue date",
        "dated",
        "bill date",
        "document date",
    ),
    "due_date": ("duedate", "due date", "due", "due on", "payment due"),
    "total": (
        "total",
        "amount",
        "invoice total",
        "gross",
        "invoice amount",
        "original amount",
        "total amount",
        "gross amount",
    ),
    "due_amount": (
        "invoiceamountdue",
        "amountdue",
        "amount due",
        "open balance",
        "outstanding",
        "balance",
        "amount outstanding",
        "outstanding amount",
        "balance due",
        "due amount",
        "unpaid",
        "remaining",
    ),
    "item_type": ("type", "transaction type", "document type", "invoice type", "doc type"),
    # bank balances
    "bank_account": ("bank account", "account", "account name", "bank", "name"),
    "bank_balance": ("balance", "closing balance", "statement balance", "balance at date", "amount", "balance £"),
    "bank_date": ("date", "balance date", "as at", "statement date"),
}

#: Columns that, when present, decide which kind a file is (with the fields above).
_KIND_NEEDS: dict[FileKind, tuple[tuple[str, ...], ...]] = {
    # Each inner tuple: any one of these fields. All tuples must match.
    "trial_balance": (("name", "code"), ("debit", "credit", "ytd_debit", "ytd_credit", "balance")),
    "chart_of_accounts": (("name",), ("type",)),
    "contacts": (("contact",), ("email", "phone", "address1", "vat_number", "postcode", "town")),
    "open_items": (("contact",), ("number", "doc_date"), ("due_amount", "total")),
    "bank_balances": (("bank_account",), ("bank_balance",)),
}

TEMPLATES: dict[FileKind, str] = {
    "trial_balance": "Code,Account,Type,Debit,Credit\n1200,Bank current account,Bank,1250.00,\n",
    "chart_of_accounts": "Code,Account,Type\n1200,Bank current account,Bank\n",
    "contacts": "Name,Kind,Email,Phone,Address,Postcode,VAT number,Terms\nSample Customer Ltd,customer,,,,,,30\n",
    "open_items": "Type,Contact,Number,Date,Due date,Total,Amount due\ninvoice,Sample Customer Ltd,INV-001,"
    "01/03/2026,31/03/2026,120.00,120.00\n",
    "bank_balances": "Account,Balance,Date\nBank current account,1250.00,31/03/2026\n",
}

# --------------------------------------------------------------------------- #
# Cells
# --------------------------------------------------------------------------- #

_MONEY_JUNK = re.compile(r"[£$€,\s ]")
_MONTHS = {
    m: i
    for i, m in enumerate(("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), start=1)
}


def pence(raw: str | None) -> int:
    """'1,234.56' -> 123456; '(12.00)', '-12', '12.00 CR', '12.00-' -> -1200; '' / '-' -> 0.

    Raises ImportFormatError for anything else, and for more than two decimal places (we never round money)."""
    s = (raw or "").strip()
    if s in ("", "-", "–", "—"):
        return 0
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1]
    upper = s.upper()
    if upper.endswith("CR"):
        negative, s = not negative, s[:-2]
    elif upper.endswith("DR"):
        s = s[:-2]
    s = _MONEY_JUNK.sub("", s)
    if s.endswith("-"):
        negative, s = not negative, s[:-1]
    if s.startswith("-"):
        negative, s = not negative, s[1:]
    if s.startswith("+"):
        s = s[1:]
    try:
        value = Decimal(s)
    except InvalidOperation as exc:
        raise ImportFormatError(f"'{raw}' is not an amount") from exc
    if not value.is_finite():
        raise ImportFormatError(f"'{raw}' is not an amount")
    exponent = value.as_tuple().exponent
    if isinstance(exponent, int) and exponent < -2 and value != value.quantize(Decimal("0.01")):
        raise ImportFormatError(f"'{raw}' has more than two decimal places; export it in pounds and pence")
    p = int(value * 100)
    return -p if negative else p


def parse_date(raw: str | None) -> date | None:
    """UK day-first or ISO; None for an empty cell. Raises ImportFormatError for anything else."""
    s = (raw or "").strip()
    if not s:
        return None
    if m := re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[T ].*)?", s):
        y, mo, d = (int(g) for g in m.groups())
    elif m := re.fullmatch(r"(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2}|\d{4})", s):
        d, mo, y = (int(g) for g in m.groups())
    elif m := re.fullmatch(r"(\d{1,2})[\s\-/]([A-Za-z]{3,9})[\s\-/,]+(\d{2}|\d{4})", s):
        d, y = int(m.group(1)), int(m.group(3))
        mo = _MONTHS.get(m.group(2)[:3].lower(), 0)
    else:
        raise ImportFormatError(f"'{raw}' is not a date (we read day/month/year)")
    if y < 100:
        y += 2000
    try:
        return date(y, mo, d)
    except ValueError as exc:
        raise ImportFormatError(f"'{raw}' is not a date (we read day/month/year)") from exc


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return f"{Decimal(repr(v)):f}" if v != int(v) else str(int(v))
    return str(v).strip()


# --------------------------------------------------------------------------- #
# Files -> grid
# --------------------------------------------------------------------------- #


def read_grid(data: bytes, name: str) -> list[list[str]]:
    """The file's rows as text cells: CSV (comma, semicolon or tab) or the first sheet of an .xlsx workbook."""
    if len(data) > MAX_FILE:
        raise ImportFormatError("The file is over 5 MB; export a shorter date range or one report at a time")
    if data[:2] == b"PK":
        return _xlsx(data)
    if name.lower().endswith((".xls", ".xlsm", ".ods", ".pdf")):
        raise ImportFormatError("Save the export as CSV or .xlsx and upload that")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1252 decodes almost anything
        raise ImportFormatError("The file is not text; save the export as CSV")
    sample = text[:4096]
    delimiter = max((",", ";", "\t"), key=sample.count)
    rows: list[list[str]] = []
    try:
        for row in csv.reader(io.StringIO(text), delimiter=delimiter):
            rows.append([c.strip() for c in row[:MAX_COLS]])
            if len(rows) > MAX_ROWS:
                raise ImportFormatError(f"The file has over {MAX_ROWS:,} rows")
    except csv.Error as exc:
        raise ImportFormatError(f"The file is not a CSV we can read ({exc})") from exc
    return rows


def _xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any unreadable workbook is the user's file, not our bug
        raise ImportFormatError("The workbook cannot be opened; save the export as CSV") from exc
    try:
        ws = wb.worksheets[0]
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True, max_col=MAX_COLS):
            rows.append([_cell(v) for v in row])
            if len(rows) > MAX_ROWS:
                raise ImportFormatError(f"The file has over {MAX_ROWS:,} rows")
        return rows
    finally:
        wb.close()


# --------------------------------------------------------------------------- #
# Header
# --------------------------------------------------------------------------- #


def _columns(header: Sequence[str], fields: Sequence[str]) -> dict[str, int]:
    """field -> column index, each column used once, the earliest-listed field winning a column it shares."""
    names = [_norm(h) for h in header]
    out: dict[str, int] = {}
    used: set[int] = set()
    for f in fields:
        for spelling in FIELDS[f]:
            idx = next((i for i, n in enumerate(names) if n == spelling and i not in used), None)
            if idx is not None:
                out[f] = idx
                used.add(idx)
                break
    return out


_KIND_FIELDS: dict[FileKind, tuple[str, ...]] = {
    "trial_balance": ("code", "ytd_debit", "ytd_credit", "debit", "credit", "type", "name", "balance"),
    "chart_of_accounts": ("code", "type", "name", "balance"),
    "contacts": (
        "contact_ref",
        "email",
        "phone",
        "address1",
        "address2",
        "town",
        "postcode",
        "vat_number",
        "contact_kind",
        "terms_days",
        "contact",
    ),
    "open_items": ("item_type", "number", "doc_date", "due_date", "due_amount", "total", "contact"),
    "bank_balances": ("bank_balance", "bank_date", "bank_account"),
}


def _fits(cols: dict[str, int], kind: FileKind) -> bool:
    return all(any(f in cols for f in group) for group in _KIND_NEEDS[kind])


def find_header(rows: list[list[str]], kind: FileKind) -> tuple[int, dict[str, int]]:
    """(row index, field -> column) of the first row in the top `HEADER_SEARCH` that names the kind's columns."""
    for i, row in enumerate(rows[:HEADER_SEARCH]):
        if not any(row):
            continue
        cols = _columns(row, _KIND_FIELDS[kind])
        if (
            kind == "trial_balance"
            and "name" not in cols
            and "code" not in cols
            and any(f in cols for f in ("debit", "credit"))
        ):
            # QuickBooks: the account column has no header; take the first column without one.
            blank = next((j for j, h in enumerate(row) if not h.strip()), None)
            if blank is not None:
                cols["name"] = blank
        if _fits(cols, kind):
            return i, cols
    wanted = " and ".join("/".join(FIELDS[group[0]][:3]) for group in _KIND_NEEDS[kind])
    raise ImportFormatError(
        f"No header row found: we looked for columns like {wanted} in the first {HEADER_SEARCH} rows"
    )


def detect_kind(rows: list[list[str]]) -> FileKind | None:
    """The kind whose header appears first, best match first: open items and trial balance before the looser ones."""
    order: tuple[FileKind, ...] = ("open_items", "trial_balance", "chart_of_accounts", "contacts", "bank_balances")
    hits: list[tuple[int, int, FileKind]] = []
    for rank, kind in enumerate(order):
        try:
            idx, cols = find_header(rows, kind)
        except ImportFormatError:
            continue
        if kind == "trial_balance" and not any(f in cols for f in ("debit", "credit", "ytd_debit", "ytd_credit")):
            if "type" in cols:
                hits.append((idx, rank + 1, "chart_of_accounts"))  # a CoA with a balance column (QuickBooks)
                continue
            if "code" not in cols:  # "Account, Balance": a bank balances list
                continue
        hits.append((idx, rank, kind))
    return min(hits)[2] if hits else None


def detect_vendor(rows: list[list[str]]) -> Vendor | None:
    """A guess from tell-tale headers; the user picks the vendor anyway."""
    heads = {_norm(c) for row in rows[:HEADER_SEARCH] for c in row if c}
    text = " ".join(heads)
    if {"code", "name", "type", "tax code"} <= heads or "ytd debit" in heads or "contactname" in heads:
        return "xero"
    if "detail type" in heads or "open balance" in heads or "account #" in heads or "quickbooks" in text:
        return "quickbooks"
    if "n/c" in heads or "a/c" in heads:
        return "sage50"
    if "ledger account" in heads or "ledger account code" in heads:
        return "sage_accounting"
    if "nominal code" in heads and "category" in heads or "freeagent" in text:
        return "freeagent"
    return None


# --------------------------------------------------------------------------- #
# Rows
# --------------------------------------------------------------------------- #


class SourceAccount(BaseModel):
    code: str | None = None
    name: str
    type: str | None = None
    debit_pence: int = 0
    credit_pence: int = 0

    @property
    def balance_pence(self) -> int:
        """Debit positive, credit negative."""
        return self.debit_pence - self.credit_pence


class ContactRow(BaseModel):
    name: str
    ref: str | None = None
    kind: Literal["customer", "supplier", "both"] | None = None
    email: str | None = None
    phone: str | None = None
    address: str | None = None
    vat_number: str | None = None
    terms_days: int | None = None


class OpenItemRow(BaseModel):
    side: Side
    contact: str
    number: str | None = None
    date: date
    due_date: date | None = None
    total_pence: int  # the document's total
    due_pence: int  # still owed at the conversion date


class BankBalanceRow(BaseModel):
    account: str
    balance_pence: int
    date: dt.date | None = None


class Parsed(BaseModel):
    kind: FileKind
    vendor: Vendor | None = None
    accounts: list[SourceAccount] = []
    contacts: list[ContactRow] = []
    items: list[OpenItemRow] = []
    banks: list[BankBalanceRow] = []
    skipped: list[str] = []  # one sentence per row we left out, with its row number


_TOTAL = re.compile(r"^(grand\s+)?totals?\b|^net\s+(profit|loss)\b|^difference\b", re.I)
_CODE_IN_BRACKETS = re.compile(r"^(.*\S)\s*\(([0-9][\w.-]{0,9})\)$")  # Xero: "Sales (200)"
_LEADING_CODE = re.compile(r"^(\d{1,6})\s*[-–:]?\s+(\S.*)$")  # "1200 Bank", "090 - Business Bank"


def _get(row: list[str], cols: dict[str, int], field: str) -> str:
    i = cols.get(field)
    return row[i].strip() if i is not None and i < len(row) else ""


def split_code(name: str, code: str | None) -> tuple[str | None, str]:
    """('Sales (200)', None) -> ('200', 'Sales'); ('1200 Bank', None) -> ('1200', 'Bank')."""
    name = name.strip()
    if code:
        return code.strip(), name
    if m := _CODE_IN_BRACKETS.match(name):
        return m.group(2), m.group(1)
    if m := _LEADING_CODE.match(name):
        return m.group(1), m.group(2)
    return None, name


def _accounts(rows: list[list[str]], start: int, cols: dict[str, int], with_amounts: bool, out: Parsed) -> None:
    ytd = "ytd_debit" in cols or "ytd_credit" in cols
    for n, row in enumerate(rows[start:], start=start + 1):
        raw_name, raw_code = _get(row, cols, "name"), _get(row, cols, "code")
        if not raw_name and not raw_code:
            continue
        if _TOTAL.match(raw_name) or _TOTAL.match(raw_code):
            continue
        code, name = split_code(raw_name or raw_code, raw_code or None)
        if not name:
            name = code or ""
        try:
            if not with_amounts:
                debit = credit = 0
            elif ytd:
                debit, credit = pence(_get(row, cols, "ytd_debit")), pence(_get(row, cols, "ytd_credit"))
            elif "debit" in cols or "credit" in cols:
                debit, credit = pence(_get(row, cols, "debit")), pence(_get(row, cols, "credit"))
            else:
                balance = pence(_get(row, cols, "balance"))
                debit, credit = max(balance, 0), max(-balance, 0)
        except ImportFormatError as exc:
            out.skipped.append(f"Row {n} ({name}): {exc}")
            continue
        # A negative debit is a credit, and the other way round (some exports sign one column).
        if debit < 0:
            debit, credit = 0, credit - debit
        if credit < 0:
            debit, credit = debit - credit, 0
        net = debit - credit
        if with_amounts and net == 0 and debit == 0 and not code:
            continue  # a section heading ("Assets", "Current liabilities") with no amount
        out.accounts.append(
            SourceAccount(
                code=code,
                name=name[:120],
                type=_get(row, cols, "type") or None,
                debit_pence=max(net, 0),
                credit_pence=max(-net, 0),
            )
        )


_SALE_WORDS = ("invoice", "sales", "accrec", "customer", "receivable", "si", "sale")
_PURCHASE_WORDS = ("bill", "bills", "purchase", "accpay", "supplier", "vendor", "payable", "pi")


def _side(raw: str, header_side: Side | None, default: Side | None) -> Side | None:
    """'Invoice', 'ACCREC', 'SI' -> sale; 'Bill', 'ACCPAY', 'PI' -> purchase; a credit note or payment -> None."""
    words = set(re.split(r"[^a-z]+", raw.strip().lower())) - {""}
    if words:
        if words & {"credit", "payment", "sc", "pc", "sr", "pr", "sa", "pa", "receipt", "journal"}:
            return None  # not an invoice or bill: left out (it is already netted off in the totals we read)
        if words & set(_PURCHASE_WORDS):
            return "purchase"
        if words & set(_SALE_WORDS):
            return "sale"
    return header_side or default


def _items(rows: list[list[str]], start: int, cols: dict[str, int], default: Side | None, out: Parsed) -> None:
    header = [_norm(h) for h in rows[start - 1]] if start else []
    contact_head = header[cols["contact"]] if "contact" in cols and cols["contact"] < len(header) else ""
    header_side: Side | None = (
        "sale"
        if contact_head.startswith("customer")
        else "purchase"
        if contact_head.startswith(("supplier", "vendor"))
        else None
    )
    for n, row in enumerate(rows[start:], start=start + 1):
        contact = _get(row, cols, "contact")
        if not contact or _TOTAL.match(contact) or _TOTAL.match(_get(row, cols, "number")):
            continue
        raw_type = _get(row, cols, "item_type")
        side = _side(raw_type, header_side, default)
        if side is None:
            if raw_type:
                out.skipped.append(f"Row {n} ({contact}): a '{raw_type}' is not an invoice or bill; left out")
            else:
                out.skipped.append(f"Row {n} ({contact}): say whether these are sales invoices or bills")
            continue
        try:
            when = parse_date(_get(row, cols, "doc_date"))
            due = parse_date(_get(row, cols, "due_date"))
            total = pence(_get(row, cols, "total")) if "total" in cols else None
            owed = pence(_get(row, cols, "due_amount")) if "due_amount" in cols else total
        except ImportFormatError as exc:
            out.skipped.append(f"Row {n} ({contact}): {exc}")
            continue
        if when is None:
            out.skipped.append(f"Row {n} ({contact}): no date")
            continue
        owed = abs(owed or 0)
        total = abs(total) if total is not None else owed
        if owed == 0:
            continue  # paid in full: nothing to carry over
        out.items.append(
            OpenItemRow(
                side=side,
                contact=contact[:120],
                number=(_get(row, cols, "number") or None),
                date=when,
                due_date=due,
                total_pence=max(total, owed),
                due_pence=owed,
            )
        )


def _contact_kind(raw: str) -> Literal["customer", "supplier", "both"] | None:
    t = raw.strip().lower()
    if not t:
        return None
    if "both" in t or ("customer" in t and "supplier" in t):
        return "both"
    if t.startswith(("supplier", "vendor", "creditor")) or t in ("s", "purchase"):
        return "supplier"
    if t.startswith(("customer", "client", "debtor")) or t in ("c", "sale", "yes", "true"):
        return "customer"
    return None


def _contacts(rows: list[list[str]], start: int, cols: dict[str, int], out: Parsed) -> None:
    for row in rows[start:]:
        name = _get(row, cols, "contact")
        if not name or _TOTAL.match(name):
            continue
        address = ", ".join(p for p in (_get(row, cols, f) for f in ("address1", "address2", "town", "postcode")) if p)
        terms_raw = re.sub(r"\D", "", _get(row, cols, "terms_days"))
        terms = int(terms_raw) if terms_raw and int(terms_raw) <= 365 else None
        out.contacts.append(
            ContactRow(
                name=name[:120],
                ref=_get(row, cols, "contact_ref") or None,
                kind=_contact_kind(_get(row, cols, "contact_kind")),
                email=_get(row, cols, "email") or None,
                phone=_get(row, cols, "phone") or None,
                address=address[:300] or None,
                vat_number=_get(row, cols, "vat_number") or None,
                terms_days=terms,
            )
        )


def _banks(rows: list[list[str]], start: int, cols: dict[str, int], out: Parsed) -> None:
    for n, row in enumerate(rows[start:], start=start + 1):
        name = _get(row, cols, "bank_account")
        if not name or _TOTAL.match(name):
            continue
        try:
            out.banks.append(
                BankBalanceRow(
                    account=name[:120],
                    balance_pence=pence(_get(row, cols, "bank_balance")),
                    date=parse_date(_get(row, cols, "bank_date")),
                )
            )
        except ImportFormatError as exc:
            out.skipped.append(f"Row {n} ({name}): {exc}")


def parse(
    data: bytes, name: str, kind: FileKind | None = None, vendor: Vendor | None = None, side: Side | None = None
) -> Parsed:
    """Read one export. `kind` None: worked out from the header. `side` is the default for open items whose rows do
    not say whether they are invoices or bills. Raises ImportFormatError with a sentence for the user."""
    rows = read_grid(data, name)
    if not any(any(c for c in r) for r in rows):
        raise ImportFormatError("The file is empty")
    kind = kind or detect_kind(rows)
    if kind is None:
        raise ImportFormatError(
            "We could not tell what this file is: pick trial balance, chart of accounts, contacts, open invoices "
            "and bills, or bank balances, and upload it again"
        )
    idx, cols = find_header(rows, kind)
    out = Parsed(kind=kind, vendor=vendor or detect_vendor(rows))
    start = idx + 1
    if kind == "trial_balance":
        _accounts(rows, start, cols, True, out)
    elif kind == "chart_of_accounts":
        _accounts(rows, start, cols, False, out)
    elif kind == "open_items":
        _items(rows, start, cols, side, out)
    elif kind == "contacts":
        _contacts(rows, start, cols, out)
    else:
        _banks(rows, start, cols, out)
    if not (out.accounts or out.contacts or out.items or out.banks):
        raise ImportFormatError(
            "No rows could be read under the header" + (f": {out.skipped[0]}" if out.skipped else "")
        )
    return out
