"""Payroll reports as data, and one renderer per format (CSV, XLSX, PDF; the web page shows the same rows as JSON).

Each report is a pure function returning a `Table` (columns, rows, totals, notes); `render(table, fmt)` turns any of
them into a file. That is the shape the shared export registry (branch claude/exports) can register later: a report
key, a function to its rows, and per-format renderers. Nothing here duplicates that framework; when it lands, register
`REPORTS` with it.

HMRC-style reports (laid out like the forms they are named after, not the forms themselves):
    payroll_summary   one pay run, a row per employee
    p11               one employee's period-by-period record for the year (P11-style deductions working sheet)
    p32               the employer payment record: each tax month's bill to HMRC (P32-style)
    pension           pension workings for one run: assessment, band, pensionable pay, contributions, relief basis
    fps               the FPS contents as readable rows, next to the XML
    journal           the run's ledger journal, line by line (for accounting-package journal exports)
    p45 / p60         the figures for a leaver's P45 and a year-end P60
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Any, Literal

from app import paye as P
from app import paye_rates as R
from app import pension_ae as AE
from app.payrun import ESTIMATE_NOTE, FILING_NOTE, Employee, PayLine, PayRun, journal
from app.rti import MonthRecord, P45Figures, P60Figures, readable_fps

Kind = Literal["text", "money", "int", "date"]
Format = Literal["csv", "xlsx", "pdf", "json"]
FORMATS: tuple[Format, ...] = ("csv", "xlsx", "pdf", "json")


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    kind: Kind = "text"
    weight: float | None = None  # relative PDF column width (numbers 1.2, text 1.6 by default)


@dataclass
class Table:
    key: str
    title: str
    subtitle: str
    columns: list[Column]
    rows: list[dict[str, Any]]
    totals: dict[str, Any] | None = None
    notes: list[str] = field(default_factory=list)
    landscape: bool = True

    def as_json(self) -> dict[str, Any]:
        return {
            "key": self.key,
            "title": self.title,
            "subtitle": self.subtitle,
            "columns": [{"key": c.key, "label": c.label, "kind": c.kind} for c in self.columns],
            "rows": self.rows,
            "totals": self.totals,
            "notes": self.notes,
        }


BASE_NOTES = [ESTIMATE_NOTE + f": {R.TAX_YEAR} rates verified on {R.VERIFIED_ON} ({R.RATES_URL}).", FILING_NOTE]


def _sum(rows: list[dict[str, Any]], cols: list[Column]) -> dict[str, Any]:
    return {c.key: sum(int(r.get(c.key) or 0) for r in rows) for c in cols if c.kind == "money"}


def _period_label(run: PayRun) -> str:
    return f"Month {run.period}" if run.frequency == "monthly" else f"Week {run.period}"


# --------------------------------------------------------------------------- #
# Reports
# --------------------------------------------------------------------------- #

SUMMARY_COLS = [
    Column("payroll_id", "Payroll ID"),
    Column("name", "Employee"),
    Column("tax_code", "Tax code"),
    Column("ni_category", "NI"),
    Column("overtime", "Overtime", "money"),
    Column("gross", "Gross pay", "money"),
    Column("sacrifice", "Salary sacrifice", "money"),
    Column("taxable", "Taxable pay", "money"),
    Column("tax", "PAYE tax", "money"),
    Column("ee_ni", "Employee NI", "money"),
    Column("er_ni", "Employer NI", "money"),
    Column("student_loan", "Student loan", "money"),
    Column("postgrad_loan", "Postgrad loan", "money"),
    Column("ee_pension", "Employee pension", "money"),
    Column("er_pension", "Employer pension", "money"),
    Column("other", "Other deductions", "money"),
    Column("net", "Net pay", "money"),
]


def _summary_row(ln: PayLine) -> dict[str, Any]:
    return {
        "payroll_id": ln.payroll_id,
        "name": ln.name,
        "tax_code": ln.tax_code,
        "ni_category": ln.ni_category,
        "overtime": ln.overtime_pence,
        "gross": ln.gross_pence,
        "sacrifice": ln.salary_sacrifice_pence,
        "taxable": ln.taxable_pay_pence,
        "tax": ln.tax_pence,
        "ee_ni": ln.ee_ni_pence,
        "er_ni": ln.er_ni_pence,
        "student_loan": ln.student_loan_pence,
        "postgrad_loan": ln.postgrad_loan_pence,
        "ee_pension": ln.ee_pension_net_pay_pence + ln.ee_pension_from_net_pence,
        "er_pension": ln.pension.employer_pence,
        "other": ln.deductions_after_tax_pence,
        "net": ln.net_pay_pence,
    }


def payroll_summary(run: PayRun) -> Table:
    rows = [_summary_row(ln) for ln in run.lines]
    notes = [
        f"Employment Allowance set against employer NI: {P.gbp(run.employment_allowance_pence)}.",
        f"Due to HMRC from this run: {P.gbp(run.hmrc_due_pence)} "
        "(the month's total is on the employer payment record).",
        *BASE_NOTES,
    ]
    return Table(
        "payroll_summary",
        f"Payroll summary {run.id}",
        f"{run.frequency.replace('_', '-')} pay, {_period_label(run)}, paid {run.pay_date:%d/%m/%Y} ({run.status})",
        SUMMARY_COLS,
        rows,
        _sum(rows, SUMMARY_COLS),
        notes,
    )


P11_COLS = [
    Column("pay_date", "Pay date", "date"),
    Column("period", "Period"),
    Column("tax_code", "Tax code"),
    Column("overtime", "Overtime", "money"),
    Column("gross", "Gross pay", "money"),
    Column("taxable", "Taxable pay", "money"),
    Column("taxable_ytd", "Taxable pay to date", "money"),
    Column("tax", "Tax", "money"),
    Column("tax_ytd", "Tax to date", "money"),
    Column("ni_category", "NI"),
    Column("ni_earnings", "Earnings for NI", "money"),
    Column("at_lel", "At LEL", "money"),
    Column("lel_to_pt", "LEL to PT", "money"),
    Column("pt_to_uel", "PT to UEL", "money"),
    Column("ee_ni", "Employee NI", "money"),
    Column("er_ni", "Employer NI", "money"),
    Column("student_loan", "Student loan", "money"),
    Column("postgrad_loan", "Postgrad loan", "money"),
    Column("ee_pension", "Employee pension", "money"),
    Column("er_pension", "Employer pension", "money"),
    Column("ssp", "SSP", "money"),
    Column("smp", "SMP", "money"),
    Column("net", "Net pay", "money"),
]


def p11(emp: Employee, runs: list[PayRun]) -> Table:
    """One employee's posted pay runs this tax year, period by period, with year-to-date tax."""
    rows = []
    for run in runs:
        if run.status != "posted":
            continue
        for ln in run.lines:
            if ln.employee_id != emp.id:
                continue
            rows.append(
                {
                    "pay_date": run.pay_date.isoformat(),
                    "period": _period_label(run),
                    "tax_code": ln.tax_code,
                    "overtime": ln.overtime_pence,
                    "gross": ln.gross_pence,
                    "taxable": ln.taxable_pay_pence,
                    "taxable_ytd": ln.ytd.taxable_pay_pence,
                    "tax": ln.tax_pence,
                    "tax_ytd": ln.ytd.tax_pence,
                    "ni_category": ln.ni_category,
                    "ni_earnings": ln.ni_earnings_pence,
                    "at_lel": ln.ni_at_lel_pence,
                    "lel_to_pt": ln.ni_lel_to_pt_pence,
                    "pt_to_uel": ln.ni_pt_to_uel_pence,
                    "ee_ni": ln.ee_ni_pence,
                    "er_ni": ln.er_ni_pence,
                    "student_loan": ln.student_loan_pence,
                    "postgrad_loan": ln.postgrad_loan_pence,
                    "ee_pension": ln.ee_pension_net_pay_pence + ln.ee_pension_from_net_pence,
                    "er_pension": ln.pension.employer_pence,
                    "ssp": ln.ssp_pence,
                    "smp": ln.smp_pence,
                    "net": ln.net_pay_pence,
                }
            )
    totals = _sum(rows, P11_COLS)
    for k in ("taxable_ytd", "tax_ytd"):
        totals[k] = rows[-1][k] if rows else 0
    notes = list(BASE_NOTES)
    if emp.p45_previous_pay_pence is not None:
        notes.insert(
            0,
            f"Previous employment (P45): pay {P.gbp(emp.p45_previous_pay_pence)}, tax "
            f"{P.gbp(emp.p45_previous_tax_pence or 0)}; used in the tax when the code is cumulative.",
        )
    if emp.brought_forward:
        notes.insert(0, "Includes figures brought forward from the payroll used before Accnomos this year.")
    return Table(
        "p11",
        f"Employee payroll record (P11-style): {emp.name}",
        f"Payroll ID {emp.payroll_id}, NI number {emp.nino or 'not known'}, tax year {R.TAX_YEAR}",
        P11_COLS,
        rows,
        totals,
        notes,
    )


P32_COLS = [
    Column("tax_month", "Tax month", "int"),
    Column("dates", "Dates"),
    Column("tax", "PAYE tax", "money"),
    Column("student_loan", "Student loans", "money"),
    Column("postgrad_loan", "Postgrad loans", "money"),
    Column("ni", "Total NI (employee + employer)", "money"),
    Column("smp_recovered", "SMP recovered", "money"),
    Column("nic_compensation", "NIC compensation", "money"),
    Column("employment_allowance", "Employment Allowance", "money"),
    Column("due", "Due to HMRC", "money"),
    Column("due_by", "Pay by (electronic)", "date"),
]


def p32(months: list[MonthRecord]) -> Table:
    rows = [
        {
            "tax_month": m.tax_month,
            "dates": f"{m.start:%d/%m/%Y}–{m.end:%d/%m/%Y}",
            "tax": m.tax_pence,
            "student_loan": m.student_loan_pence,
            "postgrad_loan": m.postgrad_loan_pence,
            "ni": m.gross_ni_pence,
            "smp_recovered": m.smp_recovered_pence,
            "nic_compensation": m.nic_compensation_pence,
            "employment_allowance": m.employment_allowance_pence,
            "due": m.due_pence,
            "due_by": m.due.isoformat(),
        }
        for m in months
    ]
    notes = [
        "Pay HMRC by the 22nd of the next tax month electronically (the 19th by post): https://www.gov.uk/pay-paye-tax.",
        "A negative amount is what HMRC owes you; it is set against later months.",
        "Months with no pay runs need an EPS saying no payment is due.",
        *BASE_NOTES,
    ]
    return Table(
        "p32",
        "Employer payment record (P32-style)",
        f"Tax year {R.TAX_YEAR}, posted pay runs only",
        P32_COLS,
        rows,
        _sum(rows, P32_COLS),
        notes,
    )


PENSION_COLS = [
    Column("name", "Employee"),
    Column("age", "Age", "int"),
    Column("category", "Assessment", weight=2.4),
    Column("member", "Member", weight=1.0),
    Column("enrolled_on", "Enrolled on", "date"),
    Column("opt_out_ends", "Opt-out window ends", "date"),
    Column("band", "Earnings band used", weight=2.2),
    Column("gross", "Earnings this period", "money"),
    Column("pensionable", "Pensionable pay", "money"),
    Column("employer", "Employer contribution", "money"),
    Column("employee_gross", "Employee contribution (gross)", "money"),
    Column("employee_deducted", "Taken from pay", "money"),
    Column("relief", "Tax relief basis", weight=2.2),
    Column("scheme_relief", "Relief claimed by scheme", "money"),
    Column("sacrifice", "Salary sacrificed", "money"),
    Column("to_provider", "Paid to provider", "money"),
    Column("notes", "Notes", weight=4.0),
]


def pension_workings(run: PayRun) -> Table:
    rows = []
    for ln in run.lines:
        p = ln.pension
        band = (
            f"{P.gbp(p.band_low_pence)}–{P.gbp(p.band_high_pence)} (qualifying earnings)"
            if p.basis == "qualifying"
            else "All pay (full-pay basis)"
        )
        rows.append(
            {
                "name": ln.name,
                "age": p.age,
                "category": AE.CATEGORY_LABELS[p.category],
                "member": "Yes" if p.member else "No",
                "enrolled_on": p.enrolled_on.isoformat() if p.enrolled_on else "",
                "opt_out_ends": p.opt_out_ends.isoformat() if p.opt_out_ends else "",
                "band": band,
                "gross": ln.gross_pence,
                "pensionable": p.pensionable_pence,
                "employer": p.employer_pence,
                "employee_gross": p.employee_gross_pence,
                "employee_deducted": p.employee_deducted_pence,
                "relief": AE.ARRANGEMENT_LABELS[p.arrangement] if p.member else "",
                "scheme_relief": p.relief_claimed_by_scheme_pence,
                "sacrifice": p.sacrifice_pence,
                "to_provider": p.to_provider_pence,
                "notes": "; ".join(p.notes),
            }
        )
    th = R.AE_THRESHOLDS[run.frequency]
    notes = [
        f"{run.frequency.replace('_', '-')} thresholds {R.TAX_YEAR}: lower level of qualifying earnings "
        f"{P.gbp(th.lower_qe)}, earnings trigger {P.gbp(th.trigger)}, upper level {P.gbp(th.upper_qe)} "
        f"({R.TPR_THRESHOLDS_URL}).",
        "Minimum on qualifying earnings: employer 3%, total 8%. Eligible jobholders: aged 22 to State Pension age, "
        "earning more than the trigger. Opt-out: within one month of enrolment, contributions refunded.",
        "Contributions are rounded to the nearest penny; check your provider's rounding.",
        *BASE_NOTES[:1],
    ]
    return Table(
        "pension",
        f"Pension workings {run.id}",
        f"Pay date {run.pay_date:%d/%m/%Y}",
        PENSION_COLS,
        rows,
        _sum(rows, PENSION_COLS),
        notes,
    )


FPS_COLS = [Column("employee", "Employee"), Column("item", "Data item"), Column("value", "Value")]


def fps_contents(run: PayRun, employees: dict[str, Employee]) -> Table:
    return Table(
        "fps",
        f"Full Payment Submission contents {run.id}",
        f"What the FPS XML says, in words. Pay date {run.pay_date:%d/%m/%Y}",
        FPS_COLS,
        readable_fps(run, employees),
        None,
        BASE_NOTES[::-1],
        landscape=False,
    )


JOURNAL_COLS = [
    Column("date", "Date", "date"),
    Column("reference", "Reference"),
    Column("account", "Account"),
    Column("description", "Description"),
    Column("debit", "Debit", "money"),
    Column("credit", "Credit", "money"),
]


def journal_rows(run: PayRun) -> Table:
    """The run's journal (what approval posts), one row per account: the rows an accounting-package journal export
    maps to its own columns."""
    rows = [
        {
            "date": run.pay_date.isoformat(),
            "reference": f"PAYROLL-{run.id}",
            "account": j.account,
            "description": f"Payroll {run.id}",
            "debit": j.debit_pence,
            "credit": j.credit_pence,
        }
        for j in journal(run)
    ]
    return Table(
        "journal",
        f"Payroll journal {run.id}",
        f"Pay date {run.pay_date:%d/%m/%Y}; posts when you approve the pay run in Review",
        JOURNAL_COLS,
        rows,
        _sum(rows, JOURNAL_COLS),
        BASE_NOTES[:1],
        landscape=False,
    )


FORM_COLS = [Column("item", "Item"), Column("value", "Value")]


def p45_table(f: P45Figures) -> Table:
    rows = [
        {"item": "Employee", "value": f.name},
        {"item": "National Insurance number", "value": f.nino or "not known"},
        {"item": "Payroll ID", "value": f.payroll_id},
        {"item": "Leaving date", "value": f"{f.leaving_date:%d/%m/%Y}"},
        {"item": "Tax code at leaving", "value": f.tax_code + (" (week 1/month 1)" if f.week1_month1 else "")},
        {"item": "Last entries", "value": f.last_period},
        {"item": "Total pay to date", "value": P.gbp(f.total_pay_to_date_pence)},
        {"item": "Total tax to date", "value": P.gbp(f.total_tax_to_date_pence)},
        {"item": "Pay in this employment", "value": P.gbp(f.this_employment_pay_pence)},
        {"item": "Tax in this employment", "value": P.gbp(f.this_employment_tax_pence)},
        {"item": "Continue student loan deductions", "value": "Yes" if f.student_loan_continue else "No"},
        {"item": "Continue postgraduate loan deductions", "value": "Yes" if f.postgrad_loan_continue else "No"},
    ]
    return Table("p45", f"P45 figures: {f.name}", "Leaving this employment", FORM_COLS, rows, None, [f.note], False)


def p60_table(f: P60Figures) -> Table:
    rows = [
        {"item": "Employee", "value": f.name},
        {"item": "National Insurance number", "value": f.nino or "not known"},
        {"item": "Payroll ID", "value": f.payroll_id},
        {"item": "Tax year", "value": f.tax_year},
        {"item": "Final tax code", "value": f.final_tax_code},
        {"item": "Pay in previous employment(s)", "value": P.gbp(f.previous_pay_pence)},
        {"item": "Tax in previous employment(s)", "value": P.gbp(f.previous_tax_pence)},
        {"item": "Pay in this employment", "value": P.gbp(f.this_employment_pay_pence)},
        {"item": "Tax in this employment", "value": P.gbp(f.this_employment_tax_pence)},
        {"item": "Total pay for year", "value": P.gbp(f.total_pay_pence)},
        {"item": "Total tax for year", "value": P.gbp(f.total_tax_pence)},
    ]
    for letter, v in sorted(f.ni.items()):
        rows += [
            {"item": f"NI letter {letter}: earnings at LEL", "value": P.gbp(v["at_lel_pence"])},
            {"item": f"NI letter {letter}: LEL to PT", "value": P.gbp(v["lel_to_pt_pence"])},
            {"item": f"NI letter {letter}: PT to UEL", "value": P.gbp(v["pt_to_uel_pence"])},
            {"item": f"NI letter {letter}: employee contributions", "value": P.gbp(v["ee_pence"])},
        ]
    rows += [
        {"item": "Student loan deductions", "value": P.gbp(f.student_loan_pence)},
        {"item": "Postgraduate loan deductions", "value": P.gbp(f.postgrad_loan_pence)},
        {"item": "SMP", "value": P.gbp(f.smp_pence)},
    ]
    note = (
        "P60 figures. Give the employee their P60 by 31 May; print it from Basic PAYE Tools or your payroll provider."
    )
    return Table(
        "p60", f"P60 figures: {f.name}", f"Tax year {f.tax_year}", FORM_COLS, rows, None, [note, FILING_NOTE], False
    )


# --------------------------------------------------------------------------- #
# Renderers
# --------------------------------------------------------------------------- #


def _cell(v: Any, kind: Kind) -> str:
    if v is None or v == "":
        return ""
    if kind == "money":
        return f"{int(v) / 100:.2f}"
    return str(v)


def _safe(s: str) -> str:
    return f"'{s}" if s[:1] in ("=", "+", "-", "@", "\t", "\r") and not _is_number(s) else s


def _is_number(s: str) -> bool:
    try:
        float(s)
    except ValueError:
        return False
    return True


def to_csv(t: Table) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow([t.title])
    w.writerow([t.subtitle])
    w.writerow([])
    w.writerow([c.label for c in t.columns])
    for r in t.rows:
        w.writerow([_safe(_cell(r.get(c.key), c.kind)) for c in t.columns])
    if t.totals:
        w.writerow(["Total" if i == 0 else _cell(t.totals.get(c.key), c.kind) for i, c in enumerate(t.columns)])
    w.writerow([])
    for n in t.notes:
        w.writerow([n])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(t: Table) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = t.key[:31]
    ws.append([t.title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([t.subtitle])
    ws.append([])
    ws.append([c.label for c in t.columns])
    for cell in ws[4]:
        cell.font = Font(bold=True)

    def val(v: Any, kind: Kind) -> Any:
        if v is None or v == "":
            return None
        if kind == "money":
            return int(v) / 100
        if kind == "int":
            return int(v)
        return _safe(str(v))

    for r in t.rows:
        ws.append([val(r.get(c.key), c.kind) for c in t.columns])
    if t.totals:
        ws.append(["Total" if i == 0 else val(t.totals.get(c.key), c.kind) for i, c in enumerate(t.columns)])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for i, c in enumerate(t.columns, start=1):
        if c.kind == "money":
            for row in ws.iter_rows(min_row=5, min_col=i, max_col=i):
                for cell in row:
                    cell.number_format = "#,##0.00"
    ws.append([])
    for n in t.notes:
        ws.append([n])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(t: Table, employer: str = "") -> bytes:
    from app.pdf import render_table

    return render_table(
        t.title,
        t.subtitle,
        [c.label for c in t.columns],
        [[_cell(r.get(c.key), c.kind) for c in t.columns] for r in t.rows],
        [("Total" if i == 0 else _cell(t.totals.get(c.key), c.kind)) for i, c in enumerate(t.columns)]
        if t.totals
        else None,
        [c.kind == "money" or c.kind == "int" for c in t.columns],
        t.notes,
        landscape=t.landscape,
        employer=employer,
        weights=[c.weight or (1.2 if c.kind in ("money", "int") else 1.6) for c in t.columns],
    )


MEDIA = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def render(t: Table, fmt: Format, employer: str = "", stem: str | None = None) -> tuple[bytes, str, str]:
    """(bytes, media type, file name). `stem` names the file (e.g. "payroll_summary-2026-27-monthly-07")."""
    name = "".join(ch if ch.isalnum() or ch in "-_" else "-" for ch in (stem or t.key)).strip("-") or t.key
    if fmt == "csv":
        return to_csv(t), MEDIA["csv"], f"{name}.csv"
    if fmt == "xlsx":
        return to_xlsx(t), MEDIA["xlsx"], f"{name}.xlsx"
    if fmt == "pdf":
        return to_pdf(t, employer), MEDIA["pdf"], f"{name}.pdf"
    raise ValueError(f"Unknown format {fmt}")
