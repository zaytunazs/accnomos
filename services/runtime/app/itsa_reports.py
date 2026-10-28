"""Self-employment reports: the SA103-style summary, a quarterly update as submitted, and HMRC's calculation.

Each report is a pure function returning a `Report` (title, notes, sections of rows); `render(report, fmt)` turns it
into CSV, XLSX or PDF bytes, and the web page draws the same rows from JSON. Nothing here reads or writes storage.
Written to register in the shared report → formats registry (branch claude/exports) when that merges: the
registry needs only `build` (rows) and the renderers below, so this is not a second export framework.

Figures labelled "HMRC" are HMRC's own numbers, copied from its response; Accnomos never presents its own tax figure
as HMRC's (docs/HONESTY.md).
"""

from __future__ import annotations

import csv
import io
from collections.abc import Sequence
from datetime import date
from decimal import Decimal
from typing import Any, Literal

from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel

from app.itsa import FIELDS, SeSummary
from app.pdf import text as pdf_text

Format = Literal["csv", "xlsx", "pdf"]
Cell = str | int | float | None
MEDIA = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


class Section(BaseModel):
    title: str
    columns: list[str]
    money: list[int] = []  # 0-based column indexes that hold pounds
    rows: list[list[Cell]]


class Report(BaseModel):
    kind: str  # "sa103" / "quarterly" / "calculation"
    title: str
    subtitle: str
    source: str  # "Accnomos, from your approved entries" or "HMRC"
    notes: list[str]
    sections: list[Section]
    filename: str  # without extension


def _gbp(pence: int) -> float:
    return float(Decimal(pence) / 100)


def _num(v: Any) -> float | None:
    return float(v) if isinstance(v, int | float) and not isinstance(v, bool) else None


# --- builders (pure) --------------------------------------------------------------------------------------------


def sa103_report(s: SeSummary, business: str | None = None) -> Report:
    """Box by box, total and disallowable side by side, then every ledger line behind each box."""
    boxes = [
        [
            f.box,
            f.label,
            f.key,
            _gbp(f.pence),
            _gbp(f.disallowable_pence) if f.group == "expenses" else None,
            f.disallowable_box,
            len(f.lines),
        ]
        for f in s.figures
    ]
    totals = [
        ["Income", _gbp(s.income_pence)],
        ["Total expenses", _gbp(s.expenses_pence)],
        ["Disallowable expenses", _gbp(s.disallowable_pence)],
        ["Allowable expenses", _gbp(s.allowable_pence)],
        ["Net profit before tax adjustments", _gbp(s.net_profit_pence)],
    ]
    lines = [
        [
            f.box,
            f.key,
            c.date.isoformat(),
            c.kind,
            c.counterparty or "",
            c.account,
            _gbp(c.pence),
            "yes" if c.disallowable else "",
            c.event_id,
        ]
        for f in s.figures
        for c in f.lines
    ]
    return Report(
        kind="sa103",
        title="Self-employment summary (SA103-style)",
        subtitle=f"{business + ' · ' if business else ''}Tax year {s.tax_year}, "
        f"{_d(s.period_start)} to {_d(s.period_end)}",
        source="Accnomos, worked out by code from your approved entries",
        notes=[s.basis, *s.warnings, *s.blockers],
        sections=[
            Section(
                title="Boxes",
                columns=[
                    "SA103F box",
                    "Description",
                    "HMRC field",
                    "Total £",
                    "Disallowable £",
                    "Disallowable box",
                    "Entries",
                ],
                money=[3, 4],
                rows=boxes,
            ),
            Section(title="Totals", columns=["", "£"], money=[1], rows=totals),
            Section(
                title="Ledger lines behind each box",
                columns=[
                    "Box",
                    "HMRC field",
                    "Date",
                    "Kind",
                    "Counterparty",
                    "Account",
                    "£",
                    "Disallowable",
                    "Entry ID",
                ],
                money=[6],
                rows=lines,
            ),
        ],
        filename=f"self-employment-{s.tax_year}-to-{s.period_end.isoformat()}",
    )


def quarterly_report(receipt: dict) -> Report:
    """A cumulative update exactly as it was sent, with HMRC's receipt details (a stored `_System/itsa` receipt)."""
    sent = receipt.get("sent") or {}
    dates = sent.get("periodDates") or {}
    labels = {f.key: (f.box, f.label) for f in FIELDS}
    rows: list[list[Cell]] = []
    for group, suffix in (("periodIncome", ""), ("periodExpenses", ""), ("periodDisallowableExpenses", "Disallowable")):
        for key, value in (sent.get(group) or {}).items():
            base = key.removesuffix(suffix) if suffix else key
            box, label = labels.get(base, (None, key))
            field = next((f for f in FIELDS if f.key == base), None)
            rows.append([group, key, field.disallowable_box if suffix and field else box, label, _num(value)])
    hmrc = receipt.get("hmrc") or {}
    rec = [
        ["Submitted at", receipt.get("submitted_at")],
        ["Confirmed by", receipt.get("confirmed_by")],
        ["HMRC correlation ID", hmrc.get("correlation_id")],
        ["Receipt ID", hmrc.get("receipt_id")],
        ["Receipt timestamp", hmrc.get("receipt_timestamp")],
        ["Business ID", receipt.get("business_id")],
        ["Matches the books when saved", "yes" if receipt.get("matches_books") else "no"],
        ["Fraud-prevention headers missing", ", ".join(hmrc.get("fraud_headers_missing") or []) or "none"],
        ["Environment", receipt.get("environment") or "sandbox"],
    ]
    end = dates.get("periodEndDate") or "?"
    return Report(
        kind="quarterly",
        title="Quarterly update as submitted",
        subtitle=f"Tax year {receipt.get('tax_year')}, {dates.get('periodStartDate')} to {end} (cumulative)",
        source="The figures Accnomos sent to HMRC, with HMRC's receipt",
        notes=["Each update covers the tax year so far and replaces the one before it."],
        sections=[
            Section(
                title="Figures sent",
                columns=["Group", "HMRC field", "SA103F box", "Description", "£"],
                money=[4],
                rows=rows,
            ),
            Section(title="Receipt", columns=["", ""], rows=[[k, v] for k, v in rec]),
        ],
        filename=f"quarterly-update-{receipt.get('tax_year')}-to-{end}",
    )


def calculation_summary(calc: dict) -> dict[str, Any]:
    """The headline numbers from HMRC's Retrieve a Self Assessment Tax Calculation (Individual Calculations 8.0),
    copied as HMRC returned them; a figure HMRC left out stays None. Never computed here."""
    calculation = calc.get("calculation") or {}
    tc = calculation.get("taxCalculation") or {}
    it = tc.get("incomeTax") or {}
    nics = tc.get("nics") or {}
    eoy = calculation.get("endOfYearEstimate") or {}
    messages = calc.get("messages") or {}
    return {
        "metadata": calc.get("metadata") or {},
        "figures": [
            ("Total income received", _num(it.get("totalIncomeReceivedFromAllSources"))),
            ("Allowances and deductions", _num(it.get("totalAllowancesAndDeductions"))),
            ("Taxable income", _num(it.get("totalTaxableIncome"))),
            ("Income Tax charged", _num(it.get("incomeTaxCharged"))),
            ("Class 2 NIC", _num((nics.get("class2Nics") or {}).get("amount"))),
            ("Class 4 NIC", _num((nics.get("class4Nics") or {}).get("totalAmount"))),
            ("Total NIC", _num(nics.get("totalNic"))),
            ("Total Income Tax and NIC due", _num(tc.get("totalIncomeTaxAndNicsDue"))),
        ],
        "end_of_year_estimate": [
            ("Estimated income for the year", _num(eoy.get("totalEstimatedIncome"))),
            ("Estimated taxable income", _num(eoy.get("totalTaxableIncome"))),
            ("Estimated Income Tax", _num(eoy.get("incomeTaxAmount"))),
            ("Estimated NIC", _num(eoy.get("totalNicAmount"))),
            ("Estimated Income Tax and NIC", _num(eoy.get("incomeTaxNicAmount"))),
        ]
        if eoy
        else [],
        "errors": [m.get("text") for m in messages.get("errors") or [] if isinstance(m, dict)],
        "warnings": [m.get("text") for m in messages.get("warnings") or [] if isinstance(m, dict)],
        "info": [m.get("text") for m in messages.get("info") or [] if isinstance(m, dict)],
    }


def calculation_report(calc: dict) -> Report:
    s = calculation_summary(calc)
    md = s["metadata"]
    notes = ["These are HMRC's figures, as HMRC returned them. Accnomos has not worked out any tax."]
    notes += [f"HMRC error: {t}" for t in s["errors"]] + [f"HMRC warning: {t}" for t in s["warnings"]]
    sections = [
        Section(title="HMRC's calculation", columns=["", "£ (HMRC)"], money=[1], rows=[list(r) for r in s["figures"]]),
    ]
    if s["end_of_year_estimate"]:
        sections.append(
            Section(
                title="HMRC's end-of-year estimate",
                columns=["", "£ (HMRC)"],
                money=[1],
                rows=[list(r) for r in s["end_of_year_estimate"]],
            )
        )
    sections.append(
        Section(
            title="About this calculation",
            columns=["", ""],
            rows=[
                [k, md.get(k)]
                for k in (
                    "calculationId",
                    "taxYear",
                    "calculationType",
                    "calculationTimestamp",
                    "periodFrom",
                    "periodTo",
                    "calculationReason",
                )
            ],
        )
    )
    return Report(
        kind="calculation",
        title="HMRC tax calculation",
        subtitle=f"Tax year {md.get('taxYear')}, {md.get('calculationType')}",
        source="HMRC (Individual Calculations API)",
        notes=notes,
        sections=sections,
        filename=f"hmrc-calculation-{md.get('taxYear')}-{md.get('calculationId')}",
    )


def _d(day: date) -> str:
    return f"{day.day} {day.strftime('%b %Y')}"


# --- renderers --------------------------------------------------------------------------------------------------


def render(report: Report, fmt: Format) -> bytes:
    return {"csv": to_csv, "xlsx": to_xlsx, "pdf": to_pdf}[fmt](report)


def _safe(v: Cell) -> Cell:
    """Spreadsheets run text that starts with = + - @ as a formula: prefix it (OWASP CSV injection)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def to_csv(report: Report) -> bytes:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow([report.title])
    w.writerow([report.subtitle])
    w.writerow([f"Source: {report.source}"])
    for n in report.notes:
        w.writerow([_safe(n)])
    for sec in report.sections:
        w.writerow([])
        w.writerow([sec.title])
        w.writerow(sec.columns)
        for row in sec.rows:
            w.writerow(["" if v is None else _safe(v) for v in row])
    return out.getvalue().encode("utf-8-sig")  # BOM: Excel opens £ correctly


def to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "About"
    for line in (report.title, report.subtitle, f"Source: {report.source}", *report.notes):
        ws.append([_safe(line)])
    ws["A1"].font = Font(bold=True, size=14)
    for sec in report.sections:
        sh = wb.create_sheet(sec.title[:31].replace("/", "-"))
        sh.append(sec.columns)
        for c in sh[1]:
            c.font = Font(bold=True)
        sh.freeze_panes = "A2"
        for row in sec.rows:
            sh.append([_safe(v) for v in row])
        for i in sec.money:
            for (cell,) in sh.iter_rows(min_row=2, min_col=i + 1, max_col=i + 1):
                cell.number_format = '"£"#,##0.00;-"£"#,##0.00'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v: Cell, money: bool) -> str:
    if v is None:
        return ""
    if money and isinstance(v, int | float):
        return f"-£{abs(v):,.2f}" if v < 0 else f"£{v:,.2f}"
    return pdf_text(str(v))


def to_pdf(report: Report) -> bytes:
    pdf = FPDF(orientation="L", format="A4")
    pdf.set_auto_page_break(True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 8, pdf_text(report.title), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, pdf_text(report.subtitle), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.cell(0, 6, pdf_text(f"Source: {report.source}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 8)
    for n in report.notes:
        pdf.multi_cell(0, 4, pdf_text(n), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    for sec in report.sections:
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 7, pdf_text(sec.title), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font("Helvetica", "", 7)
        widths = _widths(sec, pdf.epw)
        with pdf.table(col_widths=widths, text_align="LEFT", line_height=4, first_row_as_headings=True) as t:
            t.row([pdf_text(c) for c in sec.columns])
            for row in sec.rows:
                t.row([_cell(v, i in sec.money) for i, v in enumerate(row)])
    return bytes(pdf.output())


def _widths(sec: Section, total: float) -> Sequence[float]:
    lens = [
        max([len(sec.columns[i]), *(len(str(r[i] or "")) for r in sec.rows[:200])]) for i in range(len(sec.columns))
    ]
    lens = [min(max(n, 4), 40) for n in lens]
    s = sum(lens)
    return [total * n / s for n in lens]
