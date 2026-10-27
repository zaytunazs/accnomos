"""Reports: profit and loss, aged debtors and creditors, VAT summary and a 12-month trend, by code from the ledger.

Every figure is worked out from the approved ledger events on each request; no model is involved and nothing is
stored. Reports never post anything: the only file they write is the spreadsheet `POST /reports/export` saves in
07 Reports (app/routes_reports.py).

Profit and loss is **as approved**: every approved event counts on its own date (an invoice when it was raised, a
receipt when it was bought), and an undo counts against the period of the reversal. There is no cash-basis view.
Accounts are placed by their four-digit code, the chart in app/journal.py:

    4000-4999  Sales                      credit − debit
    5000-5999  Cost of sales              debit − credit   (5000 Materials, 5100 Tools & small equipment)
    6000-8999  Overheads                  debit − credit   (6200 Advertising, 7000 Gross wages … 7999 Sundry)

Anything else (1xxx assets, 2xxx liabilities such as 2200 VAT control, 3xxx capital) is balance sheet and left out,
so the VAT on a sale or purchase never reaches the P&L; VAT that cannot be reclaimed (codes N and O) is already in
the expense line, as app/journal.py posts it.

Aged balances are app.contacts.aged, unchanged. The VAT summary is accnomos_ledger.compute_vat_return per period.
The trend's "purchases" are cost of sales plus overheads, so a month's net is the P&L's net profit for that month.
"""

from __future__ import annotations

import io
import re
from collections import defaultdict
from collections.abc import Iterable, Sequence
from datetime import UTC, date, datetime, timedelta
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel

from accnomos_ledger import LedgerEvent, compute_vat_return
from app.contacts import BUCKETS, AgedBalances, AgedTable
from app.deadlines import add_months, period_id
from app.deadlines import period_label as vat_period_label
from app.facts import period_label
from app.projection import SAMPLE_LABEL, WHOLE_GBP, _table, pounds
from app.schemas import AskPeriod

Section = Literal["sales", "cost_of_sales", "overheads"]
BASIS = (
    "As approved: each entry counts on its own date (invoices when raised, receipts when bought); "
    "an undo counts on the date it was made. Not a cash-basis view."
)
#: Whoever approves the demo books' entries (tools/seed_demo.py USER). Books with any entry approved by it are the
#: sample, and an exported report says so.
SAMPLE_APPROVER = "demo-seed"
REPORT_NAMES = {
    "pl": "profit-and-loss",
    "aged": "aged-balances",
    "vat": "vat-summary",
    "trend": "monthly-trend",
    "all": "reports",
}
ReportKind = Literal["pl", "aged", "vat", "trend", "all"]
_CODE = re.compile(r"^\s*(\d{4})\b")


def account_code(account: str) -> int | None:
    """'7500 Office supplies' -> 7500; None when the name does not start with a four-digit code."""
    m = _CODE.match(account)
    return int(m.group(1)) if m else None


def section_of(account: str) -> Section | None:
    """Where an account sits on the P&L, or None for a balance-sheet account."""
    code = account_code(account)
    if code is None:
        return None
    if 4000 <= code <= 4999:
        return "sales"
    if 5000 <= code <= 5999:
        return "cost_of_sales"
    if 6000 <= code <= 8999:
        return "overheads"
    return None


def is_sample(events: Iterable[LedgerEvent]) -> bool:
    """True when these are the demo books (tools/seed_demo.py approves every entry as SAMPLE_APPROVER)."""
    return any(e.approved_by == SAMPLE_APPROVER for e in events)


# --------------------------------------------------------------------------- #
# Profit and loss
# --------------------------------------------------------------------------- #


class PlRow(BaseModel):
    account: str
    amount_pence: int  # sales positive as income, costs positive as spending


class PlSection(BaseModel):
    rows: list[PlRow]
    total_pence: int


class ProfitAndLoss(BaseModel):
    start: date
    end: date
    label: str  # "Jul–Sep 2026"
    basis: str = BASIS
    sales: PlSection
    cost_of_sales: PlSection
    gross_profit_pence: int  # sales − cost of sales
    overheads: PlSection
    net_profit_pence: int  # gross profit − overheads; negative is a loss
    entries: int  # approved entries dated in the period that touch a P&L account
    sample: bool = False


def _amounts(events: Iterable[LedgerEvent]) -> dict[str, int]:
    """Account -> signed P&L amount in pence (income credit − debit, costs debit − credit)."""
    out: dict[str, int] = defaultdict(int)
    for e in events:
        for ln in e.lines:
            sec = section_of(ln.account)
            if sec == "sales":
                out[ln.account] += ln.credit - ln.debit
            elif sec is not None:
                out[ln.account] += ln.debit - ln.credit
    return out


def _section(amounts: dict[str, int], which: Section) -> PlSection:
    rows = [
        PlRow(account=a, amount_pence=v)
        for a, v in sorted(amounts.items(), key=lambda kv: (account_code(kv[0]) or 0, kv[0]))
        if v and section_of(a) == which
    ]
    return PlSection(rows=rows, total_pence=sum(r.amount_pence for r in rows))


def profit_and_loss(events: Sequence[LedgerEvent], start: date, end: date) -> ProfitAndLoss:
    """The P&L for [start, end] from the events dated in it. Accounts that net to nothing are left out."""
    inside = [e for e in events if start <= e.date <= end]
    amounts = _amounts(inside)
    sales, cos, over = (_section(amounts, s) for s in ("sales", "cost_of_sales", "overheads"))
    gross = sales.total_pence - cos.total_pence
    return ProfitAndLoss(
        start=start,
        end=end,
        label=period_label(AskPeriod(start=start, end=end)),
        sales=sales,
        cost_of_sales=cos,
        gross_profit_pence=gross,
        overheads=over,
        net_profit_pence=gross - over.total_pence,
        entries=sum(any(section_of(ln.account) for ln in e.lines) for e in inside),
        sample=is_sample(events),
    )


# --------------------------------------------------------------------------- #
# VAT summary
# --------------------------------------------------------------------------- #


class VatSummaryPeriod(BaseModel):
    period: str  # "2026-Q3" or "2026-09", as GET /books/vat/periods names it
    start: date
    end: date
    label: str
    vat_on_sales_pence: int  # box 1
    vat_reclaimed_pence: int  # box 4
    net_vat_pence: int  # box 3 − box 4: positive to pay, negative to reclaim (box 5 is its absolute value)
    sales_ex_vat_pounds: int  # box 6, whole pounds as HMRC takes it
    purchases_ex_vat_pounds: int  # box 7
    entries: int  # entries counted in the boxes
    warnings: list[str] = []


class VatSummary(BaseModel):
    periods: list[VatSummaryPeriod]  # oldest first
    total_net_vat_pence: int
    sample: bool = False


def vat_summary(events: Sequence[LedgerEvent], spans: Sequence[tuple[date, date]], months: int = 3) -> VatSummary:
    """compute_vat_return for each (start, end) in `spans`, oldest first."""
    rows: list[VatSummaryPeriod] = []
    for start, end in sorted(spans):
        ret = compute_vat_return(list(events), period_id(start, months), start, end)
        rows.append(
            VatSummaryPeriod(
                period=period_id(start, months),
                start=start,
                end=end,
                label=vat_period_label(start, end),
                vat_on_sales_pence=int(ret.vat_due_sales * 100),
                vat_reclaimed_pence=int(ret.vat_reclaimed * 100),
                net_vat_pence=int((ret.total_vat_due - ret.vat_reclaimed) * 100),
                sales_ex_vat_pounds=ret.total_sales_ex_vat,
                purchases_ex_vat_pounds=ret.total_purchases_ex_vat,
                entries=len(ret.event_ids),
                warnings=ret.warnings,
            )
        )
    return VatSummary(periods=rows, total_net_vat_pence=sum(r.net_vat_pence for r in rows), sample=is_sample(events))


# --------------------------------------------------------------------------- #
# Monthly trend
# --------------------------------------------------------------------------- #


class TrendMonth(BaseModel):
    month: str  # "2026-10"
    label: str  # "Oct 2026"
    start: date
    end: date
    sales_pence: int
    purchases_pence: int  # cost of sales + overheads
    net_pence: int  # sales − purchases: the month's net profit


class Trend(BaseModel):
    months: list[TrendMonth]  # oldest first
    sample: bool = False


def trend_window(last_month: date, months: int = 12) -> tuple[date, date]:
    """(first day of the oldest month, last day of `last_month`'s month)."""
    first = add_months(date(last_month.year, last_month.month, 1), 1 - months)
    return first, add_months(date(last_month.year, last_month.month, 1), 1) - timedelta(days=1)


def monthly_trend(events: Sequence[LedgerEvent], last_month: date, months: int = 12) -> Trend:
    """`months` calendar months ending with the month of `last_month`, from the same accounts as the P&L."""
    first, _ = trend_window(last_month, months)
    out: list[TrendMonth] = []
    for i in range(months):
        start = add_months(first, i)
        end = add_months(start, 1) - timedelta(days=1)
        pl = profit_and_loss(events, start, end)
        costs = pl.cost_of_sales.total_pence + pl.overheads.total_pence
        out.append(
            TrendMonth(
                month=f"{start:%Y-%m}",
                label=f"{start:%b %Y}",
                start=start,
                end=end,
                sales_pence=pl.sales.total_pence,
                purchases_pence=costs,
                net_pence=pl.sales.total_pence - costs,
            )
        )
    return Trend(months=out, sample=is_sample(events))


# --------------------------------------------------------------------------- #
# Spreadsheets (07 Reports and the P&L sheet of Books.xlsx)
# --------------------------------------------------------------------------- #


def pl_sheet(ws: Worksheet, pl: ProfitAndLoss) -> None:
    """Account rows under Sales, Cost of sales and Overheads, with gross and net profit."""
    rows: list[tuple[object, ...]] = []
    bold: list[int] = []  # 0-based indexes into rows

    def add(label: str, pence: int | None, strong: bool = False) -> None:
        if strong:
            bold.append(len(rows))
        rows.append((label, pounds(pence)))

    for title, sec in (("Sales", pl.sales), ("Cost of sales", pl.cost_of_sales)):
        add(title, None, True)
        for r in sec.rows:
            add(f"  {r.account}", r.amount_pence)
        add(f"Total {title.lower()}", sec.total_pence, True)
    add("Gross profit", pl.gross_profit_pence, True)
    add("Overheads", None, True)
    for r in pl.overheads.rows:
        add(f"  {r.account}", r.amount_pence)
    add("Total overheads", pl.overheads.total_pence, True)
    add("Net profit" if pl.net_profit_pence >= 0 else "Net loss", pl.net_profit_pence, True)
    header = (f"Profit and loss · {pl.label} ({pl.start:%d/%m/%Y}–{pl.end:%d/%m/%Y})", "£")
    _table(ws, header, rows, money=(2,), widths=(48, 16))
    for i in bold:
        ws.cell(row=i + 2, column=1).font = Font(bold=True)
        ws.cell(row=i + 2, column=2).font = Font(bold=True)
    ws.append([])
    ws.append([pl.basis])
    entries = f"{pl.entries} approved {'entry' if pl.entries == 1 else 'entries'}"
    ws.append([f"{entries}. Worked out by code from _System/Ledger/events; no model involved."])


def _aged_sheet(ws: Worksheet, who: str, table: AgedTable, as_of: date) -> None:
    rows: list[tuple[object, ...]] = [
        (r.name, *(pounds(r.buckets[b]) for b in BUCKETS), pounds(r.total_pence)) for r in table.rows
    ]
    rows.append(("Total", *(pounds(table.buckets[b]) for b in BUCKETS), pounds(table.total_pence)))
    header = (f"{who} at {as_of:%d/%m/%Y}", *(f"{b} days £" for b in BUCKETS), "Total £")
    _table(ws, header, rows, money=range(2, 7), widths=(34, 14, 14, 14, 14, 14))
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append([])
    ws.append(["Unpaid documents by age from their date: no bank payment matched to them yet."])


def _vat_sheet(ws: Worksheet, summary: VatSummary) -> None:
    rows = [
        (
            r.label,
            r.start,
            r.end,
            pounds(r.vat_on_sales_pence),
            pounds(r.vat_reclaimed_pence),
            pounds(r.net_vat_pence),
            r.sales_ex_vat_pounds,
            r.purchases_ex_vat_pounds,
            r.entries,
        )
        for r in summary.periods
    ]
    header = (
        "Period",
        "From",
        "To",
        "VAT on sales (box 1) £",
        "VAT reclaimed (box 4) £",
        "Net VAT, + to pay (3 − 4) £",
        "Sales ex VAT (box 6) £",
        "Purchases ex VAT (box 7) £",
        "Entries",
    )
    _table(ws, header, rows, money=(4, 5, 6), dates=(2, 3), widths=(18, 12, 12, 20, 22, 24, 20, 24, 9))
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=8):
        for cell in row:
            cell.number_format = WHOLE_GBP


def _trend_sheet(ws: Worksheet, trend: Trend) -> None:
    rows = [(m.label, pounds(m.sales_pence), pounds(m.purchases_pence), pounds(m.net_pence)) for m in trend.months]
    rows.append(
        (
            "Total",
            *(
                pounds(sum(getattr(m, f) for m in trend.months))
                for f in ("sales_pence", "purchases_pence", "net_pence")
            ),
        )
    )
    header = ("Month", "Sales £", "Purchases (cost of sales + overheads) £", "Net £")
    _table(ws, header, rows, money=(2, 3, 4), widths=(12, 14, 38, 14))
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)


def build_report_xlsx(
    kind: ReportKind,
    *,
    pl: ProfitAndLoss | None = None,
    aged: AgedBalances | None = None,
    vat: VatSummary | None = None,
    trend: Trend | None = None,
    sample: bool = False,
    generated: datetime | None = None,
) -> tuple[bytes, list[str]]:
    """(.xlsx bytes, sheet names): an About sheet, then the sheets for `kind` ("all" = every report passed in)."""
    generated = generated or datetime.now(UTC)
    wb = Workbook()
    about = wb.active
    if not isinstance(about, Worksheet):  # a new Workbook always starts with one worksheet
        raise RuntimeError("openpyxl Workbook() has no first worksheet")
    about.title = "About"
    lines = [
        f"Accnomos report: {REPORT_NAMES[kind].replace('-', ' ')}",
        f"Generated at {generated:%d/%m/%Y %H:%M} UTC from the ledger event files in _System/Ledger/events.",
        "Worked out by code; no model involved. Saving a report posts nothing to your books.",
        BASIS,
    ]
    if sample:
        lines.insert(0, SAMPLE_LABEL)
    for text in lines:
        about.append([text])
    about.cell(row=1, column=1).font = Font(bold=True)
    about.column_dimensions["A"].width = 110

    if kind in ("pl", "all") and pl:
        pl_sheet(wb.create_sheet("P&L"), pl)
    if kind in ("aged", "all") and aged:
        _aged_sheet(wb.create_sheet("Debtors"), "Customers who owe you", aged.debtors, aged.as_of)
        _aged_sheet(wb.create_sheet("Creditors"), "Suppliers you owe", aged.creditors, aged.as_of)
    if kind in ("vat", "all") and vat:
        _vat_sheet(wb.create_sheet("VAT summary"), vat)
    if kind in ("trend", "all") and trend:
        _trend_sheet(wb.create_sheet("Trend"), trend)
    if sample:
        for ws in wb.worksheets[1:]:
            ws.append([])
            ws.append([SAMPLE_LABEL])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), wb.sheetnames


def report_filename(kind: ReportKind, today: date) -> str:
    """'2026-10-27-profit-and-loss.xlsx'."""
    return f"{today.isoformat()}-{REPORT_NAMES[kind]}.xlsx"
