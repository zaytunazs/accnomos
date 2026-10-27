"""Books.xlsx: a read-only spreadsheet view of the books, rebuilt from the files after every decision.

Decision H25: personal OneDrive has no Graph workbook API, so we build the whole .xlsx here with
openpyxl and upload it as a file. Nothing reads it back; the JSON files stay the source of truth.
"""

from __future__ import annotations

import io
from collections import defaultdict
from collections.abc import Iterable, Sequence
from datetime import UTC, date, datetime

from accnomos_ledger import compute_vat_return
from accnomos_ledger.bank import BankLine
from accnomos_ledger.events import LedgerEvent
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas import Proposal

GBP = '"£"#,##0.00;-"£"#,##0.00'
WHOLE_GBP = '"£"#,##0;-"£"#,##0'
DATE = "dd/mm/yyyy"
SAMPLE_LABEL = "SAMPLE — FICTIONAL COMPANY"
SHEETS = ("Read me", "Journal", "Accounts", "VAT", "Bank", "Proposals")
VAT_BOXES = (
    (1, "VAT due on sales"), (2, "VAT due on acquisitions from NI/EU"), (3, "Total VAT due"),
    (4, "VAT reclaimed on purchases"), (5, "Net VAT to pay or reclaim"), (6, "Total sales ex VAT"),
    (7, "Total purchases ex VAT"), (8, "Goods supplied to NI/EU"), (9, "Goods acquired from NI/EU"),
)


def pounds(pence: int | None) -> float | None:
    return None if pence is None else pence / 100


def _table(ws: Worksheet, header: Sequence[str], rows: Iterable[Sequence[object]],
           money: Sequence[int] = (), dates: Sequence[int] = (), widths: Sequence[int] = ()) -> None:
    """Bold header, frozen below it, money/date number formats on the given (1-based) columns."""
    ws.append(list(header))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(list(row))
    for col in money:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = GBP
    for col in dates:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = DATE
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def _read_me(ws: Worksheet, generated: datetime) -> None:
    rows = [
        ("Accnomos books", ""),
        (SAMPLE_LABEL, "Demo books are fictional; real books show your own entries only."),
        ("Generated at", generated.strftime("%d/%m/%Y %H:%M UTC")),
        ("", ""),
        ("Sheet", "What it shows"),
        ("Journal", "Every line of every approved ledger entry (double entry, £)."),
        ("Accounts", "Trial balance: total debits, credits and balance per account."),
        ("VAT", "The 9 VAT return boxes for the last 4 complete quarters, computed by code."),
        ("Bank", "Imported bank statement lines and the ledger entries they are matched to."),
        ("Proposals", "Documents read by the AI that are still waiting for your decision."),
        ("", ""),
        ("This file is rebuilt automatically. Edits here are overwritten; change things in the app.", ""),
    ]
    for row in rows:
        ws.append(list(row))
    for r in (1, 2, 5):
        ws.cell(row=r, column=1).font = Font(bold=True)
    ws.cell(row=5, column=2).font = Font(bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 80


def _journal(ws: Worksheet, events: Sequence[LedgerEvent]) -> None:
    rows = (
        (e.date, e.id, ln.account, pounds(ln.debit) or None, pounds(ln.credit) or None,
         ln.vat_code.value if ln.vat_code else None, e.counterparty, ", ".join(e.evidence) or None)
        for e in events for ln in e.lines
    )
    header = ("Date", "Event id", "Account", "Debit £", "Credit £", "VAT code", "Counterparty", "Evidence doc")
    _table(ws, header, rows, money=(4, 5), dates=(1,), widths=(12, 34, 32, 12, 12, 9, 28, 66))


def trial_balance(events: Iterable[LedgerEvent]) -> list[tuple[str, int, int]]:
    """(account, debits, credits) in pence, sorted by account code."""
    totals: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    for e in events:
        for ln in e.lines:
            totals[ln.account][0] += ln.debit
            totals[ln.account][1] += ln.credit
    return [(acct, dr, cr) for acct, (dr, cr) in sorted(totals.items())]


def _accounts(ws: Worksheet, events: Sequence[LedgerEvent]) -> None:
    tb = trial_balance(events)
    rows = [(a, pounds(dr), pounds(cr), pounds(dr - cr)) for a, dr, cr in tb]
    rows.append(("Total", pounds(sum(t[1] for t in tb)), pounds(sum(t[2] for t in tb)), pounds(sum(t[1] - t[2] for t in tb))))
    _table(ws, ("Account", "Debits £", "Credits £", "Balance £"), rows, money=(2, 3, 4), widths=(34, 14, 14, 14))
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)


def last_quarters(today: date, n: int = 4) -> list[tuple[date, date, str]]:
    """The last `n` complete calendar quarters, oldest first, as (start, end, period key)."""
    out: list[tuple[date, date, str]] = []
    year, q = today.year, (today.month - 1) // 3  # q = index of the current quarter (0-3)
    for _ in range(n):
        year, q = (year - 1, 3) if q == 0 else (year, q - 1)
        start = date(year, q * 3 + 1, 1)
        next_start = date(year, q * 3 + 4, 1) if q < 3 else date(year + 1, 1, 1)
        out.append((start, next_start - date.resolution, f"{year % 100:02d}A{q + 1}"))
    return out[::-1]


def _vat(ws: Worksheet, events: Sequence[LedgerEvent], today: date) -> None:
    returns = [compute_vat_return(list(events), key, s, e) for s, e, key in last_quarters(today)]
    header = ["Box", "What"] + [f"{r.period_key} ({r.start:%d/%m/%Y}–{r.end:%d/%m/%Y})" for r in returns]
    values = {
        1: [float(r.vat_due_sales) for r in returns], 2: [float(r.vat_due_acquisitions) for r in returns],
        3: [float(r.total_vat_due) for r in returns], 4: [float(r.vat_reclaimed) for r in returns],
        5: [float(r.net_vat_due) for r in returns], 6: [r.total_sales_ex_vat for r in returns],
        7: [r.total_purchases_ex_vat for r in returns], 8: [r.goods_supplied_ni_eu for r in returns],
        9: [r.acquisitions_ni_eu for r in returns],
    }
    rows = [(box, label, *values[box]) for box, label in VAT_BOXES]
    _table(ws, header, rows, widths=(6, 36, *([26] * len(returns))))
    for row in ws.iter_rows(min_row=2, min_col=3):
        for cell in row:
            cell.number_format = GBP if row[0].row <= 6 else WHOLE_GBP


def _bank(ws: Worksheet, lines: Sequence[BankLine]) -> None:
    rows = (
        (ln.date, ln.description, ln.reference, pounds(ln.amount_pence), pounds(ln.balance_pence), ln.status,
         ", ".join(ln.event_ids()) or None, ln.source_file)
        for ln in lines
    )
    header = ("Date", "Description", "Reference", "Amount £", "Balance £", "Status", "Matched event", "Statement")
    _table(ws, header, rows, money=(4, 5), dates=(1,), widths=(12, 40, 18, 12, 12, 11, 34, 50))


def _proposals(ws: Worksheet, proposals: Sequence[Proposal]) -> None:
    rows = (
        (p.created_at.replace(tzinfo=None) if p.created_at else None, p.id, p.status, p.extracted.supplier_name,
         p.extracted.date, p.extracted.gross_total, p.extracted.category_guess, "; ".join(p.reasons) or None, p.file_path)
        for p in proposals
    )
    header = ("Received", "Proposal id", "Status", "Supplier", "Document date", "Gross £", "Category", "Reasons", "File")
    _table(ws, header, rows, money=(6,), dates=(1,), widths=(12, 34, 13, 28, 13, 12, 24, 50, 60))


def build_books_xlsx(
    events: Sequence[LedgerEvent], bank_lines: Sequence[BankLine], proposals: Sequence[Proposal],
    generated: datetime | None = None,
) -> bytes:
    """The whole workbook as .xlsx bytes."""
    generated = generated or datetime.now(UTC)
    wb = Workbook()
    sheets = {name: (wb.active if i == 0 else wb.create_sheet()) for i, name in enumerate(SHEETS)}
    for name, ws in sheets.items():
        ws.title = name
    _read_me(sheets["Read me"], generated)
    _journal(sheets["Journal"], events)
    _accounts(sheets["Accounts"], events)
    _vat(sheets["VAT"], events, generated.date())
    _bank(sheets["Bank"], bank_lines)
    _proposals(sheets["Proposals"], proposals)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
