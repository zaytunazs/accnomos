"""Self-employment (MTD Income Tax) figures from the ledger, receipts and reports.

The request body shape is HMRC's: Self-Employment Business API 5.0, "Create or Amend a Self-Employment Cumulative
Period Summary" (developer.service.hmrc.gov.uk/api-documentation/docs/api/service/self-employment-business-api/5.0,
schemas/createAmendCumulativePeriodSummary/request.json, checked 28 Oct 2026). Expenses are totals and disallowables
a subset: HMRC's SA103F mapping (income-tax-mtd-changelog mapping/csv/sa103f_mapping_v3.csv).
"""

from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from accnomos_connectors import FakeStorageProvider
from accnomos_ledger import LedgerEvent, Line, VatCode
from app import itsa
from app import itsa_reports as rep
from app.books import Books
from app.main import app
from app.paths import ITSA
from app.routes_books import get_books

H = {"X-Graph-Token": "t", "X-User": "sam@sambuilds.example"}
AT = datetime(2026, 10, 1, 12, 0)

# Every key of HMRC's request schema (createAmendCumulativePeriodSummary/request.json), minus consolidatedExpenses
# and taxTakenOffTradingIncome, which Accnomos never sends.
HMRC_INCOME = {"turnover", "other"}
HMRC_EXPENSES = {
    "costOfGoods",
    "paymentsToSubcontractors",
    "wagesAndStaffCosts",
    "carVanTravelExpenses",
    "premisesRunningCosts",
    "maintenanceCosts",
    "adminCosts",
    "businessEntertainmentCosts",
    "advertisingCosts",
    "interestOnBankOtherLoans",
    "financeCharges",
    "irrecoverableDebts",
    "professionalFees",
    "depreciation",
    "otherExpenses",
}


def ev(id_: str, day: date, kind: str, *lines: Line, counterparty: str | None = None) -> LedgerEvent:
    return LedgerEvent(
        id=id_ * 32 if len(id_) == 1 else id_,
        kind=kind,
        date=day,
        lines=list(lines),
        counterparty=counterparty,  # type: ignore[arg-type]
        approved_by="sam",
        approved_at=AT,
    )


def sale(id_: str, day: date, net: int, vat: int = 0, account: str = "4000 Sales") -> LedgerEvent:
    ls = [Line(account="1100 Debtors", debit=net + vat), Line(account=account, credit=net, vat_code=VatCode.STANDARD)]
    if vat:
        ls.append(Line(account="2200 VAT control", credit=vat))
    return ev(id_, day, "sale", *ls, counterparty="Customer")


def cost(id_: str, day: date, account: str, net: int, vat: int = 0, kind: str = "purchase") -> LedgerEvent:
    ls = [Line(account=account, debit=net, vat_code=VatCode.STANDARD), Line(account="1200 Bank", credit=net + vat)]
    if vat:
        ls.append(Line(account="2200 VAT control", debit=vat))
    return ev(id_, day, kind, *ls, counterparty="Supplier")


BOOKS = [
    sale("a", date(2026, 4, 10), 1_000_00, 200_00),
    sale("b", date(2026, 9, 1), 500_50),
    ev(
        "c",
        date(2026, 5, 2),
        "sale",
        Line(account="1200 Bank", debit=25_00),
        Line(account="4900 Other income", credit=25_00),
    ),
    cost("d", date(2026, 5, 3), "5000 Materials", 300_00, 60_00),
    cost("e", date(2026, 6, 1), "7100 Rent", 400_00),
    cost("f", date(2026, 6, 2), "7700 Insurance", 50_00),
    cost("g", date(2026, 6, 3), "7400 Travel", 12_34),
    cost("h", date(2026, 6, 4), "7500 Office supplies", 9_99),
    cost("i", date(2026, 6, 5), "7600 Professional fees", 150_00),
    cost("j", date(2026, 6, 6), "7900 Bank charges", 5_00, kind="bank"),
    cost("k", date(2026, 6, 7), "6200 Advertising", 80_00),
    cost("l", date(2026, 6, 8), "7999 Client entertaining", 45_00),
    cost("m", date(2026, 6, 9), "8000 Depreciation", 100_00, kind="adjustment"),
    cost("n", date(2026, 6, 10), "7300 Motor expenses (private use, disallowable)", 20_00),
    cost("o", date(2026, 6, 11), "7999 Sundry expenses", 3_00),
    # outside the window: last tax year, and after the period end
    cost("p", date(2026, 4, 5), "7500 Office supplies", 999_00),
    cost("q", date(2026, 10, 6), "7500 Office supplies", 777_00),
]


def test_tax_year_helpers() -> None:
    assert itsa.tax_year_of(date(2026, 4, 5)) == "2025-26"
    assert itsa.tax_year_of(date(2026, 4, 6)) == "2026-27"
    assert itsa.tax_year_bounds("2026-27") == (date(2026, 4, 6), date(2027, 4, 5))
    assert itsa.tax_year_bounds("2099-00") == (date(2099, 4, 6), date(2100, 4, 5))
    for bad in ("2026-28", "2026/27", "26-27"):
        with pytest.raises(ValueError):
            itsa.tax_year_bounds(bad)


@pytest.mark.parametrize(
    ("account", "field"),
    [
        ("4000 Sales", "turnover"),
        ("4900 Other income", "other"),
        ("5000 Materials", "costOfGoods"),
        ("5100 Tools & small equipment", "costOfGoods"),
        ("6200 Advertising", "advertisingCosts"),
        ("7000 Gross wages", "wagesAndStaffCosts"),
        ("7010 Employer's NI", "wagesAndStaffCosts"),
        ("7100 Rent", "premisesRunningCosts"),
        ("7200 Utilities", "premisesRunningCosts"),
        ("7700 Insurance", "premisesRunningCosts"),
        ("7300 Motor expenses", "carVanTravelExpenses"),
        ("7310 Motor — mileage allowance", "carVanTravelExpenses"),
        ("7410 Subsistence", "carVanTravelExpenses"),
        ("7510 Software & subscriptions", "adminCosts"),
        ("7550 Telephone & internet", "adminCosts"),
        ("7520 Postage", "adminCosts"),
        ("7600 Professional fees", "professionalFees"),
        ("7900 Bank charges", "financeCharges"),
        ("7920 Loan interest", "interestOnBankOtherLoans"),
        ("7950 Bad debts written off", "irrecoverableDebts"),
        ("7999 Sundry expenses", "otherExpenses"),
        ("7999 Client entertaining", "businessEntertainmentCosts"),
        ("7999 Staff entertainment", "wagesAndStaffCosts"),
        ("5200 Subcontractors (CIS)", "paymentsToSubcontractors"),
        ("7999 Van repairs", "maintenanceCosts"),
        ("8000 Depreciation", "depreciation"),
        ("4100 Interest received", "turnover"),  # a name rule never moves income into an expense
        ("1200 Bank", None),
        ("2200 VAT control", None),
        ("2300 Owner / director loan", None),
        ("Misc", None),
    ],
)
def test_field_for(account: str, field: str | None) -> None:
    assert itsa.field_for(account) == field


def test_every_api_field_is_mapped_to_its_sa103f_box() -> None:
    assert {f.key for f in itsa.FIELDS if f.group == "income"} == HMRC_INCOME
    assert {f.key for f in itsa.FIELDS if f.group == "expenses"} == HMRC_EXPENSES
    # sa103f_mapping_v3.csv: expenses 17-30, the disallowable box is 15 higher (32-45); 24 holds two fields
    for f in itsa.FIELDS:
        if f.group == "expenses":
            assert 17 <= f.box <= 30 and f.disallowable_box == f.box + 15


def test_summary_is_cumulative_and_traceable() -> None:
    s = itsa.se_summary(BOOKS, "2026-27", date(2026, 10, 5))
    by = {f.key: f for f in s.figures}
    assert s.period_start == date(2026, 4, 6)
    assert by["turnover"].pence == 1_500_50  # net of VAT; 2200 VAT control never counts
    assert by["other"].pence == 25_00
    assert by["costOfGoods"].pence == 300_00
    assert by["premisesRunningCosts"].pence == 450_00
    assert by["adminCosts"].pence == 9_99  # 5 April (last year) and 6 October (after the end) are left out
    assert by["businessEntertainmentCosts"].disallowable_pence == 45_00
    assert by["depreciation"].pence == by["depreciation"].disallowable_pence == 100_00
    assert by["carVanTravelExpenses"].pence == 32_34 and by["carVanTravelExpenses"].disallowable_pence == 20_00
    assert s.income_pence == 1_525_50
    assert s.expenses_pence == 300_00 + 450_00 + 32_34 + 9_99 + 150_00 + 5_00 + 80_00 + 45_00 + 100_00 + 3_00
    assert s.disallowable_pence == 165_00
    assert s.allowable_pence == s.expenses_pence - 165_00
    assert s.net_profit_pence == s.income_pence - s.allowable_pence
    assert "p" * 32 not in s.event_ids and "q" * 32 not in s.event_ids
    # drill-down: the lines behind a box add up to it
    for f in s.figures:
        assert sum(c.pence for c in f.lines) == f.pence
    assert [c.event_id for c in by["premisesRunningCosts"].lines] == ["e" * 32, "f" * 32]
    assert s.can_submit and s.warnings[0] == itsa.CASH_BASIS_WARNING


def test_hmrc_body_matches_the_documented_request() -> None:
    body = itsa.se_summary(BOOKS, "2026-27", date(2026, 10, 5)).hmrc_body
    assert set(body) == {"periodDates", "periodIncome", "periodExpenses", "periodDisallowableExpenses"}
    assert body["periodDates"] == {"periodStartDate": "2026-04-06", "periodEndDate": "2026-10-05"}
    assert set(body["periodIncome"]) == HMRC_INCOME  # zero or not, HMRC wants income and expenses sent
    assert set(body["periodExpenses"]) == HMRC_EXPENSES
    assert set(body["periodDisallowableExpenses"]) == {f"{k}Disallowable" for k in HMRC_EXPENSES}
    assert body["periodIncome"] == {"turnover": 1500.5, "other": 25.0}
    assert body["periodExpenses"]["adminCosts"] == 9.99
    assert body["periodDisallowableExpenses"]["depreciationDisallowable"] == 100.0
    # multipleOf 0.01, within the schema's range
    for group in ("periodIncome", "periodExpenses", "periodDisallowableExpenses"):
        for v in body[group].values():
            assert round(v * 100) == pytest.approx(v * 100) and abs(v) <= 99999999999.99
    json.dumps(body)  # plain JSON numbers


def test_no_disallowables_no_disallowable_object_and_zeros_sent() -> None:
    body = itsa.se_summary([sale("a", date(2026, 4, 10), 100_00)], "2026-27", date(2026, 7, 5)).hmrc_body
    assert "periodDisallowableExpenses" not in body
    assert body["periodExpenses"]["costOfGoods"] == 0.0 and len(body["periodExpenses"]) == 15


def test_reversal_nets_off_and_negative_turnover_blocks() -> None:
    s = sale("a", date(2026, 5, 1), 100_00)
    undo = ev(
        "b",
        date(2026, 5, 2),
        "reversal",
        Line(account="1100 Debtors", credit=100_00),
        Line(account="4000 Sales", debit=100_00),
    )
    assert itsa.se_summary([s, undo], "2026-27", date(2026, 7, 5)).hmrc_body["periodIncome"]["turnover"] == 0.0
    out = itsa.se_summary([undo], "2026-27", date(2026, 7, 5))
    assert not out.can_submit and "below zero" in out.blockers[0]


def test_period_end_outside_tax_year_and_old_years() -> None:
    with pytest.raises(ValueError, match="outside"):
        itsa.se_summary([], "2026-27", date(2027, 4, 6))
    old = itsa.se_summary([], "2024-25", date(2024, 7, 5))
    assert not old.can_submit and "2025-26" in old.blockers[0]


def test_uncoded_pl_names_are_listed_not_dropped_quietly() -> None:
    e = ev("a", date(2026, 5, 1), "purchase", Line(account="Lunch", debit=5_00), Line(account="1200 Bank", credit=5_00))
    s = itsa.se_summary([e], "2026-27", date(2026, 7, 5))
    assert s.unmapped_accounts == ["Lunch"] and "Lunch" in s.warnings[-1]


# --- reports -------------------------------------------------------------------------------------------------


def test_sa103_report_renders_every_format() -> None:
    r = rep.sa103_report(itsa.se_summary(BOOKS, "2026-27", date(2026, 10, 5)), "Sam Builds")
    boxes = r.sections[0]
    assert boxes.rows[0][:4] == [15, "Turnover: takings, fees, sales", "turnover", 1500.5]
    lines = r.sections[2].rows
    assert len(lines) == sum(len(f.lines) for f in itsa.se_summary(BOOKS, "2026-27", date(2026, 10, 5)).figures)

    text = rep.render(r, "csv").decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0] == ["Self-employment summary (SA103-style)"] and ["Boxes"] in rows

    wb = load_workbook(io.BytesIO(rep.render(r, "xlsx")))
    assert wb.sheetnames == ["About", "Boxes", "Totals", "Ledger lines behind each box"]
    assert wb["Boxes"]["D2"].value == 1500.5 and "£" in wb["Boxes"]["D2"].number_format

    pdf = rep.render(r, "pdf")
    assert pdf.startswith(b"%PDF") and len(pdf) > 1000


def test_csv_neutralises_formulas() -> None:
    e = ev(
        "a",
        date(2026, 5, 1),
        "purchase",
        Line(account="7500 Office", debit=5_00),
        Line(account="1200 Bank", credit=5_00),
        counterparty='=HYPERLINK("x")',
    )
    text = rep.to_csv(rep.sa103_report(itsa.se_summary([e], "2026-27", date(2026, 7, 5)))).decode("utf-8-sig")
    assert "'=HYPERLINK" in text and ",=HYPERLINK" not in text


HMRC_CALC = {  # trimmed from HMRC's example, Individual Calculations 8.0 examples/retrieve/def4/retrieve_response.json
    "metadata": {
        "calculationId": "f2fb30e5-4ab6-4a29-b3c1-c7264259ff1c",
        "taxYear": "2026-27",
        "calculationType": "in-year",
        "calculationTimestamp": "2025-04-15T09:35:15.094Z",
    },
    "calculation": {
        "taxCalculation": {
            "incomeTax": {
                "totalIncomeReceivedFromAllSources": 12500,
                "totalAllowancesAndDeductions": 12500,
                "totalTaxableIncome": 12500,
                "incomeTaxCharged": 5000.99,
            },
            "nics": {"class2Nics": {"amount": 5000.99}, "class4Nics": {"totalAmount": 5000.99}, "totalNic": 5000.99},
            "totalIncomeTaxAndNicsDue": 5000.99,
        }
    },
    "messages": {"warnings": [{"id": "C55001", "text": "a warning"}]},
}


def test_calculation_report_copies_hmrc_figures_only() -> None:
    r = rep.calculation_report(HMRC_CALC)
    assert r.source.startswith("HMRC") and "Accnomos has not worked out any tax" in r.notes[0]
    figures = dict((k, v) for k, v in r.sections[0].rows)  # type: ignore[misc]
    assert figures["Total Income Tax and NIC due"] == 5000.99 and figures["Class 2 NIC"] == 5000.99
    assert "HMRC warning: a warning" in r.notes
    sparse = rep.calculation_summary({"metadata": {}, "calculation": {"taxCalculation": {"incomeTax": {}}}})
    assert dict(sparse["figures"])["Total Income Tax and NIC due"] is None  # never filled in by us
    assert rep.render(r, "pdf").startswith(b"%PDF")


# --- routes --------------------------------------------------------------------------------------------------


@pytest.fixture
async def books() -> Books:
    b = Books(FakeStorageProvider())
    await b.ensure_tree()
    for e in BOOKS:
        await b.post_event(e)
    return b


@pytest.fixture
def client(books: Books):
    async def override():
        yield books

    app.dependency_overrides[get_books] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


def _envelope(**kw) -> dict:
    return {
        "hmrc": None,
        "sent_headers": {"gov-client-connection-method": "WEB_APP_VIA_SERVER"},
        "fraud_headers_missing": [],
        "correlation_id": "corr-1",
        "receipt_id": None,
        "receipt_timestamp": None,
        **kw,
    }


def test_summary_route(client: TestClient) -> None:
    r = client.get("/books/self-employment/summary?tax_year=2026-27&period_end=2026-10-05", headers=H)
    assert r.status_code == 200, r.text
    assert r.json()["accounting_basis"] == "cash"  # HMRC's default: the two invoices are not paid yet
    assert r.json()["income_pence"] == 25_00
    r = client.get("/books/self-employment/summary?tax_year=2026-27&period_end=2026-10-05&basis=traditional", headers=H)
    assert r.json()["income_pence"] == 1_525_50 and r.json()["accounting_basis"] == "traditional"
    assert client.get("/books/self-employment/summary?tax_year=2026-28", headers=H).status_code == 422
    assert (
        client.get("/books/self-employment/summary?tax_year=2026-27&period_end=2027-05-01", headers=H).status_code
        == 422
    )


def test_cumulative_receipt_is_checked_against_the_books_and_audited(client: TestClient, books: Books) -> None:
    body = client.get("/books/self-employment/summary?tax_year=2026-27&period_end=2026-10-05", headers=H).json()
    receipt = {
        "kind": "cumulative",
        "tax_year": "2026-27",
        "nino": "TC663795B",
        "business_id": "XAIS12345678910",
        "period_end": "2026-10-05",
        "sent": body["hmrc_body"],
        "hmrc": _envelope(),
        "confirmed_at": "2026-10-28T09:00:00Z",
    }
    r = client.post("/books/self-employment/receipts", json=receipt, headers=H)
    assert r.status_code == 201, r.text
    assert r.json()["matches_books"] is True
    path = r.json()["path"]
    assert path.startswith(f"{ITSA}/2026-27/XAIS12345678910/update-to-2026-10-05-")

    tampered = {**receipt, "sent": {**body["hmrc_body"], "periodIncome": {"turnover": 1.0, "other": 0.0}}}
    assert client.post("/books/self-employment/receipts", json=tampered, headers=H).json()["matches_books"] is False

    listed = client.get("/books/self-employment/receipts?tax_year=2026-27", headers=H).json()["receipts"]
    assert [x["matches_books"] for x in listed] == [False, True]  # newest first

    audit = client.get("/audit?kind=itsa_update_submitted", headers=H)
    assert audit.status_code == 200, audit.text
    texts = [e["summary"] for e in audit.json()["entries"]]
    assert len(texts) == 2 and "did not match the books" in texts[0] and "HMRC sandbox" in texts[1]

    rep_json = client.get(f"/books/self-employment/reports/quarterly?tax_year=2026-27&path={path}", headers=H).json()
    assert rep_json["kind"] == "quarterly" and rep_json["sections"][1]["rows"][2] == ["HMRC correlation ID", "corr-1"]
    pdf = client.get(f"/books/self-employment/reports/quarterly?tax_year=2026-27&path={path}&format=pdf", headers=H)
    assert pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
    bad = client.get("/books/self-employment/reports/quarterly?tax_year=2026-27&path=_System/audit/x.json", headers=H)
    assert bad.status_code == 400


def test_receipt_validation(client: TestClient) -> None:
    base = {"tax_year": "2026-27", "nino": "TC663795B", "hmrc": _envelope()}
    assert (
        client.post("/books/self-employment/receipts", json={**base, "kind": "cumulative"}, headers=H).status_code
        == 422
    )
    final = {
        **base,
        "kind": "final_declaration",
        "calculation_id": "f2fb30e5-4ab6-4a29-b3c1-c7264259ff1c",
        "confirmed_at": "2027-05-01T09:00:00Z",
    }
    assert (
        client.post("/books/self-employment/receipts", json={**final, "declaration": "I agree"}, headers=H).status_code
        == 422
    )
    ok = client.post(
        "/books/self-employment/receipts", json={**final, "declaration": itsa.FINAL_DECLARATION}, headers=H
    )
    assert ok.status_code == 201, ok.text
    no_user = client.post(
        "/books/self-employment/receipts",
        json={**final, "declaration": itsa.FINAL_DECLARATION},
        headers={"X-Graph-Token": "t"},
    )
    assert no_user.status_code == 400
    assert (
        client.post(
            "/books/self-employment/receipts", json={**base, "kind": "calculation", "calculation_id": "x"}, headers=H
        ).status_code
        == 422
    )


def test_calculation_saved_then_reported(client: TestClient) -> None:
    cid = HMRC_CALC["metadata"]["calculationId"]
    r = client.post(
        "/books/self-employment/receipts",
        json={
            "kind": "calculation",
            "tax_year": "2026-27",
            "nino": "TC663795B",
            "calculation_id": cid,
            "hmrc": _envelope(hmrc=HMRC_CALC),
        },
        headers=H,
    )
    assert r.status_code == 201, r.text
    out = client.get(f"/books/self-employment/reports/calculation?tax_year=2026-27&calculation_id={cid}", headers=H)
    assert out.json()["source"].startswith("HMRC")
    x = client.get(
        f"/books/self-employment/reports/calculation?tax_year=2026-27&calculation_id={cid}&format=xlsx", headers=H
    )
    assert x.headers["content-disposition"].endswith('.xlsx"')
    assert client.get(
        "/books/self-employment/reports/sa103?tax_year=2026-27&period_end=2026-10-05&format=csv", headers=H
    ).content.startswith(b"\xef\xbb\xbfSelf-employment summary")
