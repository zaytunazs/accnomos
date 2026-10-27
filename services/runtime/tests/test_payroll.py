"""Pay runs from timesheets (app/payroll.py): estimates, export, confirm. Against FakeStorageProvider."""

from __future__ import annotations

import csv
import io
import sys
from datetime import date
from pathlib import Path

import pytest
from accnomos_connectors import FakeStorageProvider
from accnomos_ledger.events import LedgerEvent

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import payroll as pr  # noqa: E402
from app import payroll_rates as R  # noqa: E402
from app.books import EVENTS, Books  # noqa: E402
from app.jobs import Job, StaffMember, TimesheetEntry, add_entry, save_job, save_staff  # noqa: E402

USER = "sam@example.co.uk"


@pytest.fixture
def fake() -> FakeStorageProvider:
    return FakeStorageProvider()


@pytest.fixture
def books(fake: FakeStorageProvider) -> Books:
    return Books(fake)


def _events(fake: FakeStorageProvider) -> list[LedgerEvent]:
    return [LedgerEvent.model_validate_json(d) for p, d in fake.files.items() if p.startswith(f"{EVENTS}/")]


async def _setup(books: Books) -> tuple[StaffMember, StaffMember]:
    job = await save_job(books, Job(name="Kitchen refit"))
    ann = await save_staff(books, StaffMember(name="Ann Lee", role="Joiner", hourly_rate_pence=1600))
    bob = await save_staff(books, StaffMember(name="Bob Kaur", role="Labourer", hourly_rate_pence=1250, active=False))
    for d, who, h in (
        ("2026-10-01", ann, 8),
        ("2026-10-02", ann, 7.5),
        ("2026-10-02", bob, 4.25),
        ("2026-09-30", ann, 8),
        ("2026-11-01", bob, 3),
    ):
        await add_entry(books, TimesheetEntry(staff_id=who.id, job_id=job.id, date=d, hours=h))
    return ann, bob


# --------------------------------------------------------------------------- #
# Estimates: monthly pro-rating and thresholds
# --------------------------------------------------------------------------- #


def test_rates_table_cites_its_source_and_says_estimate() -> None:
    assert R.SOURCE_URL.startswith("https://www.gov.uk/") and R.TAX_YEAR in R.SOURCE_URL.replace("-to-20", "-")
    assert R.ESTIMATE_LABEL == "Estimate — your payroll provider calculates the actual tax and NI"
    assert any("1257L" in x for x in R.LIMITS) and any("Student loans" in x and "pensions" in x for x in R.LIMITS)


def test_typical_month_worked_example() -> None:
    # £2,000 a month: tax 20% of (2,000 − 12,570 ÷ 12, rounded down to £952) = £190.40, employee NI £76.20;
    # employer NI 15% above £5,000 ÷ 12. (HMRC's own 1257L tables give £1,048.25 free pay, so £190.20 tax:
    # the estimate is 20p high, which is why it is only ever an estimate.)
    e = pr.estimate(200_000)
    assert (e.paye_pence, e.ee_ni_pence, e.er_ni_pence, e.net_pence) == (19_040, 7_620, 23_750, 173_340)


def test_annual_thresholds_are_pro_rated_monthly() -> None:
    free = R.PERSONAL_ALLOWANCE // 12  # £1,047.50: exactly divisible
    assert free * 12 == R.PERSONAL_ALLOWANCE
    at = pr.estimate(free)
    assert (at.paye_pence, at.ee_ni_pence) == (0, 0)
    one_pound_over = pr.estimate(free + 100)
    assert (one_pound_over.paye_pence, one_pound_over.ee_ni_pence) == (20, 8)
    assert pr.estimate(free + 99).paye_pence == 0  # taxable pay is rounded down to whole pounds
    # Employer NI starts above £5,000 ÷ 12 = £416.666...
    assert pr.estimate(41_666).er_ni_pence == 0 and pr.estimate(41_667).er_ni_pence == 0
    assert pr.estimate(41_767).er_ni_pence == 15
    assert pr.estimate(0) == pr.Estimate(paye_pence=0, ee_ni_pence=0, er_ni_pence=0, net_pence=0)


@pytest.mark.parametrize(
    ("gross", "tax_rate", "ee_rate"),
    [
        (200_000, 20, 8),  # basic rate, main NI rate
        (600_000, 40, 2),  # higher rate (above £50,270 ÷ 12), above the UEL
        (1_100_000, 45, 2),  # additional rate (above £125,140 ÷ 12)
    ],
)
def test_marginal_rates_by_band(gross: int, tax_rate: int, ee_rate: int) -> None:
    lo, hi = pr.estimate(gross), pr.estimate(gross + 10_000)  # £100 more
    assert hi.paye_pence - lo.paye_pence == tax_rate * 100
    assert hi.ee_ni_pence - lo.ee_ni_pence == ee_rate * 100
    assert hi.er_ni_pence - lo.er_ni_pence == 1_500


def test_higher_rate_month_worked_example() -> None:
    e = pr.estimate(600_000)  # £6,000
    assert (e.paye_pence, e.ee_ni_pence, e.er_ni_pence) == (135_246, 28_755, 83_750)


def test_periods_and_hours() -> None:
    assert pr.period_bounds("2026-02") == (date(2026, 2, 1), date(2026, 2, 28))
    assert pr.period_bounds("2028-02")[1] == date(2028, 2, 29)
    for bad in ("2026-13", "2026-1", "../x", "2026-10-01", ""):
        with pytest.raises(pr.PayrollError):
            pr.period_bounds(bad)
    assert pr.month_quarters("37.6") == 150 and pr.month_quarters(0) == 0
    for bad in (-1, 745, "lots", "nan"):
        with pytest.raises(pr.PayrollError):
            pr.month_quarters(bad)


# --------------------------------------------------------------------------- #
# Prepare and edit
# --------------------------------------------------------------------------- #


async def test_prepare_from_the_months_timesheets(books: Books, fake: FakeStorageProvider) -> None:
    ann, bob = await _setup(books)
    run = await pr.prepare(books, "2026-10", None, USER)
    assert (run.period_start, run.period_end, run.pay_date, run.status) == (
        date(2026, 10, 1),
        date(2026, 10, 31),
        date(2026, 10, 31),
        "draft",
    )
    a, b = run.lines
    assert (a.staff_id, a.hours, a.rate_pence, a.gross_pence) == (ann.id, 15.5, 1600, 24_800)
    assert (b.staff_id, b.hours, b.gross_pence, b.notes) == (
        bob.id,
        4.25,
        5_313,
        "Marked inactive",
    )  # 4.25 × £12.50 = £53.125
    assert a.est_paye_pence == 0 and a.est_net_pence == a.gross_pence  # under the allowance
    saved = pr.PayRun.model_validate_json(fake.files[f"{pr.RUNS}/2026-10.json"])
    assert saved.estimate == R.ESTIMATE_LABEL and saved.limits == list(R.LIMITS)  # flagged in the file itself
    assert saved.confirmed_gross_pence is None

    empty = await pr.prepare(books, "2026-08", date(2026, 8, 28), USER)
    assert (
        empty.lines == []
        and empty.warnings == ["No timesheet hours in August 2026"]
        and empty.pay_date == date(2026, 8, 28)
    )


async def test_edit_recomputes_and_refuses_strangers(books: Books) -> None:
    ann, bob = await _setup(books)
    await pr.prepare(books, "2026-10", None, USER)
    run = await pr.update(
        books,
        "2026-10",
        [pr.LineEdit(staff_id=ann.id, hours="160", rate_pence=1250, notes=" ok ")],
        date(2026, 10, 28),
        USER,
    )
    (line,) = run.lines
    assert (line.gross_pence, line.notes, run.pay_date) == (200_000, "ok", date(2026, 10, 28))
    assert (line.est_paye_pence, line.est_ee_ni_pence, line.est_er_ni_pence) == (19_040, 7_620, 23_750)
    with pytest.raises(pr.PayrollError):
        await pr.update(books, "2026-10", [pr.LineEdit(staff_id="f" * 12, hours=1, rate_pence=1)], None, USER)
    with pytest.raises(pr.PayrollError):
        await pr.update(books, "2026-10", [pr.LineEdit(staff_id=ann.id, hours=1, rate_pence=1)] * 2, None, USER)
    with pytest.raises(FileNotFoundError):
        await pr.update(books, "2026-07", [], None, USER)


# --------------------------------------------------------------------------- #
# Export
# --------------------------------------------------------------------------- #


async def test_export_shape(books: Books, fake: FakeStorageProvider) -> None:
    ann, bob = await _setup(books)
    await pr.prepare(books, "2026-10", None, USER)
    path, text = await pr.export(books, "2026-10", USER)
    assert path == "05 Payroll/2026-10.csv"
    assert fake.files[path].decode("utf-8-sig") == text
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0] == [R.ESTIMATE_LABEL]
    head = rows.index(list(pr.CSV_HEADER))
    body = rows[head + 1 :]
    assert [r[0] for r in body] == ["Ann Lee", "Bob Kaur", "Total"]  # one row per employee, then totals
    assert body[0][:5] == ["Ann Lee", "Joiner", "15.5", "16.00", "248.00"]
    assert body[1][4] == "53.13" and body[2][4] == "301.13"
    assert all(len(r) == len(pr.CSV_HEADER) for r in body)
    assert all(h.startswith("Est.") for h in pr.CSV_HEADER[5:9])  # deductions are labelled as estimates


def test_export_keeps_typed_text_inert_and_big_totals_exact() -> None:
    lines = [pr.make_line(f"s{i}", 744 * 4, 1000, '=HYPERLINK("http://x")' if i == 0 else "") for i in range(17)]
    lines.append(pr.make_line("s99", 1, 1000))  # 0.25 h: total 12,648.25 h
    run = pr.PayRun(
        period_start=date(2026, 10, 1), period_end=date(2026, 10, 31), pay_date=date(2026, 10, 31), lines=lines
    )
    rows = list(csv.reader(io.StringIO(pr.export_csv(run, {"s0": ("+Ann", "@Joiner")}))))
    first = rows[rows.index(list(pr.CSV_HEADER)) + 1]
    assert (first[0], first[1], first[-1]) == ("'+Ann", "'@Joiner", '\'=HYPERLINK("http://x")')
    assert rows[-1][2] == "12648.25" and rows[-2][2] == "0.25"


# --------------------------------------------------------------------------- #
# Confirm and reopen
# --------------------------------------------------------------------------- #


async def test_confirm_posts_balanced_staff_costs(books: Books, fake: FakeStorageProvider) -> None:
    await _setup(books)
    await pr.prepare(books, "2026-10", date(2026, 10, 30), USER)
    run = await pr.confirm(books, "2026-10", 30_500, 1_234, USER)  # the provider's figures, not the estimate
    assert (run.status, run.confirmed_gross_pence, run.confirmed_er_ni_pence, run.confirmed_by) == (
        "confirmed",
        30_500,
        1_234,
        USER,
    )
    events = {e.id: e for e in _events(fake)}
    assert set(events) == set(run.event_ids) and len(events) == 2
    for ev in events.values():
        assert ev.kind == "adjustment" and ev.date == date(2026, 10, 30) and ev.approved_by == USER
        assert sum(ln.debit for ln in ev.lines) == sum(ln.credit for ln in ev.lines)
    lines = [(ln.account, ln.debit, ln.credit) for ev in events.values() for ln in ev.lines]
    assert sorted(lines) == sorted(
        [
            (pr.WAGES, 30_500, 0),
            (pr.LIABILITIES, 0, 30_500),
            (pr.EMPLOYER_NI, 1_234, 0),
            (pr.LIABILITIES, 0, 1_234),
        ]
    )
    assert "Books.xlsx" in fake.files

    with pytest.raises(pr.PayrollConflict):
        await pr.confirm(books, "2026-10", 30_500, 1_234, USER)  # never posted twice
    with pytest.raises(pr.PayrollConflict):
        await pr.prepare(books, "2026-10", None, USER)
    with pytest.raises(pr.PayrollConflict):
        await pr.update(books, "2026-10", [], None, USER)


async def test_confirm_without_employer_ni_posts_one_entry(books: Books, fake: FakeStorageProvider) -> None:
    await _setup(books)
    await pr.prepare(books, "2026-10", None, USER)
    with pytest.raises(pr.PayrollError):
        await pr.confirm(books, "2026-10", 0, 0, USER)
    run = await pr.confirm(books, "2026-10", 30_113, 0, USER)
    assert len(run.event_ids) == 1 and len(_events(fake)) == 1


async def test_failed_second_entry_undoes_the_first(books: Books, fake: FakeStorageProvider, monkeypatch) -> None:
    await _setup(books)
    await pr.prepare(books, "2026-10", None, USER)
    real, calls = books.post_event, []

    async def flaky(ev: LedgerEvent) -> str:
        calls.append(ev)
        if len(calls) == 2:
            raise OSError("network")
        return await real(ev)

    monkeypatch.setattr(books, "post_event", flaky)
    with pytest.raises(OSError):
        await pr.confirm(books, "2026-10", 30_000, 1_000, USER)
    monkeypatch.setattr(books, "post_event", real)
    evs = _events(fake)
    assert [e.kind for e in sorted(evs, key=lambda e: e.kind)] == ["adjustment", "reversal"]  # wages, then its undo
    assert (await pr.get_run(books, "2026-10")).status == "draft"


async def test_failed_save_after_posting_undoes_both_entries(
    books: Books, fake: FakeStorageProvider, monkeypatch
) -> None:
    await _setup(books)
    await pr.prepare(books, "2026-10", None, USER)

    async def broken(*_a, **_k):
        raise OSError("network")

    monkeypatch.setattr(pr, "_save", broken)
    with pytest.raises(OSError):
        await pr.confirm(books, "2026-10", 30_000, 1_000, USER)
    monkeypatch.undo()
    evs = _events(fake)
    originals = {e.id for e in evs if e.kind == "adjustment"}
    assert len(originals) == 2 and {e.reverses for e in evs if e.kind == "reversal"} == originals
    assert (await pr.get_run(books, "2026-10")).status == "draft"  # so a second confirm posts once, cleanly


async def test_reopen_reverses_and_allows_a_new_confirm(books: Books, fake: FakeStorageProvider) -> None:
    await _setup(books)
    await pr.prepare(books, "2026-10", None, USER)
    first = await pr.confirm(books, "2026-10", 99_999, 1_000, USER)  # mistyped
    run = await pr.reopen(books, "2026-10", USER)
    assert run.status == "draft" and run.confirmed_gross_pence is None and run.undone_event_ids == first.event_ids
    reversed_ids = {e.reverses for e in _events(fake) if e.kind == "reversal"}
    assert reversed_ids == set(first.event_ids)
    with pytest.raises(pr.PayrollConflict):
        await pr.reopen(books, "2026-10", USER)
    again = await pr.confirm(books, "2026-10", 30_113, 1_000, USER)
    assert again.status == "confirmed" and not set(again.event_ids) & set(first.event_ids)


async def test_books_xlsx_payroll_sheet(books: Books, fake: FakeStorageProvider) -> None:
    from openpyxl import load_workbook

    await _setup(books)
    await pr.prepare(books, "2026-10", None, USER)
    await pr.confirm(books, "2026-10", 30_500, 1_234, USER)
    ws = load_workbook(io.BytesIO(fake.files["Books.xlsx"]))["Payroll"]
    cells = [list(row) for row in ws.iter_rows(values_only=True)]
    assert cells[0][7] == "Est. PAYE £" and [R.ESTIMATE_LABEL] == [c for c in cells[-len(R.LIMITS) - 1] if c]
    flat = [c for row in cells for c in row if c is not None]
    assert "Ann Lee" in flat and "confirmed" in flat and 305.0 in flat


# --------------------------------------------------------------------------- #
# HTTP
# --------------------------------------------------------------------------- #

H = {"X-Graph-Token": "tok", "X-User": USER}


@pytest.fixture
def client(fake: FakeStorageProvider):
    from fastapi.testclient import TestClient

    from app.main import app
    from app.routes_books import get_books

    async def override() -> Books:
        return Books(fake)

    app.dependency_overrides[get_books] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


def test_payroll_routes(client, books: Books, fake: FakeStorageProvider) -> None:
    import asyncio

    ann, _ = asyncio.run(_setup(books))
    r = client.get("/payroll/runs", headers=H)
    assert r.status_code == 200 and r.json()["runs"] == [] and r.json()["estimate"] == R.ESTIMATE_LABEL
    assert client.post("/payroll/runs/prepare", headers=H, json={"period": "2026-13"}).status_code == 400
    assert (
        client.post("/payroll/runs/prepare", headers={"X-Graph-Token": "tok"}, json={"period": "2026-10"}).status_code
        == 400
    )
    r = client.post("/payroll/runs/prepare", headers=H, json={"period": "2026-10"})
    assert r.status_code == 200 and len(r.json()["lines"]) == 2 and r.json()["estimate"] == R.ESTIMATE_LABEL
    assert client.get("/payroll/runs/2026-10", headers=H).json()["status"] == "draft"
    assert client.get("/payroll/runs/2026-07", headers=H).status_code == 404

    r = client.put(
        "/payroll/runs/2026-10", headers=H, json={"lines": [{"staff_id": ann.id, "hours": 160, "rate_pence": 1250}]}
    )
    assert r.status_code == 200 and r.json()["lines"][0]["est_paye_pence"] == 19_040
    assert (
        client.put(
            "/payroll/runs/2026-10", headers=H, json={"lines": [{"staff_id": ann.id, "hours": 999, "rate_pence": 1}]}
        ).status_code
        == 400
    )

    r = client.post("/payroll/runs/2026-10/export", headers=H)
    assert r.status_code == 200 and r.json()["path"] == "05 Payroll/2026-10.csv" and R.ESTIMATE_LABEL in r.json()["csv"]

    bad = {"actual_gross_pence": 0, "actual_er_ni_pence": 0}
    assert client.post("/payroll/runs/2026-10/confirm", headers=H, json=bad).status_code == 422
    body = {"actual_gross_pence": 200_000, "actual_er_ni_pence": 23_750}
    r = client.post("/payroll/runs/2026-10/confirm", headers=H, json=body)
    assert r.status_code == 200 and r.json()["status"] == "confirmed" and len(_events(fake)) == 2
    assert client.post("/payroll/runs/2026-10/confirm", headers=H, json=body).status_code == 409
    assert client.post("/payroll/runs/2026-10/reopen", headers=H).json()["status"] == "draft"
    assert client.post("/payroll/runs/2026-10/reopen", headers=H).status_code == 409
    assert [x["period_start"] for x in client.get("/payroll/runs", headers=H).json()["runs"]] == ["2026-10-01"]
