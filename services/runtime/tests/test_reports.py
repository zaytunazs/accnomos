"""Reports (app/reports.py, app/routes_reports.py) on a hand-built ledger whose totals are worked out below.

The ledger (pence; S = standard 20 %, N = no VAT invoice):

    #   date        kind        lines                                                  net      VAT
    A   30 Jun 26   purchase    Dr 5000 Materials 8,000 S, Dr 2200 1,600, Cr 1200      8,000    1,600
    B    1 Jul 26   sale        Dr 1100 120,000, Cr 4000 Sales 100,000 S, Cr 2200      100,000  20,000   Acme, unpaid
    C   10 Jul 26   purchase    Dr 5000 Materials 25,000 S, Dr 2200 5,000, Cr 2100     25,000   5,000    Howdens, unpaid
    D   15 Aug 26   sale        Dr 1200 36,000, Cr 4000 Sales 30,000 S, Cr 2200        30,000   6,000    paid at once
    E   20 Aug 26   purchase    Dr 7500 Office 4,325 S, Dr 2200 865, Cr 1200           4,325    865
    F    5 Sep 26   purchase    Dr 7300 Motor 6,000 N, Cr 1210 Cash                    6,000    0
    G   12 Sep 26   purchase    Dr 5100 Tools 10,000 S, Dr 2200 2,000, Cr 1200         10,000   2,000
    G'  13 Sep 26   reversal    of G (every line swapped)                              10,000   2,000
    H   30 Sep 26   adjustment  Dr 7000 Wages 50,000, Dr 7010 Employer's NI 3,000, Cr 2210 53,000 (payroll)
    I    2 Oct 26   sale        Dr 1100 60,000, Cr 4000 Sales 50,000 S, Cr 2200        50,000   10,000   Acme, unpaid

P&L, 1 Jul – 30 Sep 2026 (B to H):
    Sales           4000                  100,000 + 30,000            = 130,000
    Cost of sales   5000                                                 25,000
                    5100                  10,000 − 10,000 = 0          (left out)
                    total                                                25,000
    Gross profit                          130,000 − 25,000             = 105,000
    Overheads       7000 50,000 · 7010 3,000 · 7300 6,000 · 7500 4,325 = 63,325
    Net profit                            105,000 − 63,325             = 41,675   (£416.75)
    Entries touching a P&L account: B C D E F G G' H                    = 8

VAT, calendar quarters (H is not on the return, being an adjustment):
    Q2 Apr–Jun  A:          box 1 0       box 4 1,600                net −1,600   box 6 £0     box 7 £80
    Q3 Jul–Sep  B C D E F G G':
                box 1 20,000 + 6,000 = 26,000
                box 4 5,000 + 865 + 0 + 2,000 − 2,000 = 5,865        net 20,135
                box 6 130,000 → £1,300
                box 7 25,000 + 4,325 + 6,000 + 10,000 − 10,000 = 35,325 → £353 (pence dropped)
    Q4 Oct–Dec  I:          box 1 10,000                             net 10,000   box 6 £500
    Q1 Jan–Mar  nothing.    Total net over the 4 quarters: −1,600 + 20,135 + 10,000 = 28,535

Trend, Nov 2025 – Oct 2026 (sales / purchases = cost of sales + overheads / net):
    Jun 0 / 8,000 / −8,000 · Jul 100,000 / 25,000 / 75,000 · Aug 30,000 / 4,325 / 25,675
    Sep 0 / 6,000 + 0 + 53,000 = 59,000 / −59,000 · Oct 50,000 / 0 / 50,000 · the rest 0
    Jul + Aug + Sep net = 75,000 + 25,675 − 59,000 = 41,675, the Q3 net profit.
    Twelve-month totals: sales 180,000; purchases 8,000 + 25,000 + 4,325 + 59,000 = 96,325; net 83,675.

Aged at 27 Oct 2026: debtors Acme B 120,000 (118 days, 90+) and I 60,000 (25 days, 0-30) = 180,000;
creditors Howdens C 30,000 (109 days, 90+).
"""

from __future__ import annotations

import io
from datetime import UTC, date, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from accnomos_connectors import FakeStorageProvider
from accnomos_ledger import LedgerEvent, Line, VatCode
from app import reports as rp
from app.books import AUDIT, EVENTS, Books
from app.paths import REPORTS
from app.projection import SAMPLE_LABEL

USER = "sam@example.co.uk"
H = {"X-Graph-Token": "tok", "X-User": USER}
TODAY = date(2026, 10, 27)
S, N = VatCode.STANDARD, VatCode.NO_VAT
Q3 = (date(2026, 7, 1), date(2026, 9, 30))


def _ev(
    eid: str,
    kind: str,
    when: date,
    lines: list[Line],
    net: int | None = None,
    vat: int | None = None,
    who: str | None = None,
    approved_by: str = USER,
    reverses: str | None = None,
) -> LedgerEvent:
    return LedgerEvent(
        id=eid * 32,
        kind=kind,  # type: ignore[arg-type]
        date=when,
        lines=lines,
        net=net,
        vat=vat,
        counterparty=who,
        reverses=reverses,
        approved_by=approved_by,
        approved_at=datetime(2026, 10, 1, tzinfo=UTC),
    )


def _sale(eid: str, when: date, net: int, settle: str, who: str | None = None) -> LedgerEvent:
    vat = net // 5
    lines = [
        Line(account=settle, debit=net + vat),
        Line(account="4000 Sales", credit=net, vat_code=S),
        Line(account="2200 VAT control", credit=vat),
    ]
    return _ev(eid, "sale", when, lines, net, vat, who)


def _purchase(
    eid: str, when: date, account: str, net: int, vat: int, settle: str, who: str | None = None
) -> LedgerEvent:
    lines = [Line(account=account, debit=net, vat_code=N if vat == 0 else S)]
    if vat:
        lines.append(Line(account="2200 VAT control", debit=vat))
    lines.append(Line(account=settle, credit=net + vat))
    return _ev(eid, "purchase", when, lines, net, vat, who)


def _reversal(eid: str, when: date, of: LedgerEvent) -> LedgerEvent:
    swapped = [Line(account=ln.account, debit=ln.credit, credit=ln.debit, vat_code=ln.vat_code) for ln in of.lines]
    return _ev(eid, "reversal", when, swapped, of.net, of.vat, of.counterparty, reverses=of.id)


def ledger(approved_by: str = USER) -> list[LedgerEvent]:
    g = _purchase("7", date(2026, 9, 12), "5100 Tools & small equipment", 10_000, 2_000, "1200 Bank")
    events = [
        _purchase("1", date(2026, 6, 30), "5000 Materials", 8_000, 1_600, "1200 Bank"),
        _sale("2", date(2026, 7, 1), 100_000, "1100 Debtors", "Acme Kitchens Ltd"),
        _purchase("3", date(2026, 7, 10), "5000 Materials", 25_000, 5_000, "2100 Creditors", "Howdens Joinery"),
        _sale("4", date(2026, 8, 15), 30_000, "1200 Bank"),
        _purchase("5", date(2026, 8, 20), "7500 Office supplies", 4_325, 865, "1200 Bank"),
        _purchase("6", date(2026, 9, 5), "7300 Motor expenses", 6_000, 0, "1210 Cash"),
        g,
        _reversal("8", date(2026, 9, 13), g),
        _ev(
            "9",
            "adjustment",
            date(2026, 9, 30),
            [
                Line(account="7000 Gross wages", debit=50_000),
                Line(account="7010 Employer's NI", debit=3_000),
                Line(account="2210 Payroll liabilities", credit=53_000),
            ],
        ),
        _sale("a", date(2026, 10, 2), 50_000, "1100 Debtors", "Acme Kitchens Ltd"),
    ]
    if approved_by != USER:
        events[0] = events[0].model_copy(update={"approved_by": approved_by})
    return events


# --------------------------------------------------------------------------- #
# Pure functions
# --------------------------------------------------------------------------- #


def test_accounts_are_placed_by_their_code() -> None:
    assert rp.section_of("4000 Sales") == "sales"
    assert rp.section_of("5100 Tools & small equipment") == "cost_of_sales"
    assert rp.section_of("6200 Advertising") == rp.section_of("7999 Sundry expenses") == "overheads"
    for balance_sheet in ("1100 Debtors", "1200 Bank", "2200 VAT control", "2300 Owner / director loan", "Suspense"):
        assert rp.section_of(balance_sheet) is None


def test_profit_and_loss_for_q3() -> None:
    pl = rp.profit_and_loss(ledger(), *Q3)
    assert pl.label == "Jul–Sep 2026" and pl.basis.startswith("As approved")
    assert [(r.account, r.amount_pence) for r in pl.sales.rows] == [("4000 Sales", 130_000)]
    assert [(r.account, r.amount_pence) for r in pl.cost_of_sales.rows] == [
        ("5000 Materials", 25_000)
    ]  # 5100 nets to 0
    assert pl.cost_of_sales.total_pence == 25_000 and pl.gross_profit_pence == 105_000
    assert [(r.account, r.amount_pence) for r in pl.overheads.rows] == [
        ("7000 Gross wages", 50_000),
        ("7010 Employer's NI", 3_000),
        ("7300 Motor expenses", 6_000),
        ("7500 Office supplies", 4_325),
    ]
    assert pl.overheads.total_pence == 63_325
    assert pl.net_profit_pence == 41_675
    assert pl.entries == 8 and pl.sample is False


def test_profit_and_loss_never_counts_vat_or_the_bank() -> None:
    pl = rp.profit_and_loss(ledger(), date(2026, 1, 1), date(2026, 12, 31))
    accounts = {r.account for sec in (pl.sales, pl.cost_of_sales, pl.overheads) for r in sec.rows}
    assert not {a for a in accounts if a[:1] in "123"}
    assert pl.sales.total_pence == 180_000  # B + D + I
    assert pl.net_profit_pence == 41_675 + 50_000 - 8_000  # Q3, plus I in October, less A in June


def test_an_empty_period_is_all_zeros() -> None:
    pl = rp.profit_and_loss(ledger(), date(2025, 1, 1), date(2025, 3, 31))
    assert pl.sales.rows == pl.cost_of_sales.rows == pl.overheads.rows == []
    assert pl.net_profit_pence == 0 and pl.entries == 0


def test_vat_summary_per_quarter() -> None:
    spans = [
        (date(2026, 1, 1), date(2026, 3, 31)),
        Q3,
        (date(2026, 4, 1), date(2026, 6, 30)),
        (date(2026, 10, 1), date(2026, 12, 31)),
    ]
    vs = rp.vat_summary(ledger(), spans)
    assert [p.period for p in vs.periods] == ["2026-Q1", "2026-Q2", "2026-Q3", "2026-Q4"]  # oldest first
    q1, q2, q3, q4 = vs.periods
    assert (q1.net_vat_pence, q1.entries) == (0, 0)
    assert (q2.vat_on_sales_pence, q2.vat_reclaimed_pence, q2.net_vat_pence, q2.purchases_ex_vat_pounds) == (
        0,
        1_600,
        -1_600,
        80,
    )
    assert (q3.vat_on_sales_pence, q3.vat_reclaimed_pence, q3.net_vat_pence) == (26_000, 5_865, 20_135)
    assert (q3.sales_ex_vat_pounds, q3.purchases_ex_vat_pounds, q3.entries) == (1_300, 353, 7)
    assert q3.label == "Jul–Sep 2026"
    assert (q4.net_vat_pence, q4.sales_ex_vat_pounds) == (10_000, 500)
    assert vs.total_net_vat_pence == 28_535


def test_monthly_trend_agrees_with_the_p_and_l() -> None:
    t = rp.monthly_trend(ledger(), TODAY)
    assert len(t.months) == 12 and t.months[0].month == "2025-11" and t.months[-1].month == "2026-10"
    by = {m.month: (m.sales_pence, m.purchases_pence, m.net_pence) for m in t.months}
    assert by["2026-06"] == (0, 8_000, -8_000)
    assert by["2026-07"] == (100_000, 25_000, 75_000)
    assert by["2026-08"] == (30_000, 4_325, 25_675)
    assert by["2026-09"] == (0, 59_000, -59_000)
    assert by["2026-10"] == (50_000, 0, 50_000)
    assert (
        sum(by[m][2] for m in ("2026-07", "2026-08", "2026-09")) == rp.profit_and_loss(ledger(), *Q3).net_profit_pence
    )
    assert all(v == (0, 0, 0) for k, v in by.items() if k < "2026-06")
    assert t.months[-1].label == "Oct 2026" and t.months[-1].end == date(2026, 10, 31)


def test_the_demo_books_are_the_sample() -> None:
    assert rp.is_sample(ledger()) is False
    assert rp.is_sample(ledger(approved_by=rp.SAMPLE_APPROVER)) is True
    # The flag is about the events handed in: GET /reports/pl reads only the period's, so a period without the
    # seeded entry would say False. (A, the one marked here, is in June.)
    sample = ledger(approved_by=rp.SAMPLE_APPROVER)
    assert rp.profit_and_loss(sample, *Q3).sample is True
    assert rp.profit_and_loss([e for e in sample if e.date >= Q3[0]], *Q3).sample is False
    assert rp.monthly_trend(ledger(approved_by=rp.SAMPLE_APPROVER), TODAY).sample is True


def test_the_seed_script_approves_as_the_sample_approver() -> None:
    from pathlib import Path

    seed = (Path(__file__).resolve().parents[3] / "tools" / "seed_demo.py").read_text()
    assert f'USER = "{rp.SAMPLE_APPROVER}"' in seed


def test_report_workbook_sheets_and_figures() -> None:
    events = ledger()
    data, sheets = rp.build_report_xlsx(
        "all",
        pl=rp.profit_and_loss(events, *Q3),
        vat=rp.vat_summary(events, [Q3]),
        trend=rp.monthly_trend(events, TODAY),
    )
    assert sheets == ["About", "P&L", "VAT summary", "Trend"]  # no aged balances passed in
    wb = load_workbook(io.BytesIO(data))
    assert SAMPLE_LABEL not in [c.value for c in wb["About"]["A"]]
    pl = {r[0].strip(): r[1] for r in wb["P&L"].iter_rows(min_row=2, max_col=2, values_only=True) if r[0]}
    assert pl["4000 Sales"] == 1300 and pl["Gross profit"] == 1050 and pl["Net profit"] == 416.75
    assert wb["P&L"]["B3"].number_format.startswith('"£"')
    trend = list(wb["Trend"].iter_rows(min_row=2, values_only=True))
    assert trend[-1][0] == "Total" and trend[-1][1:] == (1800, 963.25, 836.75)


def test_sample_report_is_labelled_on_every_sheet() -> None:
    events = ledger(approved_by=rp.SAMPLE_APPROVER)
    data, _ = rp.build_report_xlsx("pl", pl=rp.profit_and_loss(events, *Q3), sample=True)
    wb = load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        assert SAMPLE_LABEL in [c.value for c in ws["A"]], ws.title


def test_report_text_is_never_a_formula() -> None:
    sale = _sale("b", date(2026, 7, 2), 1_000, "1200 Bank")
    sale.lines[1] = Line(account='4000 =HYPERLINK("x")', credit=1_000, vat_code=S)
    data, _ = rp.build_report_xlsx("pl", pl=rp.profit_and_loss([sale], *Q3))
    cells = [c.value for row in load_workbook(io.BytesIO(data))["P&L"].iter_rows() for c in row]
    assert not any(isinstance(v, str) and v.startswith("=") for v in cells)


# --------------------------------------------------------------------------- #
# Routes
# --------------------------------------------------------------------------- #


@pytest.fixture
def fake() -> FakeStorageProvider:
    return FakeStorageProvider()


@pytest.fixture
async def books(fake: FakeStorageProvider) -> Books:
    b = Books(fake)
    for ev in ledger():
        await b.post_event(ev)
    return b


@pytest.fixture
def client(books: Books):
    from app.main import app
    from app.routes_books import get_books

    async def override() -> Books:
        return books

    app.dependency_overrides[get_books] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


def test_report_routes_read_only(client: TestClient, fake: FakeStorageProvider) -> None:
    before = dict(fake.files)
    pl = client.get("/reports/pl?from=2026-07-01&to=2026-09-30", headers=H)
    assert pl.status_code == 200, pl.text
    assert pl.json()["net_profit_pence"] == 41_675 and pl.json()["label"] == "Jul–Sep 2026"

    current = client.get(f"/reports/pl?today={TODAY}", headers=H).json()  # default: the VAT period we are in now
    assert (current["start"], current["end"], current["net_profit_pence"]) == ("2026-10-01", "2026-12-31", 50_000)
    stagger2 = client.get(f"/reports/pl?today={TODAY}&stagger=2", headers=H).json()
    assert (stagger2["start"], stagger2["end"]) == ("2026-08-01", "2026-10-31")

    vat = client.get(f"/reports/vat?periods=4&today={TODAY}", headers=H).json()
    assert [p["label"] for p in vat["periods"]] == ["Jan–Mar 2026", "Apr–Jun 2026", "Jul–Sep 2026", "Oct–Dec 2026"]
    assert vat["total_net_vat_pence"] == 28_535

    trend = client.get(f"/reports/trend?today={TODAY}", headers=H).json()
    assert [m["net_pence"] for m in trend["months"]][-5:] == [-8_000, 75_000, 25_675, -59_000, 50_000]

    aged = client.get(f"/reports/aged?as_of={TODAY}", headers=H).json()
    assert aged["debtors"]["total_pence"] == 180_000
    assert aged["debtors"]["rows"][0]["buckets"] == {"0-30": 60_000, "31-60": 0, "61-90": 0, "90+": 120_000}
    assert aged["creditors"]["total_pence"] == 30_000 and aged["creditors"]["buckets"]["90+"] == 30_000
    assert aged["sample"] is False

    assert fake.files == before  # nothing written, not even an audit line


@pytest.mark.parametrize(
    "query",
    ["from=2026-07-01", "to=2026-07-01", "from=2026-09-30&to=2026-07-01", "from=30/09/2026&to=2026-10-01"],
)
def test_bad_periods_are_refused(client: TestClient, query: str) -> None:
    assert client.get(f"/reports/pl?{query}", headers=H).status_code == 422


def test_export_writes_one_file_in_07_reports_and_posts_nothing(client: TestClient, fake: FakeStorageProvider) -> None:
    events_before = {p for p in fake.files if p.startswith(EVENTS)}
    r = client.post(
        f"/reports/export?today={TODAY}", json={"report": "pl", "from": "2026-07-01", "to": "2026-09-30"}, headers=H
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["path"] == f"{REPORTS}/2026-10-27-profit-and-loss.xlsx" and body["sheets"] == ["About", "P&L"]
    assert body["sample"] is False and body["size"] == len(fake.files[body["path"]])
    wb = load_workbook(io.BytesIO(fake.files[body["path"]]))
    assert wb["P&L"]["A1"].value.startswith("Profit and loss · Jul–Sep 2026")

    again = client.post(f"/reports/export?today={TODAY}", json={"report": "pl"}, headers=H).json()
    assert again["path"] == f"{REPORTS}/2026-10-27-profit-and-loss (2).xlsx"  # never overwrites

    everything = client.post(f"/reports/export?today={TODAY}", json={}, headers=H).json()
    assert everything["name"] == "2026-10-27-reports.xlsx"
    assert everything["sheets"] == ["About", "P&L", "Debtors", "Creditors", "VAT summary", "Trend"]

    assert {p for p in fake.files if p.startswith(EVENTS)} == events_before
    new = [p for p in fake.files if not p.startswith(EVENTS) and not p.startswith(REPORTS) and not p.startswith(AUDIT)]
    assert new == []
    audit = b"".join(d for p, d in fake.files.items() if p.startswith(AUDIT))
    assert audit.count(b'"report_exported"') == 3


def test_export_needs_x_user_and_a_known_report(client: TestClient) -> None:
    assert client.post("/reports/export", json={}, headers={"X-Graph-Token": "tok"}).status_code == 400
    assert client.post("/reports/export", json={"report": "balance-sheet"}, headers=H).status_code == 422
    assert client.post("/reports/export", json={"report": "pl", "extra": 1}, headers=H).status_code == 422


async def test_export_of_the_demo_books_is_labelled(fake: FakeStorageProvider) -> None:
    from app.main import app
    from app.routes_books import get_books

    b = Books(fake)
    for ev in ledger(approved_by=rp.SAMPLE_APPROVER):
        await b.post_event(ev)

    async def override() -> Books:
        return b

    app.dependency_overrides[get_books] = override
    try:
        body = TestClient(app).post(f"/reports/export?today={TODAY}", json={"report": "trend"}, headers=H).json()
    finally:
        app.dependency_overrides.clear()
    assert body["sample"] is True
    wb = load_workbook(io.BytesIO(fake.files[body["path"]]))
    assert wb["About"]["A1"].value == SAMPLE_LABEL
