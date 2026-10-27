"""tools/seed_demo.py: the plan (counts, balance, date spread) and a full run against the fake OneDrive.

The Graph run needs a token for a throwaway account: ACCNOMOS_SEED_GRAPH_TOKEN alone only reads the folder;
with ACCNOMOS_SEED_GRAPH_WRITE=1 as well it seeds that account for real.
"""

from __future__ import annotations

import importlib.util
import io
import os
import sys
from datetime import date
from pathlib import Path
from types import ModuleType

import pytest
from accnomos_connectors import FakeStorageProvider
from accnomos_ledger.bank import detect_bank, parse_bank_csv
from openpyxl import load_workbook

from app import invoicing, jobs, mileage
from app.books import INBOX, NEEDS_A_LOOK, Books
from app.deadlines import period_containing
from app.journal import is_balanced

TOOL = Path(__file__).resolve().parents[3] / "tools" / "seed_demo.py"
TODAY = date(2026, 10, 27)


@pytest.fixture(scope="module")
def seed_demo() -> ModuleType:
    spec = importlib.util.spec_from_file_location("seed_demo", TOOL)
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    sys.modules["seed_demo"] = mod  # dataclasses look their module up by name
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="module")
def plan(seed_demo: ModuleType):  # noqa: ANN201
    return seed_demo.build_plan("Bramblewood Joinery Ltd", TODAY)


# --------------------------------------------------------------------------- #
# The plan
# --------------------------------------------------------------------------- #


def test_counts(plan) -> None:  # noqa: ANN001
    assert plan.event_count == 40
    assert (len(plan.posted), len(plan.invoices), bool(plan.claim_journeys())) == (36, 3, True)
    assert len(plan.inbox) == 25
    assert [d.gold for d in plan.inbox if d.expect == "needs_a_look"] == ["057_mixed.json", "010_foreign.json"]
    kinds = {d.extracted["doc_type"] for d in plan.inbox}
    assert "sales_invoice" in kinds and kinds - {"sales_invoice"}
    assert len(plan.jobs) == 2 and {t.job_id for t in plan.timesheets} == {j.id for j in plan.jobs.values()}
    assert {d.job for d in plan.posted if d.job} == {"job1", "job2"}  # materials bought for each job
    assert len(plan.statements) == 2
    overdue = [i for i in plan.invoices if not i.paid_on and i.due < plan.today]
    assert len(overdue) == 1 and sum(i.paid_on is not None for i in plan.invoices) == 1


def test_every_journal_balances(plan) -> None:  # noqa: ANN001
    for d in plan.posted + plan.inbox:
        journal = d.journal()
        assert is_balanced(journal), d.gold
        assert sum(j.debit_pence for j in journal) == round(d.doc().gross_total * 100) or d.doc().currency != "GBP"
    for inv in plan.invoices:
        _, summary, net, vat = invoicing.price_lines(inv.lines)
        assert inv.total_pence == net + vat == sum(r.net_pence + r.vat_pence for r in summary) > 0


def test_dates_spread_over_the_last_two_periods(plan) -> None:  # noqa: ANN001
    (q1_start, q1_end), (q2_start, q2_end) = plan.periods
    assert (q1_start, q1_end, q2_start, q2_end) == (date(2026, 4, 1), date(2026, 6, 30), date(2026, 7, 1), date(2026, 9, 30))
    dated = [(d.date, d.extracted["doc_type"]) for d in plan.posted] + [(i.issue, "sales_invoice") for i in plan.invoices]
    dated.append((plan.claim_day, "purchase_invoice"))
    assert all(q1_start <= day <= q2_end for day, _ in dated)
    for lo, hi in plan.periods:
        mine = [kind for day, kind in dated if lo <= day <= hi]
        assert len(mine) >= 15
        assert "sales_invoice" in mine and len(set(mine)) > 1
        assert len({day.month for day, _ in dated if lo <= day <= hi}) == 3  # every month has entries
    assert all(plan.current[0] <= d.date <= plan.today for d in plan.inbox)
    assert all(q1_start <= t.date <= q2_end for t in plan.timesheets)
    assert all(q1_start <= j.date <= q2_end for j in plan.journeys)


@pytest.mark.parametrize("today", [date(2026, 1, 1), date(2026, 3, 31), date(2026, 12, 31), date(2027, 5, 15)])
@pytest.mark.parametrize("stagger", [1, 2, 3])
def test_other_days_and_staggers(seed_demo: ModuleType, today: date, stagger: int) -> None:
    plan = seed_demo.build_plan("Bramblewood Joinery Ltd", today, stagger)
    assert plan.event_count == 40 and len(plan.inbox) == 25
    assert plan.current == period_containing(today, stagger)
    assert plan.periods[1][1] < plan.current[0]
    assert sum(not i.paid_on and i.due < today for i in plan.invoices) == 1


def test_deterministic(seed_demo: ModuleType, plan) -> None:  # noqa: ANN001
    again = seed_demo.build_plan("Bramblewood Joinery Ltd", TODAY)
    assert seed_demo.describe(again) == seed_demo.describe(plan)
    assert again.statements == plan.statements


def test_everything_is_labelled(seed_demo: ModuleType, plan) -> None:  # noqa: ANN001
    for d in plan.posted + plan.inbox:
        assert d.text.endswith(seed_demo.LABEL)
        assert d.date.isoformat() in d.text or f"{d.date:%d/%m/%Y}" in d.text  # re-dated, not just in the fields
    for name, data in plan.statements:
        assert name.startswith("SAMPLE") and seed_demo.LABEL in data.decode()
        assert detect_bank(data) == "starling"
        assert parse_bank_csv(data, name)


def test_another_business_name(seed_demo: ModuleType) -> None:
    plan = seed_demo.build_plan("Hazelmere Carpentry Ltd", TODAY)
    sales = [d for d in plan.posted if d.extracted["doc_type"] == "sales_invoice"]
    assert len(sales) == 4 and all(d.extracted["supplier_name"] == "Hazelmere Carpentry Ltd" for d in sales)
    assert all("HAZELMERE" in d.text or "Hazelmere" in d.text for d in sales)


def test_refuses_another_apps_folder(seed_demo: ModuleType) -> None:
    assert seed_demo.root_is_accnomos("personal", "Accnomos")
    assert not seed_demo.root_is_accnomos("personal", "Graph Explorer")
    assert seed_demo.root_is_accnomos("work", "Accnomos")


async def test_wipe_asks_first(seed_demo: ModuleType) -> None:
    class Storage(FakeStorageProvider):
        _root_drive_path = "/Apps/Accnomos"
        _folders: dict = {}

        async def _request(self, method: str, url: str) -> None:
            raise AssertionError("deleted without a yes")

    storage = Storage()
    await storage.write("00 Inbox/a.png", b"x")
    assert await seed_demo._wipe(storage, ask=lambda _q: "") is False
    assert await seed_demo._wipe(storage, ask=lambda _q: "n") is False
    assert "00 Inbox/a.png" in storage.files


def test_dry_run_prints_the_plan(seed_demo: ModuleType, capsys: pytest.CaptureFixture[str]) -> None:
    assert seed_demo.main(["--dry-run"]) == 0
    out = capsys.readouterr().out
    assert "40 approved ledger events" in out and "OVERDUE" in out and "Needs a look" in out


# --------------------------------------------------------------------------- #
# A full run on the fake OneDrive
# --------------------------------------------------------------------------- #


async def test_seed_on_fake_storage(seed_demo: ModuleType) -> None:
    plan = seed_demo.build_plan("Bramblewood Joinery Ltd", date.today())  # journeys can't be logged in the future
    fake = FakeStorageProvider()
    report = await seed_demo.seed(fake, plan, say=lambda _msg: None)
    books = Books(fake)

    events = await books.list_events()
    assert report["events"] == len(events) == 40
    assert all(sum(ln.debit for ln in e.lines) == sum(ln.credit for ln in e.lines) for e in events)
    window = plan.window
    assert all(window[0] <= e.date <= window[1] for e in events)
    assert {e.kind for e in events} == {"sale", "purchase"}
    assert all(e.approved_by == seed_demo.USER for e in events)

    inbox = await books.list_files(INBOX)
    assert len(inbox) == 25 and not await books.list_files(NEEDS_A_LOOK)
    proposals = await books.list_proposals()
    assert len(proposals) == 37 and all(p.status == "posted" for p in proposals)
    assert sum(p.job_id is not None for p in proposals) == sum(d.job is not None for d in plan.posted) >= 2

    invs = {i.number: i for i in await invoicing.list_invoices(books)}
    assert [invs[n].status for n in ("INV-0001", "INV-0002", "INV-0003")] == ["paid", "sent", "sent"]
    assert [invs[n].overdue_days(date.today()) > 0 for n in ("INV-0002", "INV-0003")] == [True, False]
    assert report["emails_not_sent"] == 3

    assert len(await books.bank_statements()) == 2
    lines = await books.stored_bank_lines()
    older_end = plan.periods[0][1]
    older = [ln for ln in lines if ln.date <= older_end]
    assert [(ln.description, ln.reference) for ln in older if ln.status != "matched"] == [
        ("Savings Space", "Tax pot"), ("Starling Bank", "Account fee"),  # left for "transfer" / a bank-charge entry
    ]
    assert all(ln.status == "unmatched" for ln in lines if ln.date > older_end)
    assert await books.bank_suggestions()  # left for the judge to confirm

    assert len(await jobs.load_jobs(books)) == 2 and all(j.customer_id for j in await jobs.load_jobs(books))
    assert sum(e.invoiced_in == "INV-0001" for e in await jobs.load_entries(books)) == 12
    assert len(await mileage.load_vehicles(books)) == 1
    assert "README.txt" in fake.files and "Books.xlsx" in fake.files
    wb = load_workbook(io.BytesIO(fake.files["Books.xlsx"]))
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Entries posted"] == 40 and summary["Proposals waiting"] == 0


# --------------------------------------------------------------------------- #
# Microsoft Graph, only with a token for a throwaway account
# --------------------------------------------------------------------------- #

GRAPH_TOKEN = os.getenv("ACCNOMOS_SEED_GRAPH_TOKEN")


@pytest.mark.skipif(not GRAPH_TOKEN, reason="set ACCNOMOS_SEED_GRAPH_TOKEN (throwaway account) to run against Graph")
async def test_graph(seed_demo: ModuleType) -> None:
    from accnomos_connectors.msgraph import GraphStorageProvider

    kind = os.getenv("ACCNOMOS_SEED_GRAPH_ACCOUNT", "personal")
    async with GraphStorageProvider(GRAPH_TOKEN or "", kind) as storage:  # type: ignore[arg-type]
        await storage.ensure_root()
        await storage.list("")
        if os.getenv("ACCNOMOS_SEED_GRAPH_WRITE") != "1":
            return
        books = Books(storage)
        if await books.list_events():
            pytest.skip("the account already has books; wipe it with tools/seed_demo.py --wipe first")
        plan = seed_demo.build_plan("Bramblewood Joinery Ltd", date.today())
        report = await seed_demo.seed(storage, plan, say=print)
        assert report["events"] == 40 and report["inbox"] == 25
