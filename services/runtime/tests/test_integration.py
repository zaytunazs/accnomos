"""Every feature against one books folder: what each one writes, and one Books.xlsx rebuilt from all of them.

Each feature has its own tests; this one checks they live together. Every path written must fit the
layout in docs/ARCHITECTURE.md (the `LAYOUT` patterns below), and the rebuilt workbook must carry
every feature's sheet with its rows.
"""

from __future__ import annotations

import io
import re
import sys
from datetime import UTC, date, datetime, timedelta
from pathlib import Path

import pytest
from accnomos_connectors import FakeMailProvider, FakeStorageProvider, MailAttachment, MailMessage
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import mileage  # noqa: E402
from app.books import Books  # noqa: E402
from app.projection import SHEETS  # noqa: E402

USER = "sam@example.co.uk"
H = {"X-Graph-Token": "tok", "X-User": USER}
GOLD = Path(__file__).resolve().parents[3] / "evals" / "gold" / "bank"

# docs/ARCHITECTURE.md "The books folder". Keep the two in step.
LAYOUT = [
    r"00 Inbox/[^/]+",
    r"01 Receipts/\d{4}-\d{2}/[^/]+",
    r"02 Needs a look/[^/]+",
    r"03 Sales/[^/]+",
    r"04 Bank statements/\d{4}-\d{2}-\d{2}_[a-z]+( \(\d+\))?\.csv",
    r"06 VAT returns/[^/]+/[^/]+",
    r"_System/docs/[0-9a-f]{64}\.json",
    r"_System/proposals/[0-9a-f]+\.json",
    r"_System/audit/\d{4}-\d{2}\.jsonl",
    r"_System/audit/autopilot/\d{4}-\d{2}\.jsonl",
    r"_System/Ledger/events/\d{4}-\d{2}/\d{8}T\d{6}Z-runtime-[0-9a-f]+\.json",
    r"_System/bank/[0-9a-f]{64}\.json",
    r"_System/rules\.json",
    r"_System/Rules/rules\.json",
    r"_System/cursors\.json",
    r"_System/Intake/seen\.json",
    r"_System/Intake/log/\d{4}-\d{2}\.jsonl",
    r"_System/Contacts/contacts\.json",
    r"_System/Jobs/jobs\.json",
    r"_System/Staff/staff\.json",
    r"_System/Timesheets/\d{4}-\d{2}\.jsonl",
    r"_System/Vehicles/vehicles\.json",
    r"_System/Mileage/\d{4}-\d{2}\.jsonl",
    r"Books\.xlsx",
]


@pytest.fixture
def fake() -> FakeStorageProvider:
    return FakeStorageProvider()


@pytest.fixture
def client(fake: FakeStorageProvider):
    from accnomos_connectors import FakeContactsProvider

    from app.main import app
    from app.routes_books import get_books
    from app.routes_contacts import get_address_book
    from app.routes_mail import get_mail

    mail = FakeMailProvider()
    mail.add(
        MailMessage(id="m1", received_at=datetime.now(UTC) - timedelta(days=1), from_address="orders@screwfix.com",
                    from_name="Screwfix", subject="Your Screwfix invoice"),
        (MailAttachment(id="a1", name="Invoice 123.pdf", content_type="application/pdf", size=20_026), b"%PDF-1.7 " + b"x" * 20_017),
    )

    async def books() -> Books:
        return Books(fake)

    async def mailbox() -> FakeMailProvider:
        return mail

    async def address_book() -> FakeContactsProvider:
        return FakeContactsProvider()

    app.dependency_overrides.update({get_books: books, get_mail: mailbox, get_address_book: address_book})
    yield TestClient(app)
    app.dependency_overrides.clear()


def ok(r):  # noqa: ANN001, ANN201
    assert r.status_code == 200, r.text
    return r.json()


def test_every_feature_writes_inside_the_layout_and_books_xlsx_has_every_sheet(client: TestClient, fake: FakeStorageProvider) -> None:
    # Outlook intake
    assert len(ok(client.post("/mail/scan", headers=H, json={}))["saved"]) == 1

    # Bank: import, then turn an unmatched line into an expense and approve it
    with open(GOLD / "barclays.csv", "rb") as f:
        ok(client.post("/books/bank/import", headers=H, files={"file": ("statement.csv", f, "text/csv")}))
    line = next(u["line"] for u in ok(client.get("/books/bank/unmatched", headers=H)) if u["line"]["amount_pence"] < 0)
    expense = ok(client.post(f"/books/bank/lines/{line['id']}/expense-proposal", headers=H))
    posted = ok(client.post(f"/books/proposals/{expense['id']}/approve", headers=H, json={}))
    assert posted["status"] == "posted"

    # Jobs, staff and timesheets
    job = ok(client.post("/jobs", headers=H, json={"name": "Loft boarding"}))
    sam = ok(client.post("/staff", headers=H, json={"name": "Sam", "role": "Carpenter",
                                                     "hourly_rate_pence": 1800, "charge_rate_pence": 4000}))
    entry = ok(client.post("/timesheets", headers=H, json={"job_id": job["id"], "staff_id": sam["id"],
                                                           "date": date.today().isoformat(), "hours": 7.5}))
    ok(client.post(f"/jobs/{job['id']}/mark-invoiced", headers=H, json={"invoice_id": "INV-1", "entry_ids": [entry["id"]]}))

    # Vehicles and mileage: a claim is a proposal; approving it posts
    car = ok(client.post("/vehicles", headers=H, json={"name": "Golf", "kind": "car", "owned_by": "personal"}))
    ty = mileage.tax_year_of(date.today())
    ok(client.post("/mileage", headers=H, json={"vehicle_id": car["id"], "date": date.today().isoformat(),
                                                 "miles": 100, "purpose": "Site visit"}))
    assert ok(client.get(f"/mileage/summary?tax_year={ty}", headers=H))["allowance_pence"] == 4500
    claim = ok(client.post(f"/mileage/propose-allowance?tax_year={ty}&period={mileage.quarter_of(date.today())}",
                           headers=H))["proposal"]
    assert ok(client.post(f"/books/proposals/{claim['id']}/approve", headers=H, json={}))["status"] == "posted"

    # Contacts, then undo the bank expense (a reversal, then a rebuild)
    assert ok(client.post("/contacts/refresh", headers=H))["contacts"]
    ok(client.post(f"/books/undo/{posted['event_id']}", headers=H))
    ok(client.post("/books/rebuild", headers=H))

    stray = [p for p in fake.files if not any(re.fullmatch(pat, p) for pat in LAYOUT)]
    assert stray == [], f"written outside the documented layout: {stray}"

    wb = load_workbook(io.BytesIO(fake.files["Books.xlsx"]))
    assert tuple(wb.sheetnames) == SHEETS
    for name in ("Ledger", "Accounts", "Suppliers", "Bank", "Jobs", "Vehicles", "Proposals"):
        assert wb[name].max_row > 1, f"{name} is empty"
    assert wb["Bank"].max_row - 1 == len(ok(client.get("/books/bank/lines", headers=H)))
