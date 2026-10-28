"""Reports & exports (app/exports/*, app/routes_exports.py) on the hand-built ledger of tests/test_reports.py.

Reconciliation: the trial balance balances; the balance sheet balances; the P&L ties to the balance sheet's profit and
to the movement in retained earnings; the general ledger closes on the trial balance; cash agrees across the cash book,
cash flow and cash summary; the VAT detail adds up to the boxes. Golden files (tests/fixtures/exports/*.csv) hold the
first lines of every vendor import file exactly: headings as the vendor's template, then a few rows. Regenerate them
with UPDATE_GOLDEN=1 only after checking a change against the vendor's template.
"""

from __future__ import annotations

import csv
import io
import json
import os
import zipfile
from datetime import UTC, date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from accnomos_connectors import FakeStorageProvider
from accnomos_ledger import LedgerEvent, Line, VatCode
from accnomos_ledger.bank import BankLine
from app import contacts as cb
from app import reports as rp
from app.books import AUDIT, EVENTS, Books
from app.exports import accounts as ac
from app.exports import reconciliation as rc
from app.exports import registry as reg
from app.exports import statements as st
from app.exports import vat_reports as vr
from app.exports import writers
from app.exports.mapping import MAPPINGS, Mappings, Target, resolve
from app.exports.vendors import SOURCES, TEMPLATES, parse_address, short_ref
from app.paths import REPORTS
from app.vat_cycle import VatCycle


def _load_reports_ledger():
    """The hand-built ledger of tests/test_reports.py (whose docstring works its totals out), loaded by path: the tests
    directory is not a package."""
    import importlib.util

    spec = importlib.util.spec_from_file_location("_reports_ledger", Path(__file__).with_name("test_reports.py"))
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


_T = _load_reports_ledger()
H, Q3, TODAY, USER, _ev, _sale, ledger = _T.H, _T.Q3, _T.TODAY, _T.USER, _T._ev, _T._sale, _T.ledger

GOLDEN = Path(__file__).parent / "fixtures" / "exports"
UPDATE = os.environ.get("UPDATE_GOLDEN") == "1"
JUN = (date(2026, 6, 1), date(2026, 6, 30))
YEAR = (date(2026, 1, 1), date(2026, 12, 31))
SPANS = [Q3, JUN, YEAR, (date(2026, 8, 1), date(2026, 10, 31)), (date(2025, 1, 1), date(2025, 3, 31))]


def books_ledger() -> list[LedgerEvent]:
    """The test_reports ledger plus capital introduced on 1 Jun 2026 and drawings on 20 Sep 2026."""
    return [
        *ledger(),
        _ev(
            "b",
            "adjustment",
            date(2026, 6, 1),
            [Line(account="1200 Bank", debit=200_000), Line(account="3000 Capital", credit=200_000)],
        ),
        _ev(
            "c",
            "bank",
            date(2026, 9, 20),
            [Line(account="3010 Drawings", debit=15_000), Line(account="1200 Bank", credit=15_000)],
        ),
    ]


# --------------------------------------------------------------------------- #
# Reconciliation
# --------------------------------------------------------------------------- #


def _total(section: reg.Section, key: str) -> int:
    (row,) = [r for r in section.rows if r.style == "total"]
    return int(row.cells[key] or 0)


@pytest.mark.parametrize("span", SPANS)
def test_trial_balance_balances(span: tuple[date, date]) -> None:
    tb = ac.trial_balance(books_ledger(), *span)
    assert _total(tb, "debit") == _total(tb, "credit")
    assert sum(ac.period_balances(books_ledger(), *span).values()) == 0


@pytest.mark.parametrize("span", SPANS)
def test_balance_sheet_balances_and_ties_to_the_p_and_l(span: tuple[date, date]) -> None:
    events = books_ledger()
    _, f = ac.balance_sheet(events, *span)
    assert f.net_assets == f.equity
    assert f.profit == rp.profit_and_loss(events, *span).net_profit_pence
    # Retained earnings move by exactly the period's profit.
    before = ac.pl_net(events, None, date.fromordinal(span[0].toordinal() - 1))
    after = ac.pl_net(events, None, span[1])
    assert after - before == f.profit
    assert f.brought_forward == before


def test_q3_balance_sheet_figures() -> None:
    _, f = ac.balance_sheet(books_ledger(), *Q3)
    # Bank: −9,600 (A) +200,000 (capital) +36,000 (D) −5,190 (E) −12,000 (G) +12,000 (G') −15,000 (drawings)
    # Cash −6,000 (F); debtors 120,000 (B).
    assert f.current == 120_000 + (200_000 - 9_600 + 36_000 - 5_190 - 15_000) - 6_000
    # Creditors 30,000 (C), VAT control 20,135 + Q2's −1,600 reclaimed, payroll 53,000.
    assert f.liabilities == 30_000 + (20_135 - 1_600) + 53_000
    assert f.capital == 200_000 - 15_000
    assert f.brought_forward == -8_000 and f.profit == 41_675


@pytest.mark.parametrize("span", SPANS)
def test_general_ledger_closes_on_the_trial_balance(span: tuple[date, date]) -> None:
    events = books_ledger()
    gl = ac.general_ledger(events, *span, {})
    closing = {
        str(r.cells["account"]): int(r.cells["balance"] or 0)
        for r in gl.rows
        if r.cells.get("details") == "Closing balance"
    }
    balances = ac.period_balances(events, *span)
    assert {a: v for a, v in closing.items() if v} == balances
    rng = ac.trial_balance_range(events, *span)
    assert _total(rng, "opening") == 0 and _total(rng, "closing") == 0
    assert _total(rng, "debit") == _total(rng, "credit")


@pytest.mark.parametrize("span", SPANS)
def test_cash_agrees_everywhere(span: tuple[date, date]) -> None:
    events = books_ledger()
    _, flow = st.cash_flow(events, *span)
    book = ac.cash_book(events, *span, {})
    closing_book = sum(int(r.cells["balance"] or 0) for r in book.rows if r.cells.get("details") == "Closing balance")
    summary = st.cash_summary(events, *span)
    closing_summary = int(summary.rows[-1].cells["amount"] or 0)
    assert flow["closing"] == closing_book == closing_summary == st.cash_balance(events, span[1])
    bank = rc.bank_summary(events, *span)
    assert _total(bank, "closing") == flow["closing"]


def test_cash_flow_lines_for_q3() -> None:
    section, _ = st.cash_flow(books_ledger(), *Q3)
    by = {str(r.cells["line"]).strip(): r.cells.get("amount") for r in section.rows}
    assert by["Receipts from customers"] == 36_000  # D, paid at once
    assert by["Payments to suppliers"] == -(5_190 + 6_000)  # E, F; G and its undo net to nil
    assert by["Owner and loan funding (net)"] == -15_000  # drawings


@pytest.mark.parametrize("span", SPANS)
def test_changes_in_equity_close_on_the_balance_sheet(span: tuple[date, date]) -> None:
    events = books_ledger()
    _, closing = st.changes_in_equity(events, *span)
    _, f = ac.balance_sheet(events, *span)
    assert closing == f.equity


def test_p_and_l_comparatives() -> None:
    section, now = ac.pl_comparative(books_ledger(), Q3, reg.previous_span(*Q3), reg.year_earlier(*Q3))
    assert reg.previous_span(*Q3) == (date(2026, 4, 1), date(2026, 6, 30))  # whole months shift by months
    assert reg.year_earlier(date(2024, 2, 29), date(2024, 2, 29)) == (date(2023, 2, 28), date(2023, 2, 28))
    net = section.rows[-1].cells
    assert (net["now"], net["prev"], net["last"], net["change"]) == (41_675, -8_000, 0, 49_675)
    assert now.net_profit_pence == 41_675


def test_executive_summary_figures() -> None:
    section = st.executive_summary(books_ledger(), Q3, reg.previous_span(*Q3))
    by = {str(r.cells["measure"]): r.cells for r in section.rows}
    assert by["Income"]["now"] == "£1,300.00" and by["Net profit"]["now"] == "£416.75"
    assert by["Gross margin"]["now"] == "80.8%"  # 105,000 / 130,000
    assert by["Debtor days (debtors ÷ income × days)"]["now"] == "85 days"  # 120,000 / 130,000 × 92


def test_ledger_exceptions_flag_by_rule() -> None:
    weekend = _sale("d", date(2026, 8, 1), 50_000, "1200 Bank", "Saturday Ltd")  # a Saturday, £600 gross
    late = _ev(
        "e",
        "adjustment",
        date(2026, 3, 1),
        [Line(account="7999 Sundry expenses", debit=1_00), Line(account="Suspense", credit=1_00)],
    )
    rows = ac.ledger_exceptions([*books_ledger(), weekend, late], date(2026, 1, 1), date(2026, 12, 31)).rows
    rules = {(r.cells["rule"], r.cells["entry"] or r.cells["account"]) for r in rows}
    assert ("Weekend date", "dddddddd") in rules and ("Round amount", "dddddddd") in rules  # 60,000 = £600
    assert ("Late posting", "eeeeeeee") in rules  # dated 1 Mar, approved 1 Oct
    assert ("No account code", "eeeeeeee") in rules
    assert ("Unexpected sign", "1210 Cash") in rules  # cash paid out with none paid in


def test_vat_return_is_hmrc_s_boxes_and_wording() -> None:
    section, ret = vr.vat_return_section(books_ledger(), *Q3, 3)
    boxes = {r.cells["box"]: r for r in section.rows if r.cells["box"]}
    assert [b for b in boxes] == [f"Box {n}" for n in range(1, 10)]
    assert boxes["Box 1"].cells["description"] == "VAT due in the period on sales and other outputs"
    assert boxes["Box 4"].cells["description"] == (
        "VAT reclaimed in the period on purchases and other inputs (including acquisitions from the EU)"
    )
    assert boxes["Box 6"].cells["description"] == "Total value of sales and all other outputs excluding any VAT"
    assert [boxes[f"Box {n}"].cells["amount"] for n in range(1, 6)] == [26_000, 0, 26_000, 5_865, 20_135]
    assert [boxes[f"Box {n}"].cells["whole"] for n in range(6, 10)] == [1_300, 353, 0, 0]
    assert ret.period_key == "2026-Q3"


def test_vat_detail_adds_up_to_the_boxes() -> None:
    sales, purchases, left_out = vr.vat_detail(books_ledger(), *Q3, 3, {})
    assert (_total(sales, "vat"), _total(sales, "net")) == (26_000, 130_000)
    assert (_total(purchases, "vat"), _total(purchases, "net")) == (5_865, 35_325)  # box 7 £353, pence dropped
    assert [r.cells["kind"] for r in left_out.rows] == []  # H is an adjustment: not a candidate at all


def test_vat_reconciliation_has_no_difference() -> None:
    spans = [(date(2026, 4, 1), date(2026, 6, 30)), Q3]
    section = vr.vat_reconciliation(books_ledger(), spans, 3)
    assert [r.cells["difference"] for r in section.rows] == [0, 0]
    assert [r.cells["balance"] for r in section.rows] == [-1_600, 18_535]


def test_digital_links_trace_every_entry_on_the_return() -> None:
    events = books_ledger()
    section = vr.digital_links(events, *Q3, 3, {e.id: f"{EVENTS}/2026-07/{e.id}.json" for e in events}, {}, {})
    heads = [r for r in section.rows if r.style == "heading"]
    assert len(heads) == 7  # B C D E F G G'
    vat_lines = [r for r in section.rows if r.cells.get("account") == "2200 VAT control"]
    assert {r.cells["box"] for r in vat_lines} == {"1", "4"}
    assert all(str(r.cells["file"]).startswith(EVENTS) for r in section.rows)


def test_bank_reconciliation_and_statement_lines() -> None:
    events = books_ledger()

    def line(
        i: str,
        day: date,
        pence: int,
        desc: str,
        status: str = "matched",
        event: str | None = None,
        bal: int | None = None,
    ) -> BankLine:
        return BankLine(
            id=i,
            date=day,
            amount_pence=pence,
            description=desc,
            source_file="s.csv",
            status=status,  # type: ignore[arg-type]
            matched_event_id=event,
            balance_pence=bal,
        )

    lines = [
        line("1", date(2026, 8, 15), 36_000, "ACME", event="4" * 32),
        line("2", date(2026, 8, 20), -5_190, "NOVEX", event="5" * 32),
        line("3", date(2026, 8, 21), -1_250, "COFFEE", status="unmatched"),
        line("4", date(2026, 8, 21), -1_250, "Coffee.", status="unmatched", bal=999_999),
    ]
    section, diff = rc.bank_reconciliation(events, lines, date(2026, 8, 31))
    rows = {str(r.cells["line"]).strip(): r.cells.get("amount") for r in section.rows}
    ledger_bank = 200_000 - 9_600 + 36_000 - 5_190
    assert rows["Balance in the books (1200 Bank) at 31 Aug 2026"] == ledger_bank
    assert rows["Balance the statement should show"] == ledger_bank - 2_500
    assert diff == 999_999 - (ledger_bank - 2_500)
    assert len(rc.uncoded_lines(lines, *Q3).rows) == 3  # two lines and the total
    dupes = rc.duplicate_lines(lines, *Q3).rows
    assert [r.style for r in dupes] == ["heading", "normal", "normal"]  # COFFEE and "Coffee." are alike


# --------------------------------------------------------------------------- #
# Writers
# --------------------------------------------------------------------------- #


def _q3_doc(section: reg.Section, **kw: object) -> reg.ReportDoc:
    return reg.doc_from("t", "Test", "1 Jul 2026 to 30 Sep 2026", [section], **kw)  # type: ignore[arg-type]


def test_report_csv_xlsx_pdf_agree() -> None:
    doc = _q3_doc(ac.trial_balance(books_ledger(), *Q3))
    rows = list(csv.reader(io.StringIO(writers.to_csv(doc).decode("utf-8-sig"))))
    assert rows[:2] == [["Test"], ["1 Jul 2026 to 30 Sep 2026"]]
    header = rows.index(["Code", "Account", "Debit", "Credit"])
    sales = next(r for r in rows[header:] if r[:2] == ["4000", "Sales"])
    assert sales == ["4000", "Sales", "", "1300.00"]
    wb = load_workbook(io.BytesIO(writers.to_xlsx(doc)))
    ws = wb["Trial balance"]
    values = {r[1]: r[3] for r in ws.iter_rows(min_row=2, values_only=True) if r[1]}
    assert values["Sales"] == 1300 and ws["C2"].number_format.startswith('"£"')
    assert writers.to_pdf(doc).startswith(b"%PDF")


def test_no_text_is_ever_a_formula() -> None:
    evil = _sale("f", date(2026, 7, 3), 1_000, "1100 Debtors", '=HYPERLINK("http://x","y")')
    events = [*books_ledger(), evil]
    doc = _q3_doc(ac.journal_listing(events, *Q3))
    text = writers.to_csv(doc).decode("utf-8-sig")
    assert "=HYPERLINK" in text and "\n=HYPERLINK" not in text and ",=HYPERLINK" not in text
    for row in csv.reader(io.StringIO(text)):
        assert not any(cell.startswith(("=", "+", "@")) for cell in row), row
    cells = [c.value for row in load_workbook(io.BytesIO(writers.to_xlsx(doc))).worksheets[1].iter_rows() for c in row]
    assert not any(isinstance(v, str) and v.startswith("=") for v in cells)


def test_sample_books_are_labelled() -> None:
    doc = _q3_doc(ac.trial_balance(books_ledger(), *Q3), sample=True)
    assert writers.to_csv(doc).decode("utf-8-sig").startswith("SAMPLE — FICTIONAL COMPANY")


# --------------------------------------------------------------------------- #
# Registry
# --------------------------------------------------------------------------- #


def test_registry_has_every_report_in_its_group() -> None:
    specs = reg.catalogue()
    groups = {g for g, _ in reg.GROUPS}
    assert all(s.group in groups for s in specs.values())
    by_group = {g: {s.id for s in specs.values() if s.group == g} for g in groups}
    assert {
        "balance-sheet",
        "profit-and-loss",
        "cash-flow",
        "changes-in-equity",
        "executive-summary",
        "cash-summary",
        "management-report",
    } <= by_group["financial"]
    assert {
        "aged-receivables-detail",
        "aged-payables-summary",
        "receivable-invoice-detail",
        "payable-invoice-summary",
        "contact-transactions",
        "income-expenses-by-contact",
        "customer-invoice-activity",
        "supplier-bill-activity",
        "supplier-days-to-pay",
    } <= by_group["payables"]
    assert {
        "account-summary",
        "bank-reconciliation",
        "bank-summary",
        "uncoded-statement-lines",
        "duplicate-statement-lines",
    } <= by_group["reconciliations"]
    assert {
        "general-ledger",
        "general-ledger-summary",
        "general-ledger-exceptions",
        "journal-report",
        "trial-balance",
        "trial-balance-range",
        "vat-return",
        "vat-reconciliation",
        "vat-detail",
        "vat-digital-links",
    } <= by_group["taxes"]
    assert by_group["transactions"] >= {"account-transactions"}
    assert {"accountant-pack", "ledger-json"} <= by_group["pack"]
    assert "ec-sales-list" not in specs and not set(reg.RESERVED) & set(specs)
    for s in specs.values():
        if s.group == "export" and s.raw is None:
            assert s.id in SOURCES and "2026" in SOURCES[s.id]


def test_reserved_and_duplicate_ids_are_refused() -> None:
    async def build(ctx: reg.ExportContext) -> reg.ReportDoc:
        raise AssertionError

    with pytest.raises(ValueError, match="kept for accounting-tools"):
        reg.register(reg.ExportSpec("fixed-asset-schedule", "financial", "x", "x", build))
    with pytest.raises(ValueError, match="registered twice"):
        reg.register(reg.ExportSpec("trial-balance", "taxes", "x", "x", build))


# --------------------------------------------------------------------------- #
# Vendor files: golden headings and rows
# --------------------------------------------------------------------------- #


@pytest.fixture
def fake() -> FakeStorageProvider:
    return FakeStorageProvider()


@pytest.fixture
async def books(fake: FakeStorageProvider) -> Books:
    b = Books(fake)
    for ev in books_ledger():
        await b.post_event(ev)
    contacts, _ = await cb.refresh(b)
    acme = next(c for c in contacts if c.name == "Acme Kitchens Ltd")
    howdens = next(c for c in contacts if c.name == "Howdens Joinery")
    edits = {
        acme.id: {
            "email": "accounts@acme.example",
            "phone": "0117 496 0000",
            "address": "Unit 4, Harbour Road\nBristol BS1 4RW",
            "vrn": "GB123456789",
            "terms_days": 14,
        },
        howdens.id: {"address": "12 Mill Lane, Leeds, LS1 1AA", "company_number": "01234567"},
    }
    await cb.save_contacts(b, [c.model_copy(update=edits.get(c.id, {})) for c in contacts])
    return b


def _ctx(b: Books, span: tuple[date, date] = Q3, mode: str | None = None) -> reg.ExportContext:
    return reg.ExportContext(
        books=b, today=TODAY, start=span[0], end=span[1], as_of=span[1], cycle=VatCycle(), mode=mode
    )


VENDOR_IDS = [t.id for t, _ in TEMPLATES]


@pytest.mark.parametrize("export_id", VENDOR_IDS)
async def test_vendor_file_matches_its_golden(books: Books, export_id: str) -> None:
    spec = reg.catalogue()[export_id]
    mode = spec.modes[0][0] if spec.modes else None
    doc = await spec.build(_ctx(books, YEAR if "invoices" in export_id or "bills" in export_id else Q3, mode))
    assert doc.import_file is not None
    data = writers.to_csv(doc).decode("utf-8")
    head = "".join(data.splitlines(keepends=True)[:6])
    path = GOLDEN / f"{export_id}.csv"
    if UPDATE:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(head.encode())
    assert head == path.read_bytes().decode(), f"{export_id} differs from {path}"
    (section,) = doc.sections
    t = next(t for t, _ in TEMPLATES if t.id == export_id)
    assert [c.label for c in section.columns] == t.headings
    if t.header:
        assert data.splitlines()[0] == ",".join(t.headings)


def test_golden_files_have_no_stale_ones() -> None:
    assert {p.stem for p in GOLDEN.glob("*.csv")} == set(VENDOR_IDS)


async def test_xero_journals_balance_and_carry_no_vat(books: Books) -> None:
    doc = await reg.catalogue()["xero-manual-journals"].build(_ctx(books, mode="entries"))
    rows = doc.sections[0].rows
    by_narration: dict[str, int] = {}
    for r in rows:
        by_narration[str(r.cells["c0"])] = by_narration.get(str(r.cells["c0"]), 0) + round(
            float(str(r.cells["c5"])) * 100
        )
        assert r.cells["c4"] == "No VAT"
    assert all(v == 0 for v in by_narration.values())  # every journal nets to nil: debits positive, credits negative
    balances = await reg.catalogue()["xero-manual-journals"].build(_ctx(books, mode="balances"))
    assert len({r.cells["c0"] for r in balances.sections[0].rows}) == 1
    assert sum(round(float(str(r.cells["c5"])) * 100) for r in balances.sections[0].rows) == 0


async def test_sage50_journals_are_jd_jc_t9(books: Books) -> None:
    doc = await reg.catalogue()["sage50-audit-trail"].build(_ctx(books, mode="entries"))
    rows = doc.sections[0].rows
    dr = sum(round(float(str(r.cells["c7"])) * 100) for r in rows if r.cells["c0"] == "JD")
    cr = sum(round(float(str(r.cells["c7"])) * 100) for r in rows if r.cells["c0"] == "JC")
    assert dr == cr and {r.cells["c8"] for r in rows} == {"T9"}
    assert {r.cells["c2"] for r in rows if r.cells["c2"] == "1230"}  # 1210 Cash goes to Sage's petty cash


async def test_invoice_lines_add_up_to_the_document(books: Books) -> None:
    doc = await reg.catalogue()["xero-sales-invoices"].build(_ctx(books, YEAR, "period"))
    nets = [round(float(str(r.cells["c18"])) * 100) for r in doc.sections[0].rows]
    vats = [round(float(str(r.cells["c22"])) * 100) for r in doc.sections[0].rows]
    assert sorted(zip(nets, vats, strict=True)) == [(50_000, 10_000), (100_000, 20_000)]  # Acme's two invoices
    open_doc = await reg.catalogue()["xero-sales-invoices"].build(_ctx(books, Q3, "open"))
    assert len(open_doc.sections[0].rows) == 2  # both still unpaid, whatever the period


async def test_freeagent_bank_has_no_header_or_quotes(books: Books) -> None:
    doc = await reg.catalogue()["freeagent-bank-statement"].build(_ctx(books))
    text = writers.to_csv(doc).decode()
    assert '"' not in text and not text.startswith("date")
    first = text.splitlines()[0].split(",")
    assert len(first) == 3 and first[0] == "15/08/2026" and first[1] == "360.00"


def test_address_parsing_and_short_refs() -> None:
    a = parse_address("Unit 4, Harbour Road\nBristol BS1 4RW")
    assert (a.lines, a.town, a.postcode) == (["Unit 4", "Harbour Road"], "Bristol", "BS1 4RW")
    taken: set[str] = set()
    assert [short_ref("Acme Kitchens Ltd", taken), short_ref("Acme Kitchens Ltd", taken)] == ["ACMEKI01", "ACMEKI02"]


def test_mapping_defaults_and_overrides() -> None:
    m = Mappings()
    assert (resolve("xero", "7500 Office supplies", m).code, resolve("xero", "7500 Office supplies", m).quality) == (
        "461",
        "default",
    )
    assert resolve("xero", "7650 Cleaning", m).quality == "suggested"
    assert resolve("sage50", "1210 Cash", m).code == "1230"
    assert resolve("qbo", "7500 Office supplies", m).quality == "own"
    assert resolve("xero", "Suspense", m).quality == "unmapped"
    m.accounts["xero"] = {"7500 Office supplies": Target(code="453", name="Office Expenses")}
    assert resolve("xero", "7500 Office supplies", m).model_dump(include={"code", "quality"}) == {
        "code": "453",
        "quality": "user",
    }


# --------------------------------------------------------------------------- #
# Routes
# --------------------------------------------------------------------------- #


@pytest.fixture
def client(books: Books):
    from app.main import app
    from app.routes_books import get_books

    async def override() -> Books:
        return books

    app.dependency_overrides[get_books] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


def test_catalogue_lists_groups_and_reports(client: TestClient) -> None:
    r = client.get("/exports", headers=H)
    assert r.status_code == 200, r.text
    body = r.json()
    assert [g["id"] for g in body["groups"]][:5] == [
        "financial",
        "payables",
        "reconciliations",
        "taxes",
        "transactions",
    ]
    ids = {x["id"] for x in body["reports"]}
    assert {"trial-balance", "vat-return", "xero-manual-journals", "accountant-pack"} <= ids
    tb = next(x for x in body["reports"] if x["id"] == "trial-balance")
    assert tb["formats"] == ["csv", "xlsx", "pdf"] and tb["params"] == ["period"]
    assert body["favourites"] == [] and body["packages"]["qbo"] == "QuickBooks Online"


def test_preview_and_download_write_nothing(client: TestClient, fake: FakeStorageProvider) -> None:
    before = dict(fake.files)
    r = client.get(f"/exports/trial-balance?from=2026-07-01&to=2026-09-30&today={TODAY}", headers=H)
    assert r.status_code == 200, r.text
    body = r.json()
    assert (
        body["doc"]["title"] == "Trial balance"
        and body["filenames"]["csv"] == "trial-balance-2026-07-01-to-2026-09-30.csv"
    )
    default = client.get(f"/exports/vat-return?today={TODAY}", headers=H).json()
    assert (default["from"], default["to"]) == ("2026-10-01", "2026-12-31")  # the VAT period we are in
    for fmt, magic in (("csv", "﻿Test".encode()[:3]), ("xlsx", b"PK"), ("pdf", b"%PDF")):
        f = client.get(f"/exports/balance-sheet/file?format={fmt}&from=2026-07-01&to=2026-09-30", headers=H)
        assert f.status_code == 200 and f.content.startswith(magic), fmt
        assert (
            f.headers["content-disposition"] == f'attachment; filename="balance-sheet-2026-07-01-to-2026-09-30.{fmt}"'
        )
    assert fake.files == before


def test_save_writes_07_reports_and_an_audit_line(client: TestClient, fake: FakeStorageProvider) -> None:
    events = {p for p in fake.files if p.startswith(EVENTS)}
    body = {"format": "pdf", "from": "2026-07-01", "to": "2026-09-30"}
    r = client.post("/exports/vat-return/save", json=body, headers=H)
    assert r.status_code == 200, r.text
    assert r.json()["path"] == f"{REPORTS}/vat-return-2026-07-01-to-2026-09-30.pdf"
    again = client.post("/exports/vat-return/save", json=body, headers=H).json()
    assert again["path"] == f"{REPORTS}/vat-return-2026-07-01-to-2026-09-30 (2).pdf"
    assert {p for p in fake.files if p.startswith(EVENTS)} == events
    audit = b"".join(d for p, d in fake.files.items() if p.startswith(AUDIT))
    assert audit.count(b'"report_saved"') == 2
    assert client.post("/exports/vat-return/save", json=body, headers={"X-Graph-Token": "tok"}).status_code == 400


def test_bad_requests_are_refused(client: TestClient) -> None:
    assert client.get("/exports/nope", headers=H).status_code == 404
    assert client.get("/exports/trial-balance/file?format=zip", headers=H).status_code == 422
    assert client.get("/exports/trial-balance?from=2026-07-01", headers=H).status_code == 422
    assert client.get("/exports/trial-balance?from=2026-07-01&to=2026-06-01", headers=H).status_code == 422
    assert client.get("/exports/xero-manual-journals?mode=everything", headers=H).status_code == 422
    assert client.post("/exports/trial-balance/save", json={"format": "csv", "extra": 1}, headers=H).status_code == 422


def test_favourites_round_trip(client: TestClient, fake: FakeStorageProvider) -> None:
    r = client.put("/exports/favourites", json={"ids": ["vat-return", "trial-balance", "vat-return"]}, headers=H)
    assert r.json() == {"ids": ["vat-return", "trial-balance"]}
    assert json.loads(fake.files["_System/exports/favourites.json"])["ids"] == ["vat-return", "trial-balance"]
    assert client.get("/exports", headers=H).json()["favourites"] == ["vat-return", "trial-balance"]
    assert client.put("/exports/favourites", json={"ids": ["nope"]}, headers=H).status_code == 422


def test_mappings_change_the_import_file(client: TestClient, fake: FakeStorageProvider) -> None:
    view = client.get("/exports/mappings/xero", headers=H).json()
    office = next(a for a in view["accounts"] if a["account"] == "7500 Office supplies")
    assert (office["code"], office["quality"]) == ("461", "default")
    r = client.put(
        "/exports/mappings/xero",
        json={"accounts": {"7500 Office supplies": {"code": "453", "name": "Office Expenses"}}},
        headers=H,
    )
    assert r.status_code == 200, r.text
    assert MAPPINGS in fake.files
    rows = client.get("/exports/xero-manual-journals?from=2026-08-01&to=2026-08-31", headers=H).json()["doc"][
        "sections"
    ][0]["rows"]
    assert "453" in {row["cells"]["c3"] for row in rows}
    client.put("/exports/mappings/xero", json={"accounts": {"7500 Office supplies": None}}, headers=H)
    office = next(
        a
        for a in client.get("/exports/mappings/xero", headers=H).json()["accounts"]
        if a["account"] == "7500 Office supplies"
    )
    assert office["quality"] == "default"
    assert client.get("/exports/mappings/myob", headers=H).status_code == 404


def test_accountant_pack_and_ledger_json(client: TestClient) -> None:
    r = client.get("/exports/accountant-pack/file?format=zip&from=2026-01-01&to=2026-12-31&today=2026-10-27", headers=H)
    assert r.status_code == 200, r.text
    names = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    for want in (
        "README.txt",
        "trial-balance.csv",
        "general-ledger.pdf",
        "balance-sheet.csv",
        "profit-and-loss.csv",
        "aged-receivables-detail.csv",
        "aged-payables-detail.csv",
        "contact-list.csv",
        "ledger.json",
        "vat/2026-Q3/vat-return.pdf",
        "vat/2026-Q3/vat-detail.csv",
        "vat/2026-Q2/vat-digital-links.csv",
    ):
        assert want in names, want
    assert "vat/2026-Q4/vat-return.csv" not in names  # Q4 ends inside the period but after today
    readme = zipfile.ZipFile(io.BytesIO(r.content)).read("README.txt").decode()
    assert "Period: 1 Jan 2026 to 31 Dec 2026." in readme and "As approved" in readme
    data = client.get("/exports/ledger-json/file?format=json", headers=H).json()
    assert (
        data["format"] == "accnomos-ledger" and data["version"] == 1 and set(data["schema"]) >= {"events", "accounts"}
    )
    assert data["totals"]["debits"] == data["totals"]["credits"] and data["totals"]["events"] == len(books_ledger())
    assert all(sum(ln["debit"] for ln in e["lines"]) == sum(ln["credit"] for ln in e["lines"]) for e in data["events"])
    xero = client.get("/exports/xero-all/file?format=zip&from=2026-07-01&to=2026-09-30", headers=H)
    assert "xero-manual-journals.csv" in zipfile.ZipFile(io.BytesIO(xero.content)).namelist()


def test_every_report_previews_and_renders(client: TestClient) -> None:
    from app import security

    for spec in reg.catalogue().values():
        if security.LIMITER:  # some 150 calls in one test: more than one client's burst
            security.LIMITER.reset()
        r = client.get(f"/exports/{spec.id}?from=2026-01-01&to=2026-12-31&today=2026-10-27", headers=H)
        assert r.status_code == 200, (spec.id, r.text)
        for fmt in spec.formats:
            f = client.get(
                f"/exports/{spec.id}/file?format={fmt}&from=2026-07-01&to=2026-09-30&today=2026-10-27", headers=H
            )
            assert f.status_code == 200 and f.content, (spec.id, fmt, f.status_code, f.text[:300])


def test_generated_at_is_utc() -> None:
    assert _q3_doc(reg.Section(title="x", columns=[])).generated_at.tzinfo == UTC
    assert USER  # the test_reports helpers approve as this user
    assert VatCode.STANDARD.value == "S"
