"""Jobs, staff and timesheets: quarter-hour rounding, pence maths, invoice-line grouping, job summaries
(with and without tagged purchases) and the /jobs, /staff, /timesheets routes, over FakeStorageProvider."""

from __future__ import annotations

import json
import sys
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from accnomos_connectors import FakeStorageProvider
from accnomos_ledger.events import LedgerEvent, Line

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import jobs  # noqa: E402
from app.approve import ApprovalError, approve, undo  # noqa: E402
from app.books import Books  # noqa: E402
from app.jobs import Job, JobsConflict, JobsError, StaffMember, TimesheetEntry  # noqa: E402
from app.schemas import ExtractedDoc, ExtractedLine, Proposal  # noqa: E402

USER = "sam@example.co.uk"
H = {"X-Graph-Token": "tok", "X-User": USER}
JOB = Job(id="00000000a001", name="Kitchen refit, 12 Elm Road", customer_id="c1")
SAM = StaffMember(id="00000000b001", name="Sam", role="Electrician", hourly_rate_pence=2000, charge_rate_pence=4500)
ALEX = StaffMember(id="00000000b002", name="Alex", role="Apprentice", hourly_rate_pence=1133, charge_rate_pence=None)


def entry(staff: StaffMember, hours: float, day: int = 3, **kw: object) -> TimesheetEntry:
    return TimesheetEntry(staff_id=staff.id, job_id=JOB.id, date=date(2026, 10, day), hours=hours, **kw)


# --------------------------------------------------------------------------- #
# Hours and pence
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    ("given", "want"),
    [
        (1, "1"),
        (1.25, "1.25"),
        ("7.5", "7.5"),
        (0.125, "0.25"),
        (0.37, "0.25"),
        (0.375, "0.5"),
        (2.1, "2"),
        (23.9, "24"),
    ],
)
def test_hours_round_to_the_nearest_quarter_half_up(given: object, want: str) -> None:
    assert jobs.round_hours(given) == Decimal(want)  # type: ignore[arg-type]


@pytest.mark.parametrize("bad", [0, 0.1, -1, 24.2, "lots", float("nan"), float("inf")])
def test_hours_out_of_range_are_refused(bad: object) -> None:
    with pytest.raises(JobsError):
        jobs.round_hours(bad)  # type: ignore[arg-type]


def test_entries_store_rounded_hours() -> None:
    assert entry(SAM, 1.3).hours == 1.25
    with pytest.raises(ValueError, match="at least 0.25"):
        entry(SAM, 0.05)


def test_pence_are_whole_and_round_half_up() -> None:
    assert jobs.pence_for(4, 2000) == 2000  # 1h at £20
    assert jobs.pence_for(1, 1133) == 283  # 0.25h × £11.33 = 283.25p
    assert jobs.pence_for(1, 1134) == 284  # 283.5p rounds up
    assert jobs.pence_for(3, 1001) == 751  # 750.75p
    assert jobs.pence_for(30, 4500) == 33750  # 7.5h at £45


def test_costs_are_per_person_totals_not_summed_rounded_entries() -> None:
    # Four 0.25h entries at £11.33: per-entry rounding would give 4 × 283 = 1132p; the hour costs 1133p.
    s = jobs.summarise(JOB, [entry(ALEX, 0.25, day=d) for d in (1, 2, 3, 4)], [ALEX])
    assert s.labour_cost_pence == 1133
    assert s.hours == 1


# --------------------------------------------------------------------------- #
# Invoice lines
# --------------------------------------------------------------------------- #


def test_invoice_lines_group_unbilled_billable_hours_by_staff() -> None:
    entries = [
        entry(SAM, 7.5, day=1),
        entry(SAM, 2.25, day=6),
        entry(ALEX, 8, day=1),
        entry(SAM, 1, day=2, billable=False),  # not billable: left out
        entry(SAM, 3, day=2, invoiced_in="INV-0001"),  # already invoiced: left out
        TimesheetEntry(staff_id=SAM.id, job_id="00000000a999", date=date(2026, 10, 1), hours=5),  # other job
    ]
    prop = jobs.invoice_lines(JOB, entries, [SAM, ALEX])
    assert [(ln.staff_id, ln.qty, ln.unit_price_pence, ln.net_pence) for ln in prop.lines] == [
        (ALEX.id, 8, 1133, 9064),  # no charge rate: charged at cost, with a warning
        (SAM.id, 9.75, 4500, 43875),
    ]
    assert prop.lines[1].description == "Kitchen refit, 12 Elm Road: Electrician (Sam), labour 01/10/2026–06/10/2026"
    assert prop.lines[1].entry_ids == [entries[0].id, entries[1].id]
    assert prop.net_pence == 9064 + 43875
    assert set(prop.entry_ids) == {entries[0].id, entries[1].id, entries[2].id}
    assert prop.customer_id == "c1"
    assert any("Alex has no charge rate" in w for w in prop.warnings)
    assert all(ln.vat_code == "S" and ln.unit == "hour" for ln in prop.lines)


def test_invoice_lines_with_nothing_to_bill() -> None:
    prop = jobs.invoice_lines(JOB, [entry(SAM, 2, billable=False)], [SAM])
    assert prop.lines == [] and prop.net_pence == 0
    assert prop.warnings == ["No unbilled billable hours on this job"]


# --------------------------------------------------------------------------- #
# Job summary
# --------------------------------------------------------------------------- #


def test_summary_without_purchases() -> None:
    entries = [entry(SAM, 8), entry(SAM, 2, billable=False), entry(SAM, 4, invoiced_in="INV-0002"), entry(ALEX, 6)]
    s = jobs.summarise(JOB, entries, [SAM, ALEX], customer_name="Mrs Patel")
    assert (s.hours, s.billable_hours, s.unbilled_hours) == (20, 18, 14)
    assert s.labour_cost_pence == 14 * 2000 + 6 * 1133  # 34798
    assert s.billable_value_pence == 12 * 4500 + 6 * 1133  # 60798
    assert s.invoiced_value_pence == 4 * 4500
    assert s.unbilled_value_pence == s.billable_value_pence - 18000
    assert s.materials_pence == 0 and s.purchases == []
    assert s.margin_pence == 60798 - 34798
    assert s.customer_name == "Mrs Patel"
    assert [p.name for p in s.by_staff] == ["Alex", "Sam"]


def _purchase(
    eid: str, net: int, vat: int, proposal_id: str, kind: str = "purchase", reverses: str | None = None
) -> LedgerEvent:
    lines = [
        Line(account="5000 Materials", debit=net),
        Line(account="2200 VAT control", debit=vat),
        Line(account="1200 Bank", credit=net + vat),
    ]
    return LedgerEvent(
        id=eid,
        kind=kind,
        date=date(2026, 10, 5),
        lines=lines,
        approved_by=USER,  # type: ignore[arg-type]
        approved_at=datetime(2026, 10, 5, tzinfo=UTC),
        proposal_id=proposal_id,
        net=net,
        vat=vat,
        counterparty="Screwfix",
        reverses=reverses,
    )


def _proposal(pid: str, event_id: str | None, job_id: str | None, status: str = "posted") -> Proposal:
    return Proposal.model_validate(
        {
            "id": pid,
            "doc_sha256": pid,
            "extracted": _doc().model_dump(),
            "checks": [],
            "status": status,
            "reasons": [],
            "suggested_filename": "x.jpg",
            "journal_draft": [],
            "model": "stub",
            "elapsed_s": 0,
            "event_id": event_id,
            "job_id": job_id,
        }
    )


def test_summary_with_purchases_counts_live_tagged_net_only() -> None:
    events = [
        _purchase("e1", 10000, 2000, "p1"),  # tagged: counts at net £100 (VAT is reclaimed, not a cost)
        _purchase("e2", 5000, 1000, "p2"),  # tagged but undone: does not count
        _purchase("r2", 5000, 1000, "p2", kind="reversal", reverses="e2"),
        _purchase("e3", 7000, 1400, "p3"),  # not tagged
        _purchase("e4", 2500, 500, "p4"),  # another job
    ]
    proposals = [
        _proposal("p1", "e1", JOB.id),
        _proposal("p2", "e2", JOB.id),
        _proposal("p3", "e3", None),
        _proposal("p4", "e4", "00000000a999"),
    ]
    bought = jobs.job_purchases(JOB.id, proposals, events)
    assert [(p["event_id"], p["net_pence"], p["supplier"]) for p in bought] == [("e1", 10000, "Screwfix")]
    s = jobs.summarise(JOB, [entry(SAM, 8)], [SAM], bought)
    assert s.materials_pence == 10000
    assert s.margin_pence == 8 * 4500 - 8 * 2000 - 10000


# --------------------------------------------------------------------------- #
# Storage
# --------------------------------------------------------------------------- #


@pytest.fixture
def fake() -> FakeStorageProvider:
    return FakeStorageProvider()


@pytest.fixture
async def books(fake: FakeStorageProvider) -> Books:
    b = Books(fake)
    await jobs.save_job(b, JOB)
    await jobs.save_staff(b, SAM)
    await jobs.save_staff(b, ALEX)
    return b


def _sheet(fake: FakeStorageProvider, month: str) -> list[dict]:
    return [json.loads(x) for x in fake.files.get(f"{jobs.TIMESHEETS}/{month}.jsonl", b"").decode().splitlines()]


async def test_entries_are_filed_by_month_and_move_when_redated(fake: FakeStorageProvider, books: Books) -> None:
    e = await jobs.add_entry(books, entry(SAM, 3))
    assert [r["id"] for r in _sheet(fake, "2026-10")] == [e.id]
    moved = await jobs.update_entry(books, e.id, {"date": "2026-11-02", "hours": 2.6})
    assert moved.hours == 2.5
    assert _sheet(fake, "2026-10") == [] and _sheet(fake, "2026-11")[0]["hours"] == 2.5
    assert [x.id for x in await jobs.load_entries(books)] == [e.id]


async def test_invoiced_entries_are_frozen(books: Books) -> None:
    a = await jobs.add_entry(books, entry(SAM, 3))
    b = await jobs.add_entry(books, entry(SAM, 1, billable=False))
    with pytest.raises(JobsError, match="not billable"):
        await jobs.mark_invoiced(books, [a.id, b.id], "INV-0003")
    assert (await jobs.get_entry(books, a.id)).invoiced_in is None  # all or nothing
    await jobs.mark_invoiced(books, [a.id], "INV-0003")
    await jobs.mark_invoiced(books, [a.id], "INV-0003")  # a retried send is harmless
    with pytest.raises(JobsConflict):
        await jobs.mark_invoiced(books, [a.id], "INV-0004")
    with pytest.raises(JobsConflict):
        await jobs.update_entry(books, a.id, {"hours": 5})
    with pytest.raises(JobsConflict):
        await jobs.delete_entry(books, a.id)
    with pytest.raises(JobsError):
        await jobs.update_entry(books, b.id, {"invoiced_in": "INV-0005"})


async def test_closed_jobs_and_unknown_staff_take_no_hours(books: Books) -> None:
    with pytest.raises(JobsError, match="No such staff"):
        await jobs.add_entry(
            books, TimesheetEntry(staff_id="00000000bfff", job_id=JOB.id, date=date(2026, 10, 1), hours=1)
        )
    await jobs.save_job(books, JOB.model_copy(update={"status": "closed"}))
    with pytest.raises(JobsConflict, match="closed"):
        await jobs.add_entry(books, entry(SAM, 1))


async def test_jobs_and_staff_with_hours_cannot_be_deleted(books: Books) -> None:
    await jobs.add_entry(books, entry(SAM, 1))
    with pytest.raises(JobsConflict):
        await jobs.delete_job(books, JOB.id)
    with pytest.raises(JobsConflict):
        await jobs.delete_staff(books, SAM.id)
    await jobs.delete_staff(books, ALEX.id)
    assert [s.id for s in await jobs.load_staff(books)] == [SAM.id]


# --------------------------------------------------------------------------- #
# Tagging a receipt to a job on approval
# --------------------------------------------------------------------------- #


def _doc(**kw: object) -> ExtractedDoc:
    base: dict[str, object] = dict(
        doc_type="receipt",
        supplier_name="Screwfix",
        supplier_vat_number="GB 232 4567 89",
        customer_name=None,
        invoice_number=None,
        date="2026-10-21",
        currency="GBP",
        lines=[ExtractedLine(description="Cable", qty=1, unit_price=120.0, net=100.0, vat_rate=20, vat=20.0)],
        net_total=100.0,
        vat_total=20.0,
        gross_total=120.0,
        payment_method="card",
        card_last4=None,
        category_guess="Materials",
        confidence=0.95,
    )
    base.update(kw)
    return ExtractedDoc.model_validate(base)


async def _open(books: Books, pid: str = "abcdef0123456789") -> Proposal:
    p = _proposal(pid, None, None, status="proposed")
    await books.save_proposal(p)
    return p


async def test_approve_tags_a_receipt_and_the_summary_counts_it_until_undone(books: Books) -> None:
    p = await _open(books)
    with pytest.raises(ApprovalError, match="No such job"):
        await approve(books, p.id, USER, {"job_id": "00000000afff"})
    posted = await approve(books, p.id, USER, {"job_id": JOB.id})
    assert posted.job_id == JOB.id
    [s] = await jobs.summaries(books)
    assert s.materials_pence == 10000
    await undo(books, posted.event_id or "", USER)
    [s] = await jobs.summaries(books)
    assert s.materials_pence == 0


async def test_sales_invoices_cannot_be_tagged(books: Books) -> None:
    p = _proposal("abcdef0123456780", None, None, status="proposed")
    p.extracted = _doc(doc_type="sales_invoice", customer_name="Mrs Patel")
    await books.save_proposal(p)
    with pytest.raises(ApprovalError, match="Only purchases"):
        await approve(books, p.id, USER, {"job_id": JOB.id})


async def test_a_broken_jobs_file_refuses_the_tag_not_the_request(fake: FakeStorageProvider, books: Books) -> None:
    p = await _open(books)
    fake.files[jobs.JOBS] = b"[{not json"
    with pytest.raises(ApprovalError, match="jobs file could not be read"):
        await approve(books, p.id, USER, {"job_id": JOB.id})
    assert (await approve(books, p.id, USER)).status == "posted"  # untagged approval still works


# --------------------------------------------------------------------------- #
# Routes
# --------------------------------------------------------------------------- #


@pytest.fixture
def client(fake: FakeStorageProvider):
    from app.main import app
    from app.routes_books import get_books

    async def override() -> Books:
        return Books(fake)

    app.dependency_overrides[get_books] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


def test_routes_end_to_end(client: TestClient, fake: FakeStorageProvider) -> None:
    job = client.post("/jobs", json={"name": "Loft boarding", "customer_id": "c9"}, headers=H).json()
    sam = client.post(
        "/staff",
        json={"name": "Sam", "role": "Carpenter", "hourly_rate_pence": 1800, "charge_rate_pence": 4000},
        headers=H,
    ).json()
    add = {"staff_id": sam["id"], "job_id": job["id"], "date": "2026-10-20", "note": "Joists"}
    e1 = client.post("/timesheets", json={**add, "hours": 7.6}, headers=H).json()
    e2 = client.post("/timesheets", json={**add, "hours": 1, "billable": False}, headers=H).json()
    assert e1["hours"] == 7.5
    assert client.post("/timesheets", json={**add, "hours": 0}, headers=H).status_code == 400

    s = client.get(f"/jobs/{job['id']}/summary", headers=H).json()
    assert (s["hours"], s["labour_cost_pence"], s["billable_value_pence"]) == (8.5, 15300, 30000)

    prop = client.post(f"/jobs/{job['id']}/invoice-lines", headers=H).json()
    assert [(ln["qty"], ln["net_pence"]) for ln in prop["lines"]] == [(7.5, 30000)]
    assert prop["entry_ids"] == [e1["id"]]
    assert all(r["invoiced_in"] is None for r in _sheet(fake, "2026-10"))  # proposing marks nothing

    r = client.post(
        f"/jobs/{job['id']}/mark-invoiced", json={"invoice_id": "INV-0007", "entry_ids": prop["entry_ids"]}, headers=H
    )
    assert r.status_code == 200, r.text
    assert client.put(f"/timesheets/{e1['id']}", json={"hours": 2}, headers=H).status_code == 409
    assert client.put(f"/timesheets/{e2['id']}", json={"colour": "red"}, headers=H).status_code == 400
    assert client.put(f"/timesheets/{e2['id']}", json={"hours": 2}, headers=H).json()["hours"] == 2
    assert client.post(f"/jobs/{job['id']}/invoice-lines", headers=H).json()["lines"] == []

    assert [j["id"] for j in client.get("/jobs/list?status=open", headers=H).json()] == [job["id"]]
    assert client.get("/jobs/list?status=closed", headers=H).json() == []
    [listed] = client.get("/jobs", headers=H).json()
    assert listed["invoiced_value_pence"] == 30000 and listed["unbilled_hours"] == 0
    assert client.delete(f"/jobs/{job['id']}", headers=H).status_code == 409
    assert client.get("/jobs/nothex/summary", headers=H).status_code == 404
    assert client.post("/jobs", json={"name": "x"}, headers={"X-Graph-Token": "tok"}).status_code == 400  # no X-User
    assert {e["id"] for e in client.get(f"/timesheets?job_id={job['id']}", headers=H).json()} == {e1["id"], e2["id"]}
    assert "Books.xlsx" in fake.files  # rebuilt after marking invoiced


def test_mark_invoiced_refuses_empty_lists_and_other_jobs_entries(
    client: TestClient, fake: FakeStorageProvider
) -> None:
    a = client.post("/jobs", json={"name": "Kitchen"}, headers=H).json()
    b = client.post("/jobs", json={"name": "Bathroom"}, headers=H).json()
    sam = client.post("/staff", json={"name": "Sam", "hourly_rate_pence": 1800}, headers=H).json()
    e = client.post(
        "/timesheets", json={"staff_id": sam["id"], "job_id": b["id"], "date": "2026-10-20", "hours": 2}, headers=H
    ).json()
    url = f"/jobs/{a['id']}/mark-invoiced"
    assert client.post(url, json={"invoice_id": "INV-1", "entry_ids": []}, headers=H).status_code == 422
    r = client.post(url, json={"invoice_id": "INV-1", "entry_ids": [e["id"]]}, headers=H)
    assert r.status_code == 400 and "another job" in r.json()["detail"]
    assert all(row["invoiced_in"] is None for row in _sheet(fake, "2026-10"))
    assert client.post("/jobs", json={"name": "x", "site_address": "y" * 501}, headers=H).status_code == 422


async def test_books_xlsx_jobs_sheet(fake: FakeStorageProvider, books: Books) -> None:
    import io

    from openpyxl import load_workbook

    await jobs.add_entry(books, entry(SAM, 8))
    await jobs.add_entry(books, entry(ALEX, 2.5))
    fake.files[jobs.CONTACTS] = json.dumps(
        [{"id": "c1", "kind": "customer", "name": "Mrs Patel", "normalised_key": "mrs patel"}]
    ).encode()
    p = await _open(books)
    await approve(books, p.id, USER, {"job_id": JOB.id})  # rebuilds Books.xlsx
    ws = load_workbook(io.BytesIO(fake.files["Books.xlsx"]))["Jobs"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == (
        "Job",
        "Customer",
        "Status",
        "Hours",
        "Labour cost £",
        "Materials £",
        "Billable value £",
        "Invoiced £",
        "Margin £",
    )
    alex = jobs.pence_for(10, 1133)  # 2.5h at cost (no charge rate)
    labour, value = 8 * 2000 + alex, 8 * 4500 + alex
    assert rows[1] == (
        JOB.name,
        "Mrs Patel",
        "open",
        10.5,
        labour / 100,
        100.0,
        value / 100,
        0.0,
        (value - labour - 10000) / 100,
    )
    assert rows[2][0] == "Total"


async def test_a_broken_jobs_file_leaves_the_rest_of_books_xlsx(fake: FakeStorageProvider, books: Books) -> None:
    fake.files[jobs.STAFF] = b'[{"id": "x"}]'  # no name or rate
    await books.rebuild()
    assert "Books.xlsx" in fake.files
