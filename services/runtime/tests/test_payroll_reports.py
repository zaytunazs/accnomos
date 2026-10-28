"""Payroll reports (app/payroll_reports.py) in every format, and the /payroll/paye routes end to end."""

from __future__ import annotations

import base64
import csv
import io
from collections.abc import AsyncIterator, Iterator
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_payrun import ann, run_and_post, setup  # sibling module: pytest puts tests/ on the path
from test_rti import FPS_XSD, validate

from accnomos_connectors import FakeStorageProvider
from app import payroll_reports as REP
from app import payrun as PR
from app import rti
from app.books import Books


@pytest.fixture
def books() -> Books:
    return Books(FakeStorageProvider())


async def posted(books: Books) -> tuple[PR.PayRun, dict[str, PR.Employee]]:
    await setup(books, ann(student_loan_plan="1"), ann(payroll_id="E002", forenames="Bo", annual_salary_pence=900_000))
    run = await run_and_post(books, date(2026, 10, 31))
    return run, {e.id: e for e in await PR.load_employees(books)}


async def test_payroll_summary_totals_match_the_run(books: Books) -> None:
    run, _ = await posted(books)
    t = REP.payroll_summary(run)
    assert [r["name"] for r in t.rows] == ["Ann Lee", "Bo Lee"]
    assert t.totals and t.totals["net"] == run.total("net_pay_pence") and t.totals["tax"] == run.total("tax_pence")
    assert any(PR.FILING_NOTE in n for n in t.notes) and any("Estimate — check before filing" in n for n in t.notes)


async def test_p11_p32_pension_fps_journal(books: Books) -> None:
    run, emps = await posted(books)
    a = next(e for e in emps.values() if e.payroll_id == "E001")
    p11 = REP.p11(a, await PR.list_runs(books))
    (row,) = p11.rows
    assert row["period"] == "Month 7" and row["tax_ytd"] == run.lines[0].ytd.tax_pence
    p32 = REP.p32(
        rti.employer_payment_record(await PR.list_runs(books), await PR.load_employer(books), date(2026, 4, 6))
    )
    assert len(p32.rows) == 12 and p32.rows[6]["due_by"] == "2026-11-22"
    pen = REP.pension_workings(run)
    assert pen.rows[0]["category"].startswith("Eligible jobholder") and "qualifying earnings" in pen.rows[0]["band"]
    assert pen.rows[0]["relief"].startswith("Relief at source")
    assert pen.totals and pen.totals["to_provider"] == sum(ln.pension.to_provider_pence for ln in run.lines)
    fps = REP.fps_contents(run, emps)
    assert {"employee": "Ann Lee", "item": "Tax code", "value": "1257L"} in fps.rows
    j = REP.journal_rows(run)
    assert j.totals and j.totals["debit"] == j.totals["credit"]


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "pdf"])
async def test_every_report_renders_in_every_format(books: Books, fmt: str) -> None:
    run, emps = await posted(books)
    a = next(iter(emps.values()))
    tables = [
        REP.payroll_summary(run),
        REP.p11(a, await PR.list_runs(books)),
        REP.pension_workings(run),
        REP.fps_contents(run, emps),
        REP.journal_rows(run),
        REP.p60_table(rti.p60_figures(a, run.lines[0])),
    ]
    for t in tables:
        data, media, name = REP.render(t, fmt, "Acme Ltd")  # type: ignore[arg-type]
        assert name.endswith(f".{fmt}") and data
        if fmt == "csv":
            rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
            assert rows[0] == [t.title] and rows[3] == [c.label for c in t.columns]
        elif fmt == "xlsx":
            ws = load_workbook(io.BytesIO(data)).active
            assert ws is not None and ws["A1"].value == t.title
        else:
            assert data[:5] == b"%PDF-" and media == "application/pdf"


def test_csv_guards_typed_in_text() -> None:
    t = REP.Table(
        "x", "T", "S", [REP.Column("name", "Name"), REP.Column("n", "N", "money")], [{"name": "=1+1", "n": -500}]
    )
    rows = list(csv.reader(io.StringIO(REP.to_csv(t).decode("utf-8-sig"))))
    assert rows[4] == ["'=1+1", "-5.00"]


# --------------------------------------------------------------------------- #
# Routes
# --------------------------------------------------------------------------- #


@pytest.fixture
def client() -> Iterator[tuple[object, FakeStorageProvider]]:
    from fastapi.testclient import TestClient

    from app.main import app
    from app.routes_books import get_books

    fake = FakeStorageProvider()

    async def override() -> AsyncIterator[Books]:
        yield Books(fake)

    app.dependency_overrides[get_books] = override
    try:
        yield TestClient(app, headers={"X-Graph-Token": "t", "X-User": "owner@example.co.uk"}), fake
    finally:
        app.dependency_overrides.clear()


def test_routes_end_to_end(client: tuple[object, FakeStorageProvider], tmp_path: Path) -> None:
    c, fake = client
    from fastapi.testclient import TestClient

    assert isinstance(c, TestClient)
    s = c.get("/payroll/paye/setup").json()
    assert s["submission"]["available"] is False and "Basic PAYE Tools" in s["filing_note"]
    assert c.put("/payroll/paye/employer", json={"rti_test_submission_enabled": True}).status_code == 400
    r = c.put(
        "/payroll/paye/employer",
        json={
            "name": "Acme Ltd",
            "paye_office": "123",
            "paye_reference": "AB45678",
            "accounts_office_ref": "123PA00045678",
        },
    )
    assert r.status_code == 200, r.text
    emp = ann(pension_exempt=True).model_dump(mode="json")
    r = c.post("/payroll/paye/employees", json=emp)
    assert r.status_code == 200, r.text
    eid = r.json()["id"]
    assert c.post("/payroll/paye/employees", json={**emp, "nino": "bad"}).status_code == 400
    run = c.post("/payroll/paye/payruns/prepare", json={"frequency": "monthly", "pay_date": "2026-10-31"}).json()
    rid = run["id"]
    assert run["status"] == "draft" and run["journal"]
    assert c.get(f"/payroll/paye/payruns/{rid}/fps").status_code == 409  # not approved yet
    slip = c.get(f"/payroll/paye/payruns/{rid}/payslips/{eid}").json()
    assert base64.b64decode(slip["base64"])[:5] == b"%PDF-" and slip["saved_to"] is None  # drafts aren't filed
    edited = c.put(
        f"/payroll/paye/payruns/{rid}", json={"inputs": [{"employee_id": eid, "bonus_pence": 10_000}]}
    ).json()
    assert edited["lines"][0]["gross_pence"] == 310_000
    p = c.post(f"/payroll/paye/payruns/{rid}/propose").json()
    assert p["run"]["status"] == "proposed" and p["proposal_id"]
    assert f"_System/proposals/{p['proposal_id']}.json" in fake.files  # waiting in Review
    done = c.post(f"/payroll/paye/payruns/{rid}/approve").json()
    assert done["status"] == "posted"
    f = c.get(f"/payroll/paye/payruns/{rid}/fps").json()
    xml = base64.b64decode(f["file"]["base64"])
    validate(xml, FPS_XSD, tmp_path)
    assert f["file"]["saved_to"] == "05 Payroll/2026-27/RTI/FPS-2026-27-monthly-07.xml"
    assert "Direct submission from Accnomos is currently not available" in f["file"]["note"]
    assert any(r["item"] == "Tax code" for r in f["contents"]["rows"])
    e = c.get("/payroll/paye/eps/7?employment_allowance=true").json()
    assert base64.b64decode(e["file"]["base64"]).startswith(b"<?xml")
    for provider in ("generic", "nest", "peoples_pension", "now_pensions"):
        assert c.get(f"/payroll/paye/payruns/{rid}/pension/{provider}").status_code == 200
    smart = c.get(f"/payroll/paye/payruns/{rid}/pension/smart_pension")  # this scheme is relief at source
    assert smart.status_code == 400 and "net pay arrangement only" in smart.json()["detail"]
    for key, q in (
        ("payroll_summary", f"run={rid}"),
        ("p32", ""),
        ("p11", f"employee={eid}"),
        ("p60", f"employee={eid}"),
    ):
        j = c.get(f"/payroll/paye/reports/{key}?{q}").json()
        assert j["rows"], key
        x = c.get(f"/payroll/paye/reports/{key}?{q}&format=xlsx").json()
        assert x["name"].endswith(".xlsx")
    assert c.get(f"/payroll/paye/reports/p45?employee={eid}").status_code == 400  # not a leaver
    assert c.get("/payroll/paye/rti/status").json()["submission_available"] is False
    assert any(path.startswith("05 Payroll/2026-27/payslips/") for path in fake.files) is False
    c.get(f"/payroll/paye/payruns/{rid}/payslips/{eid}")
    assert any(path.startswith("05 Payroll/2026-27/payslips/") for path in fake.files)
