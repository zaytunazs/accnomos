"""Bank import + matching routes and the Books.xlsx projection, against FakeStorageProvider."""

from __future__ import annotations

import io
import json
from datetime import UTC, date, datetime

import pytest
from accnomos_connectors import FakeStorageProvider
from accnomos_ledger import LedgerEvent, Line, VatCode
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.approve import approve
from app.books import AUDIT, BANK_LINES, WORKBOOK, BankMatchError, Books
from app.projection import SHEETS, build_books_xlsx, last_quarters
from app.schemas import ExtractedDoc, Proposal

USER = "sam@example.co.uk"
STATEMENT = b"""Date,Description,Amount,Balance
21/10/2026,CARD PAYMENT SCREWFIX 1234,-38.47,961.53
22/10/2026,ACME KITCHENS INV 42,"1,200.00","2,161.53"
23/10/2026,COFFEE,-2.50,"2,159.03"
"""


def _event(eid: str, gross: int, when: date, who: str, sale: bool = False) -> LedgerEvent:
    net = round(gross / 1.2)
    lines = (
        [Line(account="1100 Debtors", debit=gross), Line(account="4000 Sales", credit=net, vat_code=VatCode.STANDARD),
         Line(account="2200 VAT control", credit=gross - net)]
        if sale else
        [Line(account="5100 Tools & small equipment", debit=net, vat_code=VatCode.STANDARD),
         Line(account="2200 VAT control", debit=gross - net), Line(account="1200 Bank", credit=gross)]
    )
    return LedgerEvent(id=eid, kind="sale" if sale else "purchase", date=when, lines=lines, net=net, vat=gross - net,
                       counterparty=who, evidence=["ab" * 32], approved_by=USER, approved_at=datetime(2026, 10, 25, tzinfo=UTC))


EVENTS = [
    _event("e" * 32, 3847, date(2026, 10, 21), "Screwfix"),
    _event("f" * 32, 120000, date(2026, 10, 22), "Acme Kitchens Ltd", sale=True),
    _event("a" * 32, 1200, date(2026, 8, 3), "Toolstation"),
]


@pytest.fixture
def fake() -> FakeStorageProvider:
    return FakeStorageProvider()


@pytest.fixture
async def books(fake: FakeStorageProvider) -> Books:
    b = Books(fake)
    for ev in EVENTS:
        await b.post_event(ev)
    return b


@pytest.fixture
def client(fake: FakeStorageProvider, books: Books):
    from app.main import app
    from app.routes_books import get_books

    async def override() -> Books:
        return books

    app.dependency_overrides[get_books] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


H = {"X-Graph-Token": "tok", "X-User": USER}


def _audit(fake: FakeStorageProvider) -> list[dict]:
    return [json.loads(x) for p, d in sorted(fake.files.items()) if p.startswith(f"{AUDIT}/") for x in d.decode().splitlines()]


def test_import_suggest_match_unmatch_routes(client: TestClient, fake: FakeStorageProvider) -> None:
    r = client.post("/books/bank/import", headers=H, files={"file": ("oct.csv", STATEMENT, "text/csv")})
    assert r.status_code == 200, r.text
    assert r.json()["filed_path"] == "03 Bank statements/2026-10/oct.csv" and r.json()["new_lines"] == 3
    assert fake.files["03 Bank statements/2026-10/oct.csv"] == STATEMENT
    assert len([p for p in fake.files if p.startswith(f"{BANK_LINES}/")]) == 1
    again = client.post("/books/bank/import", headers=H, files={"file": ("copy.csv", STATEMENT, "text/csv")}).json()
    assert again["duplicate_file"] and again["new_lines"] == 0

    sugg = client.get("/books/bank/suggestions", headers=H).json()
    assert sorted(s["event_id"] for s in sugg) == ["e" * 32, "f" * 32] and all(s["kind"] == "exact" for s in sugg)
    assert client.get("/books/bank/summary", headers=H).json() == {"total": 3, "matched": 0, "suggested": 2, "unmatched": 1}

    screwfix = next(s for s in sugg if s["event_id"] == "e" * 32)
    r = client.post(f"/books/bank/lines/{screwfix['line_id']}/match", headers=H, json={"event_id": "e" * 32})
    assert r.status_code == 200 and r.json()["status"] == "matched"
    assert client.post(f"/books/bank/lines/{screwfix['line_id']}/match", headers=H, json={"event_id": "e" * 32}).status_code == 409
    coffee = client.get("/books/bank/lines?status=unmatched", headers=H).json()[0]
    assert coffee["description"] == "COFFEE"
    r = client.post(f"/books/bank/lines/{coffee['id']}/match", headers=H, json={"event_id": "e" * 32})
    assert r.status_code == 409 and "already matched" in r.json()["detail"]
    assert client.post("/books/bank/lines/0000000000000000/match", headers=H, json={"event_id": "x"}).status_code == 404
    assert client.get("/books/bank/summary", headers=H).json() == {"total": 3, "matched": 1, "suggested": 1, "unmatched": 1}
    assert [ln["matched_event_id"] for ln in client.get("/books/bank/lines?status=matched", headers=H).json()] == ["e" * 32]

    r = client.post(f"/books/bank/lines/{screwfix['line_id']}/unmatch", headers=H)
    assert r.status_code == 200 and r.json()["status"] == "unmatched" and r.json()["matched_event_id"] is None
    assert client.post(f"/books/bank/lines/{screwfix['line_id']}/unmatch", headers=H).status_code == 409
    actions = [a["action"] for a in _audit(fake)]
    assert actions == ["bank_imported", "bank_matched", "bank_unmatched"]
    assert WORKBOOK in fake.files  # rebuilt after import/match/unmatch
    assert not any(p.startswith("_System/Ledger/events/") and "bank" in p for p in fake.files)  # matching posts nothing


def test_import_rejects_non_statement(client: TestClient) -> None:
    r = client.post("/books/bank/import", headers=H, files={"file": ("x.csv", b"name,email\nSam,x\n", "text/csv")})
    assert r.status_code == 400


async def test_confirm_match_validates_event(books: Books) -> None:
    await books.import_bank_csv(STATEMENT, "oct.csv")
    line = (await books.list_bank_lines())[0]
    with pytest.raises(BankMatchError, match="No posted ledger entry"):
        await books.confirm_match(line.id, "9" * 32, USER)
    with pytest.raises(FileNotFoundError):
        await books.confirm_match("../../etc", "e" * 32, USER)


async def test_overlapping_statements_do_not_duplicate_lines(books: Books) -> None:
    await books.import_bank_csv(STATEMENT, "oct.csv")
    overlap = STATEMENT + b'24/10/2026,TOOLSTATION,-5.00,"2,154.03"\n'
    result = await books.import_bank_csv(overlap, "oct-later.csv")
    assert (result.lines, result.new_lines) == (4, 1)
    assert len(await books.stored_bank_lines()) == 4


# --------------------------------------------------------------------------- #
# Books.xlsx
# --------------------------------------------------------------------------- #


def _proposal() -> Proposal:
    doc = ExtractedDoc(
        doc_type="receipt", supplier_name="Wickes", supplier_vat_number=None, customer_name=None, invoice_number=None,
        date="2026-10-26", currency="GBP", lines=[], net_total=10, vat_total=2, gross_total=12, payment_method="card",
        card_last4=None, category_guess="Materials", confidence=0.6,
    )
    return Proposal(id="c" * 16, doc_sha256="cd" * 32, extracted=doc, checks=[], status="needs_a_look",
                    reasons=["No VAT number"], suggested_filename="x.jpg", journal_draft=[], model="stub", elapsed_s=1.0,
                    created_at=datetime(2026, 10, 26, 9, tzinfo=UTC))


def test_last_quarters() -> None:
    qs = last_quarters(date(2026, 10, 27))
    assert [k for _, _, k in qs] == ["25A4", "26A1", "26A2", "26A3"]
    assert qs[0][:2] == (date(2025, 10, 1), date(2025, 12, 31)) and qs[-1][1] == date(2026, 9, 30)


async def test_build_books_xlsx_has_six_sheets_with_expected_cells(books: Books) -> None:
    await books.import_bank_csv(STATEMENT, "oct.csv")
    lines = await books.list_bank_lines()
    data = build_books_xlsx(EVENTS, lines, [_proposal()], generated=datetime(2026, 10, 27, 12, tzinfo=UTC))
    wb = load_workbook(io.BytesIO(data))
    assert tuple(wb.sheetnames) == SHEETS

    readme = [c.value for c in wb["Read me"]["A"]]
    assert "SAMPLE — FICTIONAL COMPANY" in readme and wb["Read me"]["B3"].value == "27/10/2026 12:00 UTC"

    j = wb["Journal"]
    assert [c.value for c in j[1]] == ["Date", "Event id", "Account", "Debit £", "Credit £", "VAT code", "Counterparty", "Evidence doc"]
    assert j["A1"].font.bold and j.freeze_panes == "A2" and j.column_dimensions["C"].width == 32
    assert j.max_row == 1 + 9
    first = [c.value for c in j[2]]
    assert first[1:4] == ["e" * 32, "5100 Tools & small equipment", 32.06] and first[5:7] == ["S", "Screwfix"]
    assert j["D2"].number_format.startswith('"£"')

    acc = {r[0]: r[1:] for r in wb["Accounts"].iter_rows(min_row=2, values_only=True)}
    assert acc["1200 Bank"] == (0, 50.47, -50.47) and acc["Total"][2] == 0
    assert acc["1100 Debtors"] == (1200, 0, 1200)

    vat = wb["VAT"]
    assert vat["C1"].value.startswith("25A4") and vat["F1"].value.startswith("26A3")
    q3 = {r[0]: r[5] for r in vat.iter_rows(min_row=2, values_only=True)}
    assert q3[4] == 2.0 and q3[7] == 10 and q3[1] == 0  # Toolstation in Aug; October events are in 26A4

    bank = list(wb["Bank"].iter_rows(min_row=2, values_only=True))
    assert len(bank) == 3 and bank[0][3] == -38.47 and bank[0][5] == "suggested" and bank[0][6] == "e" * 32

    (prop,) = wb["Proposals"].iter_rows(min_row=2, values_only=True)
    assert prop[1:4] == ("c" * 16, "needs_a_look", "Wickes") and prop[5] == 12


async def test_rebuild_after_approve_writes_books_xlsx(fake: FakeStorageProvider) -> None:
    from app.books import PROPOSALS

    books = Books(fake)
    p = _proposal()
    p.status, p.id = "proposed", "d" * 16
    p.extracted = p.extracted.model_copy(update={"supplier_vat_number": "GB 232 4567 89"})
    fake.files[f"{PROPOSALS}/{p.id}.json"] = p.model_dump_json().encode()

    await approve(books, p.id, USER)

    wb = load_workbook(io.BytesIO(fake.files[WORKBOOK]))
    assert wb["Journal"].max_row > 1 and wb["Journal"]["B2"].value == p.id
    assert wb["Proposals"].max_row == 1  # no longer open
    assert (await books.workbook_ref()) is not None


async def test_rebuild_failure_never_blocks_the_action(fake: FakeStorageProvider, monkeypatch) -> None:
    import app.projection as projection

    def boom(*_a: object, **_k: object) -> bytes:
        raise RuntimeError("disk full")

    monkeypatch.setattr(projection, "build_books_xlsx", boom)
    books = Books(fake)
    assert await books.rebuild_safely(USER) is None
    assert _audit(fake)[-1]["action"] == "books_rebuild_failed" and "disk full" in _audit(fake)[-1]["refs"]["error"]


def test_rebuild_route_and_summary_books_url(client: TestClient, fake: FakeStorageProvider) -> None:
    assert client.get("/books/summary", headers=H).json()["books_url"] is None
    r = client.post("/books/rebuild", headers=H)
    assert r.status_code == 200 and r.json()["path"] == WORKBOOK and r.json()["size"] > 0
    assert WORKBOOK in fake.files
    assert "books_url" in client.get("/books/summary", headers=H).json()
