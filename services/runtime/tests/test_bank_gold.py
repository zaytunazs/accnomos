"""The gold bank statements (evals/gold/bank/) through the /books/bank/* routes: upload, statements,
suggested matches, confirm, the three unmatched-line actions and "reconciled to"."""

from __future__ import annotations

import io
import json
from datetime import UTC, date, datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from accnomos_connectors import FakeStorageProvider
from accnomos_ledger import LedgerEvent, Line, VatCode
from app.books import BANK, WORKBOOK, Books

GOLD = Path(__file__).resolve().parents[3] / "evals" / "gold" / "bank"
EXPECTED = json.loads((GOLD / "expected.json").read_text())
USER = "sam@example.co.uk"
H = {"X-Graph-Token": "tok", "X-User": USER}


def _event(spec: dict) -> LedgerEvent:
    gross = spec["gross_pence"]
    if spec["kind"] == "sale":
        lines = [
            Line(account="1100 Debtors", debit=gross),
            Line(account="4000 Sales", credit=gross, vat_code=VatCode.STANDARD),
        ]
    else:
        lines = [
            Line(account="7999 Sundry expenses", debit=gross, vat_code=VatCode.STANDARD),
            Line(account="1200 Bank", credit=gross),
        ]
    return LedgerEvent(
        id=spec["id"],
        kind=spec["kind"],
        date=date.fromisoformat(spec["date"]),
        lines=lines,
        net=gross,
        vat=0,
        counterparty=spec["counterparty"],
        evidence=["ab" * 32],
        approved_by=USER,
        approved_at=datetime(2026, 10, 26, tzinfo=UTC),
    )


@pytest.fixture
def fake() -> FakeStorageProvider:
    return FakeStorageProvider()


@pytest.fixture
async def books(fake: FakeStorageProvider) -> Books:
    b = Books(fake)
    for spec in EXPECTED["events"]:
        await b.post_event(_event(spec))
    return b


@pytest.fixture
def client(books: Books):
    from app.main import app
    from app.routes_books import get_books

    async def override() -> Books:
        return books

    app.dependency_overrides[get_books] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


def _upload(client: TestClient, name: str, as_name: str | None = None) -> dict:
    r = client.post(
        "/books/bank/import", headers=H, files={"file": (as_name or name, (GOLD / name).read_bytes(), "text/csv")}
    )
    assert r.status_code == 200, r.text
    return r.json()


@pytest.mark.parametrize("name", sorted(EXPECTED["statements"]))
def test_upload_files_as_date_and_bank_and_is_idempotent(
    client: TestClient, fake: FakeStorageProvider, name: str
) -> None:
    want = EXPECTED["statements"][name]
    first = _upload(client, name, "Statement export (1).csv")
    path = f"{BANK}/{want['last_date']}_{want['bank']}.csv"
    assert (first["filed_path"], first["bank"], first["new_lines"]) == (path, want["bank"], len(want["lines"]))
    assert fake.files[path] == (GOLD / name).read_bytes()

    again = _upload(client, name, "renamed.csv")  # same bytes, same sha256: nothing new
    assert again["duplicate_file"] and again["filed_path"] is None and again["new_lines"] == 0
    assert [p for p in fake.files if p.startswith(f"{BANK}/")] == [path]
    (stmt,) = client.get("/books/bank/statements", headers=H).json()
    assert (stmt["path"], stmt["bank"], stmt["lines"], stmt["last_date"]) == (
        path,
        want["bank"],
        len(want["lines"]),
        want["last_date"],
    )


def test_gold_suggestions_unmatched_actions_and_reconciled_to(client: TestClient, fake: FakeStorageProvider) -> None:
    for name in EXPECTED["statements"]:
        _upload(client, name)
    want = [w for s in EXPECTED["statements"].values() for w in s["lines"]]
    by_event = {m["event_id"]: m for m in client.get("/books/bank/suggestions", headers=H).json()}
    assert sorted(by_event) == sorted(w["match"] for w in want if "match" in w)
    assert {e: m["kind"] for e, m in by_event.items()} == {w["match"]: w["kind"] for w in want if "match" in w}
    travis = by_event["gold-bc-travis"]  # a penny out and five days apart, carried by the name
    assert (travis["line"]["amount_pence"], travis["events"]) == (
        -4501,
        [{"id": "gold-bc-travis", "date": "2026-10-15", "counterparty": "Travis Perkins", "amount_pence": -4500}],
    )

    unmatched = client.get("/books/bank/unmatched", headers=H).json()
    got = sorted((u["line"]["description"], u["line"]["amount_pence"], u["suggested"]["action"]) for u in unmatched)
    assert got == sorted((w["description"], w["amount_pence"], w["action"]) for w in want if "action" in w)
    assert client.get("/books/bank/summary", headers=H).json()["reconciled_to"] is None

    # Confirm every suggestion, then take each unmatched line's suggested action.
    for m in by_event.values():
        r = client.post(f"/books/bank/lines/{m['line_id']}/match", headers=H, json={"event_ids": m["event_ids"]})
        assert r.status_code == 200 and r.json()["status"] == "matched"
    proposals = []
    for u in unmatched:
        lid, action = u["line"]["id"], u["suggested"]["action"]
        if action == "create_expense_proposal":
            r = client.post(f"/books/bank/lines/{lid}/expense-proposal", headers=H)
            assert r.status_code == 200, r.text
            proposals.append((lid, r.json()))
        else:
            r = client.post(
                f"/books/bank/lines/{lid}/{'transfer' if action == 'mark_transfer' else 'ignore'}", headers=H
            )
            assert r.status_code == 200 and r.json()["status"] == (
                "transfer" if action == "mark_transfer" else "ignored"
            )

    # Expense proposals wait on /review; the lines stay open until the proposal is posted.
    summary = client.get("/books/bank/summary", headers=H).json()
    assert summary["unmatched"] == len(proposals) and summary["reconciled_to"] == "2026-10-01"  # Pret, 2 Oct, is open
    lid, p = proposals[0]
    assert p["status"] == "needs_a_look" and p["extracted"]["vat_total"] == 0
    again = client.post(f"/books/bank/lines/{lid}/expense-proposal", headers=H).json()
    assert again["id"] == p["id"]  # asking twice does not make a second proposal

    # Once approved on /review, each proposal's bank leg matches its own line.
    for _, p in proposals:
        assert client.post(f"/books/approve/{p['id']}", headers=H, json={}).status_code == 200
    sugg = client.get("/books/bank/suggestions", headers=H).json()
    assert sorted(s["line_id"] for s in sugg) == sorted(lid for lid, _ in proposals)
    assert all(s["kind"] == "exact" for s in sugg)
    for s in sugg:
        assert (
            client.post(
                f"/books/bank/lines/{s['line_id']}/match", headers=H, json={"event_ids": s["event_ids"]}
            ).status_code
            == 200
        )
    summary = client.get("/books/bank/summary", headers=H).json()
    assert summary["unmatched"] == summary["suggested"] == 0 and summary["reconciled_to"] == "2026-10-25"

    # Books.xlsx Bank sheet: every line, its status, the matched event and any expense proposal.
    rows = list(load_workbook(io.BytesIO(fake.files[WORKBOOK]))["Bank"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(want) and {r[5] for r in rows} == {"matched", "transfer", "ignored"}
    assert all(r[6] for r in rows if r[5] == "matched")
    assert sorted(r[8] for r in rows if r[8]) == sorted(p["id"] for _, p in proposals)
    summary_sheet = {
        r[0]: r[1] for r in load_workbook(io.BytesIO(fake.files[WORKBOOK]))["Summary"].iter_rows(values_only=True)
    }
    assert summary_sheet["Bank reconciled to"] == datetime(2026, 10, 25)


def test_line_actions_refuse_and_reopen(client: TestClient) -> None:
    _upload(client, "starling.csv")
    unmatched = {
        u["line"]["description"]: u["line"]["id"] for u in client.get("/books/bank/unmatched", headers=H).json()
    }
    savings, payee = unmatched["Starling Savings Space"], unmatched["Unknown Payee"]
    assert client.post(f"/books/bank/lines/{savings}/transfer", headers=H).json()["status"] == "transfer"
    r = client.post(f"/books/bank/lines/{savings}/ignore", headers=H)
    assert r.status_code == 409 and "transfer" in r.json()["detail"]
    assert client.post(f"/books/bank/lines/{savings}/unmatch", headers=H).json()["status"] == "unmatched"
    assert client.post(f"/books/bank/lines/{payee}/expense-proposal", headers=H).status_code == 409  # money in
    assert client.post("/books/bank/lines/0000000000000000/ignore", headers=H).status_code == 404
    assert client.get("/books/bank/lines?status=transfer", headers=H).json() == []


def test_rejected_expense_proposal_can_be_made_again(client: TestClient) -> None:
    _upload(client, "starling.csv")
    ee = next(
        u["line"]["id"]
        for u in client.get("/books/bank/unmatched", headers=H).json()
        if u["line"]["description"] == "EE LIMITED"
    )
    first = client.post(f"/books/bank/lines/{ee}/expense-proposal", headers=H).json()
    status = {u["line"]["id"]: u["proposal_status"] for u in client.get("/books/bank/unmatched", headers=H).json()}
    assert status[ee] == "needs_a_look"

    assert client.post(f"/books/proposals/{first['id']}/reject", headers=H, json={"reason": "wrong"}).status_code == 200
    status = {u["line"]["id"]: u["proposal_status"] for u in client.get("/books/bank/unmatched", headers=H).json()}
    assert status[ee] == "rejected"  # the page offers "Create expense proposal" again
    second = client.post(f"/books/bank/lines/{ee}/expense-proposal", headers=H).json()
    assert second["id"] != first["id"] and second["status"] == "needs_a_look"
    line = next(ln for ln in client.get("/books/bank/lines", headers=H).json() if ln["id"] == ee)
    assert line["proposal_id"] == second["id"]
