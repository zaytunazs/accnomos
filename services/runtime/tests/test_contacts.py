"""Contacts from the ledger: dedupe, manual edits, statements, aged balances, Outlook import. Fake providers only."""

from __future__ import annotations

import json
from datetime import UTC, date, datetime

import pytest
from accnomos_connectors import AddressBookContact, FakeContactsProvider, FakeStorageProvider
from accnomos_ledger import LedgerEvent, Line, VatCode
from fastapi.testclient import TestClient

from app import contacts as cb
from app.books import AUDIT, Books
from app.schemas import ExtractedDoc, Proposal

USER = "sam@example.co.uk"
H = {"X-Graph-Token": "tok", "X-User": USER}
TODAY = date(2026, 10, 27)
YEAR = date(2026, 4, 6)
_n = 0


def _id(prefix: str, length: int = 32) -> str:
    """Unique hex id (event and proposal ids must be hex to be looked up)."""
    global _n
    _n += 1
    return f"{prefix}{_n:0{length - 1}x}"


def _event(gross: int, when: date, who: str, *, sale: bool = False, on_account: bool = True,
           proposal_id: str | None = None, eid: str | None = None) -> LedgerEvent:
    net = round(gross / 1.2)
    lines = (
        [Line(account="1100 Debtors", debit=gross), Line(account="4000 Sales", credit=net, vat_code=VatCode.STANDARD),
         Line(account="2200 VAT control", credit=gross - net)]
        if sale else
        [Line(account="5000 Materials", debit=net, vat_code=VatCode.STANDARD),
         Line(account="2200 VAT control", debit=gross - net),
         Line(account="2100 Creditors" if on_account else "1200 Bank", credit=gross)]
    )
    return LedgerEvent(id=eid or _id("e"), kind="sale" if sale else "purchase", date=when, lines=lines, net=net,
                       vat=gross - net, counterparty=who, approved_by=USER, proposal_id=proposal_id,
                       approved_at=datetime(2026, 10, 25, tzinfo=UTC))


def _reversal(e: LedgerEvent) -> LedgerEvent:
    return LedgerEvent(
        id=_id("c"), kind="reversal", date=e.date, reverses=e.id, net=e.net, vat=e.vat, counterparty=e.counterparty,
        approved_by=USER, approved_at=datetime(2026, 10, 26, tzinfo=UTC), proposal_id=e.proposal_id,
        lines=[Line(account=ln.account, debit=ln.credit, credit=ln.debit, vat_code=ln.vat_code) for ln in e.lines],
    )


def _proposal(supplier: str, vrn: str | None, *, sale: bool = False, invoice: str | None = None,
              category: str = "Materials") -> Proposal:
    doc = ExtractedDoc(
        doc_type="sales_invoice" if sale else "purchase_invoice", supplier_name=supplier, supplier_vat_number=vrn,
        customer_name=None, invoice_number=invoice, date="2026-09-01", currency="GBP", lines=[], net_total=10,
        vat_total=2, gross_total=12, payment_method=None, card_last4=None, category_guess=category, confidence=0.9,
    )
    return Proposal(id=_id("a", 16), doc_sha256="cd" * 32, extracted=doc, checks=[], status="posted", reasons=[],
                    suggested_filename="x.pdf", journal_draft=[], model="stub", elapsed_s=1.0)


def _ledger(events: list[LedgerEvent], proposals: list[Proposal] = (), bank_lines=()) -> cb.Ledger:
    return cb.Ledger(events=events, proposals={p.id: p for p in proposals}, bank_lines=list(bank_lines))


# --------------------------------------------------------------------------- #
# Dedupe
# --------------------------------------------------------------------------- #


def test_dedupe_by_normalised_name_and_vrn() -> None:
    p1 = _proposal("Screwfix Direct Ltd.", "GB 232 4567 89")
    p2 = _proposal("Screwfix", "232456789")  # same VAT number, different name
    p3 = _proposal("Smith Builders", "GB111111111")
    p4 = _proposal("Smith Builders", "GB222222222")  # same name, another business
    own = _proposal("Acme Kitchens Ltd", "GB999999999", sale=True)  # on a sales invoice the VRN is the user's own
    events = [
        _event(3847, date(2026, 9, 1), "Screwfix Direct Ltd.", proposal_id=p1.id),
        _event(1200, date(2026, 9, 2), "SCREWFIX DIRECT"),
        _event(960, date(2026, 9, 3), "Screwfix", proposal_id=p2.id),
        _event(5000, date(2026, 9, 4), "Smith Builders", proposal_id=p3.id),
        _event(7000, date(2026, 9, 5), "Smith Builders", proposal_id=p4.id),
        _event(120000, date(2026, 9, 6), "Acme Kitchens Ltd", sale=True, proposal_id=own.id),
        _event(600, date(2026, 9, 7), "ACME KITCHENS", on_account=False),
    ]
    book = cb.merge_ledger([], _ledger(events, [p1, p2, p3, p4, own]))
    assert len(book) == 4
    screwfix = next(c for c in book if c.vrn == "GB232456789")
    assert screwfix.name == "Screwfix Direct Ltd." and screwfix.normalised_key == "screwfix-direct"
    assert screwfix.kind == "supplier" and screwfix.aliases == ["screwfix"] and screwfix.source == "ledger"
    assert sorted(c.vrn for c in book if c.name == "Smith Builders") == ["GB111111111", "GB222222222"]
    acme = next(c for c in book if c.normalised_key == "acme-kitchens")
    assert acme.kind == "both" and acme.vrn is None
    assert screwfix.default_category == "Materials" and screwfix.default_vat_code == "S"
    # Stable ids: the same ledger gives the same ids, so a list shown before contacts.json exists keeps them.
    assert [c.id for c in cb.merge_ledger([], _ledger(events, [p1, p2, p3, p4, own]))] == [c.id for c in book]
    by_contact = cb.assign(book, cb.documents(_ledger(events, [p1, p2, p3, p4, own])))
    assert len(by_contact[screwfix.id]) == 3 and len(by_contact[acme.id]) == 2


def test_reversal_adds_nobody_and_nets_off() -> None:
    e = _event(2400, date(2026, 9, 1), "Toolstation")
    ledger = _ledger([e, _reversal(e)])
    (c,) = cb.merge_ledger([], ledger)
    t = cb.with_totals([c], ledger, YEAR, TODAY)[0].totals
    assert t.documents == 0 and t.ytd_purchases_net_pence == 0 and t.open_payable_pence == 0


# --------------------------------------------------------------------------- #
# Manual edits survive a refresh
# --------------------------------------------------------------------------- #


def test_manual_edits_survive_refresh() -> None:
    first = _ledger([_event(3000, date(2026, 9, 1), "Howdens")])
    (c,) = cb.merge_ledger([], first)
    cb.apply_edit(c, {"name": "Howdens Joinery Co", "email": " accounts@howdens.example ", "default_category": "Rent",
                      "vrn": "gb 123 4567 89"}, [c])
    assert c.normalised_key == "howdens-joinery" and c.aliases == ["howdens"] and c.email == "accounts@howdens.example"
    assert c.vrn == "GB123456789" and sorted(c.edited) == ["default_category", "email", "name", "vrn"]

    later = _ledger([*first.events, _event(1800, date(2026, 10, 1), "Howdens"),
                     _event(54000, date(2026, 10, 2), "Howdens", sale=True)])
    (again,) = cb.merge_ledger([c], later)  # the old name still finds the renamed contact
    assert again.id == c.id and again.name == "Howdens Joinery Co" and again.default_category == "Rent"
    assert again.email == "accounts@howdens.example" and again.vrn == "GB123456789"
    assert again.kind == "both"  # not edited, so the ledger keeps it up to date
    assert again.default_vat_code == "S"


def test_bad_edits_are_refused() -> None:
    (c,) = cb.merge_ledger([], _ledger([_event(3000, date(2026, 9, 1), "Howdens")]))
    other = c.model_copy(update={"id": "0" * 16, "vrn": "GB123456789"})
    for edits, msg in [({"name": "  "}, "empty"), ({"kind": "friend"}, "Kind"), ({"vrn": "12345"}, "VAT number"),
                       ({"vrn": "GB123456789"}, "Another contact"), ({"email": "nope"}, "email"),
                       ({"default_category": "Snacks"}, "category"), ({"default_vat_code": "X"}, "VAT code"),
                       ({"balance": 1}, "Can't edit")]:
        with pytest.raises(cb.ContactEditError, match=msg):
            cb.apply_edit(c.model_copy(deep=True), edits, [c, other])


# --------------------------------------------------------------------------- #
# Statements and totals
# --------------------------------------------------------------------------- #


def _bank_line(when: date, amount: int, event_ids: list[str]) -> cb.BankLine:
    return cb.BankLine(id=_id("b", 16), date=when, amount_pence=amount, description="ACME KITCHENS INV 42",
                       source_file="03 Bank statements/2026-09/sep.csv", status="matched",
                       matched_event_id=event_ids[0], matched_event_ids=event_ids)


def test_statement_running_balance_payments_and_reversal() -> None:
    inv = _proposal("Me Ltd", None, sale=True, invoice="INV-42")
    s1 = _event(120000, date(2026, 8, 31), "Acme Kitchens Ltd", sale=True, proposal_id=inv.id)
    s2 = _event(54000, date(2026, 9, 15), "Acme Kitchens Ltd", sale=True)
    s3 = _event(6000, date(2026, 7, 1), "Acme Kitchens", sale=True)
    card = _event(3847, date(2026, 9, 16), "Acme Kitchens Ltd", on_account=False)  # bought from them by card
    ledger = _ledger([s3, _reversal(s3), s1, s2, card], [inv], [_bank_line(date(2026, 9, 20), 120000, [s1.id])])
    ledger.events.sort(key=lambda e: (e.date, e.approved_at))
    (acme,) = cb.merge_ledger([], ledger)
    docs = cb.assign([acme], cb.documents(ledger))[acme.id]

    st = cb.statement(acme, docs, ledger)
    rows = [(ln.date.isoformat(), ln.details, ln.amount_pence, ln.paid_pence, ln.balance_pence) for ln in st.lines]
    assert rows == [
        ("2026-07-01", "Sales invoice", 6000, 0, 6000),
        ("2026-07-01", "Undo: sales invoice", -6000, 0, 0),
        ("2026-08-31", "Sales invoice", 120000, 0, 120000),
        ("2026-09-15", "Sales invoice", 54000, 0, 174000),
        ("2026-09-16", "Purchase, paid at the time", -3847, 3847, 174000),
        ("2026-09-20", "Bank: ACME KITCHENS INV 42", 0, -120000, 54000),
    ]
    assert st.lines[2].reference == "INV-42" and st.opening_balance_pence == 0 and st.closing_balance_pence == 54000

    sept = cb.statement(acme, docs, ledger, date(2026, 9, 1), date(2026, 9, 16))
    assert sept.opening_balance_pence == 120000 and [ln.balance_pence for ln in sept.lines] == [174000, 174000]
    assert sept.closing_balance_pence == 174000

    t = cb.totals(docs, ledger, YEAR, TODAY)
    assert t.open_receivable_pence == 54000 and t.open_payable_pence == 0 and t.balance_pence == 54000
    assert t.ytd_sales_net_pence == 100000 + 45000 and t.ytd_sales_vat_pence == 20000 + 9000
    assert t.ytd_purchases_net_pence == 3206 and t.documents == 3 and t.last_activity == date(2026, 9, 16)
    assert cb.totals(docs, ledger, date(2026, 9, 1), TODAY).ytd_sales_net_pence == 45000


def test_uk_tax_year_start() -> None:
    assert cb.uk_tax_year_start(date(2026, 4, 5)) == date(2025, 4, 6)
    assert cb.uk_tax_year_start(date(2026, 4, 6)) == date(2026, 4, 6)


# --------------------------------------------------------------------------- #
# Aged balances
# --------------------------------------------------------------------------- #


def test_bucket_edges() -> None:
    assert [cb.bucket(d) for d in (0, 30, 31, 60, 61, 90, 91, 400)] == \
        ["0-30", "0-30", "31-60", "31-60", "61-90", "61-90", "90+", "90+"]


def test_aged_debtors_and_creditors() -> None:
    paid = _event(9000, date(2026, 8, 1), "Acme Kitchens Ltd", sale=True)
    events = [
        _event(1000, date(2026, 9, 27), "Howdens", eid="1" * 32),  # 30 days
        _event(2000, date(2026, 9, 26), "Howdens"),  # 31
        _event(3000, date(2026, 7, 29), "Howdens"),  # 90
        _event(4000, date(2026, 7, 28), "Howdens"),  # 91
        _event(500, date(2026, 10, 1), "Screwfix", on_account=False),  # paid by card: never owed
        _event(54000, date(2026, 9, 15), "Acme Kitchens Ltd", sale=True),  # 42
        paid,
        _event(7000, date(2026, 11, 2), "Acme Kitchens Ltd", sale=True),  # after the as-of date
        _event(800, date(2026, 6, 1), "Bloggs & Co", sale=True),  # 148
    ]
    ledger = _ledger(events, bank_lines=[_bank_line(date(2026, 8, 20), 9000, [paid.id])])
    contacts = cb.merge_ledger([], ledger)
    a = cb.aged(contacts, ledger, TODAY)

    (howdens,) = a.creditors.rows
    assert howdens.name == "Howdens" and howdens.buckets == {"0-30": 1000, "31-60": 2000, "61-90": 3000, "90+": 4000}
    assert howdens.total_pence == 10000 and [d.age_days for d in howdens.documents] == [30, 31, 90, 91]
    assert a.creditors.total_pence == 10000 and a.creditors.buckets["90+"] == 4000

    assert [r.name for r in a.debtors.rows] == ["Acme Kitchens Ltd", "Bloggs & Co"]
    assert a.debtors.rows[0].buckets == {"0-30": 0, "31-60": 54000, "61-90": 0, "90+": 0}
    assert a.debtors.buckets == {"0-30": 0, "31-60": 54000, "61-90": 0, "90+": 800} and a.debtors.total_pence == 54800


# --------------------------------------------------------------------------- #
# Outlook import mapping
# --------------------------------------------------------------------------- #


ADDRESS_BOOK = [
    AddressBookContact(id="o1", display_name="Sam Patel", company_name="Acme Kitchens Ltd",
                       emails=["sam@acme.example", "other@acme.example"], business_phones=["020 7946 0000"]),
    AddressBookContact(id="o2", display_name="Howdens", emails=["new@howdens.example"], business_phones=["0161 000 0000"]),
    AddressBookContact(id="o3", display_name="Mum", emails=["mum@example.com"]),
    AddressBookContact(id="o4", display_name="", company_name="", emails=["x@example.com"]),
]


def test_import_maps_and_never_overwrites() -> None:
    ledger = _ledger([_event(120000, date(2026, 9, 1), "Acme Kitchens", sale=True), _event(3000, date(2026, 9, 2), "Howdens")])
    book = cb.merge_ledger([], ledger)
    howdens = next(c for c in book if c.name == "Howdens")
    howdens.email = "accounts@howdens.example"

    merged, r = cb.import_address_book(book, ADDRESS_BOOK, add_new=False)
    assert r.model_dump() == {"read": 4, "matched": 2, "updated": 2, "added": 0, "skipped": 2}
    acme = next(c for c in merged if c.normalised_key == "acme-kitchens")
    assert acme.email == "sam@acme.example" and acme.phone == "020 7946 0000"  # matched by company name
    h = next(c for c in merged if c.name == "Howdens")
    assert h.email == "accounts@howdens.example" and h.phone == "0161 000 0000"  # existing email kept
    assert len(merged) == 2

    merged, r = cb.import_address_book(book, ADDRESS_BOOK, add_new=True)
    assert (r.added, r.skipped) == (1, 1)
    mum = next(c for c in merged if c.name == "Mum")
    assert mum.source == "outlook" and mum.kind == "customer" and mum.email == "mum@example.com"


# --------------------------------------------------------------------------- #
# Routes, against fake storage and a fake address book
# --------------------------------------------------------------------------- #


@pytest.fixture
def fake() -> FakeStorageProvider:
    return FakeStorageProvider()


@pytest.fixture
async def books(fake: FakeStorageProvider) -> Books:
    b = Books(fake)
    p = _proposal("Howdens Joinery Ltd", "GB123456789", invoice="H-77")
    await b.save_proposal(p)
    for ev in [_event(24000, date(2026, 9, 30), "Howdens Joinery Ltd", proposal_id=p.id, eid="d" * 32),
               _event(120000, date(2026, 8, 31), "Acme Kitchens Ltd", sale=True, eid="f" * 32)]:
        await b.post_event(ev)
    return b


@pytest.fixture
def address_book() -> FakeContactsProvider:
    return FakeContactsProvider(ADDRESS_BOOK)


@pytest.fixture
def client(books: Books, address_book: FakeContactsProvider):
    from app.main import app
    from app.routes_books import get_books
    from app.routes_contacts import get_address_book

    async def override_books() -> Books:
        return books

    async def override_address_book() -> FakeContactsProvider:
        return address_book

    app.dependency_overrides[get_books] = override_books
    app.dependency_overrides[get_address_book] = override_address_book
    yield TestClient(app)
    app.dependency_overrides.clear()


def _audit(fake: FakeStorageProvider) -> list[str]:
    return [json.loads(x)["action"] for p, d in sorted(fake.files.items()) if p.startswith(f"{AUDIT}/") for x in d.decode().splitlines()]


def test_contact_routes(client: TestClient, fake: FakeStorageProvider, address_book: FakeContactsProvider) -> None:
    r = client.get("/contacts?year_start=2026-04-06", headers=H)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["saved"] is False and cb.CONTACTS not in fake.files  # a GET writes nothing
    by_name = {c["name"]: c for c in body["contacts"]}
    howdens, acme = by_name["Howdens Joinery Ltd"], by_name["Acme Kitchens Ltd"]
    assert howdens["vrn"] == "GB123456789" and howdens["totals"]["open_payable_pence"] == 24000
    assert acme["kind"] == "customer" and acme["totals"]["open_receivable_pence"] == 120000
    assert acme["totals"]["ytd_sales_net_pence"] == 100000

    assert client.post("/contacts/refresh", headers={"X-Graph-Token": "tok"}).status_code == 400  # X-User needed
    r = client.post("/contacts/refresh", headers=H)
    assert r.status_code == 200 and r.json()["saved"] is True
    stored = json.loads(fake.files[cb.CONTACTS])
    assert {c["name"] for c in stored} == {"Howdens Joinery Ltd", "Acme Kitchens Ltd"}
    assert all("totals" not in c and "balance_pence" not in json.dumps(c) for c in stored)  # no money stored

    r = client.put(f"/contacts/{acme['id']}", headers=H, json={"phone": "01234 567890", "address": "1 High St"})
    assert r.status_code == 200 and r.json()["phone"] == "01234 567890" and sorted(r.json()["edited"]) == ["address", "phone"]
    assert client.put(f"/contacts/{acme['id']}", headers=H, json={"email": "nope"}).status_code == 400
    assert client.put("/contacts/0123456789abcdef", headers=H, json={"phone": "1"}).status_code == 404
    assert client.put("/contacts/..%2F..", headers=H, json={"phone": "1"}).status_code == 404

    st = client.get(f"/contacts/{howdens['id']}/statement?from=2026-09-01&to=2026-12-31", headers=H).json()
    assert [(ln["details"], ln["reference"], ln["balance_pence"]) for ln in st["lines"]] == [("Purchase", "H-77", -24000)]
    assert client.get(f"/contacts/{howdens['id']}/statement?from=2026-10-01&to=2026-09-01", headers=H).status_code == 400
    assert client.get("/contacts/0123456789abcdef/statement", headers=H).status_code == 404

    aged = client.get("/contacts/aged?as_of=2026-10-27", headers=H).json()
    assert aged["creditors"]["buckets"]["0-30"] == 24000 and aged["debtors"]["buckets"]["31-60"] == 120000

    r = client.post("/contacts/import-outlook", headers=H, json={"add_new": False})
    assert r.status_code == 200 and r.json()["matched"] == 1 and address_book.calls == 1
    after = {c["name"]: c for c in json.loads(fake.files[cb.CONTACTS])}
    assert after["Acme Kitchens Ltd"]["email"] == "sam@acme.example"
    assert after["Acme Kitchens Ltd"]["phone"] == "01234 567890"  # the user's edit is kept
    assert _audit(fake) == ["contacts_refreshed", "contact_edited", "contacts_imported"]


async def test_books_xlsx_uses_saved_contacts(books: Books, fake: FakeStorageProvider) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    contacts, _ = await cb.refresh(books, USER)
    howdens = next(c for c in contacts if c.vrn)
    await cb.edit(books, howdens.id, {"name": "Howdens"}, USER)
    await books.rebuild()
    wb = load_workbook(BytesIO(fake.files["Books.xlsx"]))
    sup = {r[0]: r[1:] for r in wb["Suppliers"].iter_rows(min_row=2, values_only=True)}
    assert sup["Howdens"][0] == "GB123456789" and sup["Howdens"][3] == 240  # unpaid, on account
    cus = {r[0]: r[1:] for r in wb["Customers"].iter_rows(min_row=2, values_only=True)}
    assert cus["Acme Kitchens Ltd"][3] == 1200
